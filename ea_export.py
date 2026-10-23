
from dataclasses import dataclass, field
from datetime import date
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side



@dataclass
class EaIncomeRow:
    number: str
    date: date
    contact: str
    subject: str
    net: float
    vat: float
    gross: float
    vat_by_rate: dict = field(default_factory=dict)


@dataclass
class EaExpenseRow:
    number: str
    date: date
    supplier: str
    subject: str
    category_code: str
    category_name: str
    ear_group: str
    net: float
    vat: float
    gross: float
    vat_rate: float
    vat_deductible: float



def _r(x) -> float:
    return round(float(x), 2)


def compute_ea_summary(income_rows, expense_rows) -> dict:
    income_net = _r(sum(r.net for r in income_rows))
    income_vat = _r(sum(r.vat for r in income_rows))
    income_gross = _r(sum(r.gross for r in income_rows))

    inc_rate_net: dict = {}
    inc_rate_vat: dict = {}
    for r in income_rows:
        for rate, parts in r.vat_by_rate.items():
            k = _r(rate)
            inc_rate_net[k] = inc_rate_net.get(k, 0.0) + float(parts.get("net", 0.0))
            inc_rate_vat[k] = inc_rate_vat.get(k, 0.0) + float(parts.get("vat", 0.0))
    income_by_rate = [
        {"rate": k, "net": _r(inc_rate_net[k]), "vat": _r(inc_rate_vat[k])}
        for k in sorted(inc_rate_net)
    ]

    expense_net = _r(sum(r.net for r in expense_rows))
    expense_vat = _r(sum(r.vat for r in expense_rows))
    expense_gross = _r(sum(r.gross for r in expense_rows))
    expense_vorsteuer = _r(sum(r.vat_deductible for r in expense_rows))

    exp_rate_net: dict = {}
    exp_rate_vat: dict = {}
    exp_rate_vst: dict = {}
    for r in expense_rows:
        k = _r(r.vat_rate)
        exp_rate_net[k] = exp_rate_net.get(k, 0.0) + r.net
        exp_rate_vat[k] = exp_rate_vat.get(k, 0.0) + r.vat
        exp_rate_vst[k] = exp_rate_vst.get(k, 0.0) + r.vat_deductible
    expense_by_rate = [
        {"rate": k, "net": _r(exp_rate_net[k]), "vat": _r(exp_rate_vat[k]),
         "vorsteuer": _r(exp_rate_vst[k])}
        for k in sorted(exp_rate_net)
    ]

    cat_acc: dict = {}
    for r in expense_rows:
        key = r.category_code or "(ohne)"
        acc = cat_acc.setdefault(key, {
            "code": key,
            "name": r.category_name or "Ohne Kategorie",
            "ear_group": r.ear_group or "(ohne)",
            "net": 0.0, "vat": 0.0, "gross": 0.0, "count": 0,
        })
        acc["net"] += r.net
        acc["vat"] += r.vat
        acc["gross"] += r.gross
        acc["count"] += 1
    expense_by_category = []
    for acc in cat_acc.values():
        expense_by_category.append({
            "code": acc["code"], "name": acc["name"], "ear_group": acc["ear_group"],
            "net": _r(acc["net"]), "vat": _r(acc["vat"]), "gross": _r(acc["gross"]),
            "count": acc["count"],
        })
    expense_by_category.sort(key=lambda a: (a["ear_group"], -a["net"]))

    grp_acc: dict = {}
    for a in expense_by_category:
        g = grp_acc.setdefault(a["ear_group"], {"ear_group": a["ear_group"],
                                                "net": 0.0, "gross": 0.0})
        g["net"] += a["net"]
        g["gross"] += a["gross"]
    expense_by_ear_group = [
        {"ear_group": g["ear_group"], "net": _r(g["net"]), "gross": _r(g["gross"])}
        for g in sorted(grp_acc.values(), key=lambda x: x["ear_group"])
    ]

    return {
        "income_count": len(income_rows),
        "income_net": income_net,
        "income_vat": income_vat,
        "income_gross": income_gross,
        "income_by_rate": income_by_rate,
        "expense_count": len(expense_rows),
        "expense_net": expense_net,
        "expense_vat": expense_vat,
        "expense_gross": expense_gross,
        "expense_vorsteuer": expense_vorsteuer,
        "expense_by_rate": expense_by_rate,
        "expense_by_category": expense_by_category,
        "expense_by_ear_group": expense_by_ear_group,
        "ueberschuss_net": _r(income_net - expense_net),
        "ueberschuss_gross": _r(income_gross - expense_gross),
        "ust_zahllast": _r(income_vat - expense_vorsteuer),
    }



_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="00707A", end_color="00707A", fill_type="solid")
_TITLE_FONT = Font(bold=True, size=13, color="00707A")
_LABEL_FONT = Font(bold=True, size=11)
_SUBTOTAL_FONT = Font(bold=True, size=11)
_SUBTOTAL_FILL = PatternFill(start_color="F1EFE9", end_color="F1EFE9", fill_type="solid")
_RESULT_FILL = PatternFill(start_color="FDF6E3", end_color="FDF6E3", fill_type="solid")
_EUR = "#,##0.00 €"
_THIN = Side(style="thin", color="D9D5CC")
_BORDER = Border(bottom=_THIN)


def _pct(rate: float) -> str:
    r = float(rate)
    return f"{int(r)} %" if r == int(r) else f"{r:g} %"


def _autowidth(ws, max_width=44):
    for col in ws.columns:
        length = 0
        letter = None
        for cell in col:
            if letter is None:
                letter = cell.column_letter
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        if letter:
            ws.column_dimensions[letter].width = min(length + 2, max_width)


def _header_row(ws, headers, row=1):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL


def build_ea_workbook(income_rows, expense_rows, summary, von: date, bis: date) -> Workbook:
    wb = Workbook()

    ws = wb.active
    ws.title = "Einnahmen"
    headers = ["Beleg-Nr", "Zahlungsdatum", "Kunde", "Betreff", "Netto", "USt", "Brutto"]
    _header_row(ws, headers)
    rows_sorted = sorted(income_rows, key=lambda r: (r.date or date.min, r.number))
    ridx = 2
    for r in rows_sorted:
        ws.cell(row=ridx, column=1, value=r.number)
        ws.cell(row=ridx, column=2, value=r.date.isoformat() if r.date else None)
        ws.cell(row=ridx, column=3, value=r.contact)
        ws.cell(row=ridx, column=4, value=r.subject)
        ws.cell(row=ridx, column=5, value=r.net).number_format = _EUR
        ws.cell(row=ridx, column=6, value=r.vat).number_format = _EUR
        ws.cell(row=ridx, column=7, value=r.gross).number_format = _EUR
        ridx += 1
    ws.cell(row=ridx, column=4, value="Summe Einnahmen").font = _SUBTOTAL_FONT
    for col, val in ((5, summary["income_net"]), (6, summary["income_vat"]), (7, summary["income_gross"])):
        c = ws.cell(row=ridx, column=col, value=val)
        c.font = _SUBTOTAL_FONT
        c.fill = _SUBTOTAL_FILL
        c.number_format = _EUR
    _autowidth(ws)

    wa = wb.create_sheet("Ausgaben")
    headers = ["Beleg-Nr", "Zahlungsdatum", "Lieferant", "Betreff", "Kategorie",
               "EAR-Gruppe", "Netto", "USt-Satz", "USt", "Vorsteuer abz.", "Brutto"]
    _header_row(wa, headers)
    rows_sorted = sorted(expense_rows, key=lambda r: (r.date or date.min, r.number))
    ridx = 2
    for r in rows_sorted:
        wa.cell(row=ridx, column=1, value=r.number)
        wa.cell(row=ridx, column=2, value=r.date.isoformat() if r.date else None)
        wa.cell(row=ridx, column=3, value=r.supplier)
        wa.cell(row=ridx, column=4, value=r.subject)
        wa.cell(row=ridx, column=5, value=r.category_name)
        wa.cell(row=ridx, column=6, value=r.ear_group)
        wa.cell(row=ridx, column=7, value=r.net).number_format = _EUR
        wa.cell(row=ridx, column=8, value=_pct(r.vat_rate))
        wa.cell(row=ridx, column=9, value=r.vat).number_format = _EUR
        wa.cell(row=ridx, column=10, value=r.vat_deductible).number_format = _EUR
        wa.cell(row=ridx, column=11, value=r.gross).number_format = _EUR
        ridx += 1
    wa.cell(row=ridx, column=4, value="Summe Ausgaben").font = _SUBTOTAL_FONT
    for col, val in ((7, summary["expense_net"]), (9, summary["expense_vat"]),
                     (10, summary["expense_vorsteuer"]), (11, summary["expense_gross"])):
        c = wa.cell(row=ridx, column=col, value=val)
        c.font = _SUBTOTAL_FONT
        c.fill = _SUBTOTAL_FILL
        c.number_format = _EUR
    _autowidth(wa)

    ws3 = wb.create_sheet("Summen")
    _build_summary_tab(ws3, summary, von, bis)

    return wb


def _money(ws, row, col, val):
    c = ws.cell(row=row, column=col, value=val)
    c.number_format = _EUR
    return c


def _build_summary_tab(ws, summary, von: date, bis: date):
    t = ws.cell(row=1, column=1, value="Einnahmen-Ausgaben-Uebersicht")
    t.font = _TITLE_FONT
    ws.cell(row=2, column=1,
            value=f"Zeitraum {von.isoformat()} bis {bis.isoformat()} "
                  f"(Zufluss-/Abfluss-Prinzip, nach Zahlungsdatum)").font = Font(italic=True, size=10, color="6B6455")

    r = 4
    ws.cell(row=r, column=1, value="Einnahmen").font = _LABEL_FONT
    r += 1
    ws.cell(row=r, column=1, value=f"Anzahl Belege")
    ws.cell(row=r, column=2, value=summary["income_count"])
    r += 1
    for label, key in (("Netto", "income_net"), ("USt", "income_vat"), ("Brutto", "income_gross")):
        ws.cell(row=r, column=1, value=label)
        _money(ws, r, 2, summary[key])
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="USt nach Satz (Einnahmen)").font = _LABEL_FONT
    r += 1
    _header_row(ws, ["Satz", "Bemessungsgrundlage (Netto)", "USt"], row=r)
    r += 1
    for e in summary["income_by_rate"]:
        ws.cell(row=r, column=1, value=_pct(e["rate"]))
        _money(ws, r, 2, e["net"])
        _money(ws, r, 3, e["vat"])
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Ausgaben").font = _LABEL_FONT
    r += 1
    ws.cell(row=r, column=1, value="Anzahl Belege")
    ws.cell(row=r, column=2, value=summary["expense_count"])
    r += 1
    for label, key in (("Netto", "expense_net"), ("USt", "expense_vat"),
                       ("Vorsteuer abziehbar", "expense_vorsteuer"), ("Brutto", "expense_gross")):
        ws.cell(row=r, column=1, value=label)
        _money(ws, r, 2, summary[key])
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Vorsteuer nach Satz (Ausgaben)").font = _LABEL_FONT
    r += 1
    _header_row(ws, ["Satz", "Bemessungsgrundlage (Netto)", "USt", "Vorsteuer abz."], row=r)
    r += 1
    for e in summary["expense_by_rate"]:
        ws.cell(row=r, column=1, value=_pct(e["rate"]))
        _money(ws, r, 2, e["net"])
        _money(ws, r, 3, e["vat"])
        _money(ws, r, 4, e["vorsteuer"])
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Ausgaben nach Kategorie").font = _LABEL_FONT
    r += 1
    _header_row(ws, ["EAR-Gruppe", "Kategorie", "Anzahl", "Netto", "USt", "Brutto"], row=r)
    r += 1
    current_group = None
    group_lookup = {g["ear_group"]: g for g in summary["expense_by_ear_group"]}
    for a in summary["expense_by_category"]:
        if current_group is not None and a["ear_group"] != current_group:
            g = group_lookup.get(current_group)
            r = _write_group_subtotal(ws, r, current_group, g)
        current_group = a["ear_group"]
        ws.cell(row=r, column=1, value=a["ear_group"])
        ws.cell(row=r, column=2, value=a["name"])
        ws.cell(row=r, column=3, value=a["count"])
        _money(ws, r, 4, a["net"])
        _money(ws, r, 5, a["vat"])
        _money(ws, r, 6, a["gross"])
        r += 1
    if current_group is not None:
        r = _write_group_subtotal(ws, r, current_group, group_lookup.get(current_group))

    r += 1
    ws.cell(row=r, column=1, value="Ergebnis").font = _LABEL_FONT
    r += 1
    for label, key in (("Ueberschuss (Netto) = EA-Bemessungsgrundlage", "ueberschuss_net"),
                       ("Ueberschuss (Brutto) = Cash-Ueberschuss", "ueberschuss_gross"),
                       ("USt-Zahllast (USt minus Vorsteuer)", "ust_zahllast")):
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = _SUBTOTAL_FONT
        lc.fill = _RESULT_FILL
        vc = _money(ws, r, 2, summary[key])
        vc.font = _SUBTOTAL_FONT
        vc.fill = _RESULT_FILL
        r += 1

    _autowidth(ws, max_width=52)


def _write_group_subtotal(ws, r, group_name, g):
    if not g:
        return r
    lc = ws.cell(row=r, column=2, value=f"Zwischensumme {group_name}")
    lc.font = _SUBTOTAL_FONT
    lc.fill = _SUBTOTAL_FILL
    nc = _money(ws, r, 4, g["net"])
    nc.font = _SUBTOTAL_FONT
    nc.fill = _SUBTOTAL_FILL
    gc = _money(ws, r, 6, g["gross"])
    gc.font = _SUBTOTAL_FONT
    gc.fill = _SUBTOTAL_FILL
    return r + 1
