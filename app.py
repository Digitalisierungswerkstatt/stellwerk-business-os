import calendar
import io
import os
import re
import secrets
import subprocess
import sys
import shutil
import tempfile
import uuid
from collections import defaultdict
from datetime import datetime, date, timedelta, timezone
from types import SimpleNamespace
from urllib.parse import urlparse
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file, make_response, abort
from markupsafe import Markup, escape
from sqlalchemy import event, func
from sqlalchemy.engine import Engine
from sqlalchemy.orm import joinedload
from werkzeug.utils import secure_filename

import search_fts
import ea_export

from models import (db, Contact, Task, TaskComment, task_tags, Interaction, EmailTemplate, StatusChange,
                    SalesGoal, FollowUpRule, TimeEntry, WorkDay, ExpectedPayment,
                    Document, LineItemGroup, LineItem, NumberSequence,
                    ServiceTemplate, GroupTemplate, GroupTemplateItem, ConditionsTemplate,
                    PIPELINE_STAGES, DEAD_PIPELINE_STAGES, ANREDE_OPTIONS, ANREDE_EN_OPTIONS,
                    KONTAKT_WIE_OPTIONS, SPRACHE_OPTIONS, DEAL_PROBABILITY_DEFAULTS,
                    LOST_REASONS, QUELLEN_OPTIONS, BRANCHEN_OPTIONS, PRIORITAET_OPTIONS,
                    TIER_OPTIONS, PIPELINE_PHASE_OPTIONS,
                    FOLLOWUP_DEFAULTS, GOAL_PERIOD_TYPES,
                    TAETIGKEIT_KONTAKT_OPTIONS, TAETIGKEIT_ALLGEMEIN_OPTIONS,
                    DOC_TYPES, QUOTE_STATUSES, QUOTE_WON_STATUSES, INVOICE_STATUSES,
                    UNIT_OPTIONS, BILLING_TYPES, DUNNING_LEVELS,
                    BOARD_STAGES, BOARD_PRIORITY_OPTIONS, BOARD_EFFORT_OPTIONS,
                    BOARD_IMPACT_OPTIONS, CREATED_BY_OPTIONS, RECURRENCE_PRESETS,
                    TASK_STALE_THRESHOLD_DAYS, BOARD_SCHEDULED_LEAD_DAYS,
                    SERVICE_CATEGORIES, PAYMENT_TYPES, PAYMENT_STATUSES,
                    DEFAULT_HOURLY_RATE,
                    CompanySettings, DocumentStatusLog, Retainer,
                    WikiPage, WikiEntry, WikiFileLink,
                    LEISTUNGSBEREICHE, ENTRY_TYPES, ENTRY_STATUSES,
                    AccountToolType, CustomerAccount,
                    ACCOUNT_STATUSES, ACCOUNT_CATEGORIES,
                    IncomingInvoice, IncomingLineItem, ExpenseCategory,
                    CONTACT_TYPES, INCOMING_INVOICE_STATUSES, PAYMENT_METHODS,
                    BELEG_STATUSES,
                    BankImport, BankTransaction, BANK_MATCH_STATUSES)
from database import init_db, import_from_excel, get_pipeline_stats
import stellwerk_write
from bank_import import parse_csv, BankFormatError

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE_PATH = os.environ.get('STELLWERK_DB_PATH', os.path.join(BASE_DIR, 'data', 'crm.db'))

SYSTEM_CHROME_CANDIDATES = [
    r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe',
    r'C:\Program Files\Google\Chrome\Application\chrome.exe',
]


def _resolve_chrome_path():
    candidates = list(SYSTEM_CHROME_CANDIDATES)

    cache_root = os.path.join(os.path.expanduser('~'), '.cache', 'puppeteer', 'chrome')
    if os.path.isdir(cache_root):
        def version_key(name):
            match = re.search(r'-(\d[\d.]*)$', name)
            return [int(part) for part in match.group(1).split('.')] if match else []

        try:
            builds = sorted(os.listdir(cache_root), key=version_key, reverse=True)
        except OSError:
            builds = []
        candidates += [os.path.join(cache_root, b, 'chrome-win64', 'chrome.exe')
                       for b in builds]

    return next((path for path in candidates if os.path.exists(path)), None)


def _export_chrome_path():
    if os.environ.get('PUPPETEER_EXECUTABLE_PATH'):
        return
    chrome = _resolve_chrome_path()
    if chrome:
        os.environ['PUPPETEER_EXECUTABLE_PATH'] = chrome


_export_chrome_path()

def _secret_key():
    key = os.environ.get('STELLWERK_SECRET_KEY')
    if key:
        return key
    pfad = os.path.join(BASE_DIR, 'data', '.secret_key')
    if os.path.exists(pfad):
        with open(pfad, encoding='utf-8') as f:
            key = f.read().strip()
    if not key:
        key = secrets.token_hex(32)
        os.makedirs(os.path.dirname(pfad), exist_ok=True)
        with open(pfad, 'w', encoding='utf-8') as f:
            f.write(key)
    return key


app = Flask(__name__)
app.config['SECRET_KEY'] = _secret_key()
app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{DATABASE_PATH}'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['TEMPLATES_AUTO_RELOAD'] = True


_LOKALE_HOSTS = {'localhost', '127.0.0.1', '::1'}


def _lokaler_host(netloc):
    try:
        return (urlparse('//' + netloc).hostname or '').lower() in _LOKALE_HOSTS
    except ValueError:
        return False


@app.before_request
def _herkunft_pruefen():
    if not _lokaler_host(request.host):
        abort(403)
    if request.method in ('GET', 'HEAD', 'OPTIONS'):
        return None
    site = request.headers.get('Sec-Fetch-Site')
    if site is not None:
        if site not in ('same-origin', 'none'):
            abort(403)
        return None
    origin = request.headers.get('Origin')
    if origin is not None and origin.rstrip('/') != request.host_url.rstrip('/'):
        abort(403)
    return None


@app.after_request
def _nur_selbst_einbettbar(resp):
    resp.headers.setdefault('X-Frame-Options', 'SAMEORIGIN')
    resp.headers.setdefault('Content-Security-Policy', "frame-ancestors 'self'")
    resp.headers.setdefault('X-Content-Type-Options', 'nosniff')
    return resp


@event.listens_for(Engine, 'connect')
def _set_sqlite_pragmas(dbapi_connection, connection_record):
    cursor = dbapi_connection.cursor()
    cursor.execute('PRAGMA foreign_keys=ON')
    cursor.execute('PRAGMA journal_mode=WAL')
    cursor.close()

init_db(app)

from blueprints.services import services_bp
app.register_blueprint(services_bp)
from blueprints.content_calendar import content_calendar_bp
app.register_blueprint(content_calendar_bp)


def _get_default_hourly_rate():
    return float(CompanySettings.get('default_hourly_rate') or DEFAULT_HOURLY_RATE)

def _auto_backup():
    if os.environ.get('STELLWERK_SKIP_BACKUP'):
        return
    from pathlib import Path
    from backup import PROD_DB_PATH, create_backup, latest_backup
    db_path = Path(DATABASE_PATH)
    if not db_path.exists() or db_path.resolve() != PROD_DB_PATH.resolve():
        return
    latest = latest_backup()
    if latest is None or datetime.fromtimestamp(latest.stat().st_mtime).date() < date.today():
        create_backup()

try:
    _auto_backup()
except Exception as _backup_err:
    print(f'Auto-Backup uebersprungen: {_backup_err}', file=sys.stderr)


def _autopost_on_start():
    from content_publisher import run_due
    with app.app_context():
        try:
            import content_outbox
            content_outbox.print_report(content_outbox.sync(lock_wait=10))
        except Exception as _outbox_err:
            print(f'  Cloud-Abgleich uebersprungen: {_outbox_err}', file=sys.stderr)
        if CompanySettings.get('content_autopost_enabled') != '1':
            return
        results = run_due(dry_run=False)
    if not results:
        return
    ok = sum(1 for _e, good, _m in results if good)
    print(f'  Content-Poster: {ok} von {len(results)} Eintraegen veroeffentlicht')
    for entry, good, message in results:
        if not good:
            print(f'    FEHL #{entry.id} {entry.title[:40]}: {message}', file=sys.stderr)



@app.errorhandler(404)
def not_found(e):
    is_json = (request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html) \
              or request.content_type == 'application/json'
    if is_json:
        return jsonify({'success': False, 'error': 'Nicht gefunden'}), 404
    flash('Die angeforderte Seite wurde nicht gefunden.', 'error')
    return redirect(url_for('dashboard'))


@app.errorhandler(500)
def server_error(e):
    db.session.rollback()
    if request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html:
        return jsonify({'error': 'Interner Serverfehler'}), 500
    flash('Ein interner Fehler ist aufgetreten. Bitte versuche es erneut.', 'error')
    return redirect(url_for('dashboard'))



@app.template_filter('format_euro')
def format_euro_filter(value):
    if value is None:
        return '0,00'
    val = float(value)
    return f'{val:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


@app.template_filter('format_decimal')
def format_decimal_filter(value):
    if value is None:
        return '0,00'
    val = float(value)
    return f'{val:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')



def _generate_recurring_time_entries():
    today = date.today()
    recurring = TimeEntry.query.filter(
        TimeEntry.recurring_days.isnot(None),
        TimeEntry.timer_start.is_(None)
    ).all()
    created = 0
    for te in recurring:
        last = te.last_recurred or te.datum
        while last + timedelta(days=te.recurring_days) <= today:
            next_date = last + timedelta(days=te.recurring_days)
            exists = TimeEntry.query.filter_by(
                contact_id=te.contact_id, datum=next_date, taetigkeit=te.taetigkeit,
                dauer_minuten=te.dauer_minuten
            ).first()
            if not exists:
                new_entry = TimeEntry(
                    contact_id=te.contact_id, datum=next_date,
                    dauer_minuten=te.dauer_minuten, taetigkeit=te.taetigkeit,
                    notiz=f'[Wiederkehrend] {te.notiz or ""}'.strip()
                )
                db.session.add(new_entry)
                created += 1
            last = next_date
        te.last_recurred = last
    if created:
        db.session.commit()
    return created



def _set_followup_reminders():
    rules = {r.stage: r for r in FollowUpRule.query.filter_by(enabled=True).all()}
    if not rules:
        return 0

    today = date.today()
    set_count = 0

    for stage, rule in rules.items():
        if stage in ('won', 'lost', 'on_hold', 'closed_no_response'):
            continue
        threshold_date = today - timedelta(days=rule.days_threshold)

        contacts = Contact.query.filter(
            Contact.status == stage,
            db.or_(Contact.kontaktiert < threshold_date, Contact.kontaktiert.is_(None))
        ).all()

        for contact in contacts:
            if contact.erinnerung and contact.erinnerung >= today:
                continue

            stage_label = dict(PIPELINE_STAGES).get(stage, stage)
            if contact.kontaktiert:
                days_ago = (today - contact.kontaktiert).days
                text = f'Follow-up: {stage_label} seit {days_ago} Tagen offen'
            else:
                text = f'Follow-up: {stage_label} ohne Erstkontakt'

            contact.erinnerung = today + timedelta(days=3)
            contact.erinnerungstext = text
            set_count += 1

    if set_count:
        db.session.commit()
    return set_count


_generate_followup_tasks = _set_followup_reminders



@app.route('/')
def dashboard():
    followup_created = _set_followup_reminders()
    _update_document_statuses()
    _generate_recurring_time_entries()

    stats = get_pipeline_stats()
    total_contacts = Contact.query.count()
    today = date.today()

    dash_periode = request.args.get('periode', 'monat')
    if dash_periode == 'monat':
        dash_period_start = today.replace(day=1)
        dash_period_label = today.strftime('%B %Y')
    elif dash_periode == 'quartal':
        q_month = ((today.month - 1) // 3) * 3 + 1
        dash_period_start = today.replace(month=q_month, day=1)
        q_num = (today.month - 1) // 3 + 1
        dash_period_label = f'Q{q_num} {today.year}'
    elif dash_periode == 'jahr':
        dash_period_start = today.replace(month=1, day=1)
        dash_period_label = str(today.year)
    else:
        dash_period_start = None
        dash_period_label = 'Gesamt'

    if dash_period_start:
        dash_start_dt = datetime.combine(dash_period_start, datetime.min.time())
        dash_new_contacts = Contact.query.filter(Contact.created_at >= dash_start_dt).count()
        dash_won_changes = StatusChange.query.filter(
            StatusChange.to_status == 'won', StatusChange.changed_at >= dash_start_dt
        ).all()
        dash_won_ids = [sc.contact_id for sc in dash_won_changes]
        dash_won_contacts = Contact.query.filter(
            Contact.id.in_(dash_won_ids), Contact.status == 'won'
        ).all() if dash_won_ids else []
        dash_won_value = sum(c.deal_value or 0 for c in dash_won_contacts)
        dash_won_count = len(dash_won_contacts)
        dash_lost_changes = StatusChange.query.filter(
            StatusChange.to_status == 'lost', StatusChange.changed_at >= dash_start_dt
        ).all()
        dash_lost_count = len(set(sc.contact_id for sc in dash_lost_changes))
        dash_time_minutes = db.session.query(
            func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
        ).filter(TimeEntry.datum >= dash_period_start, TimeEntry.timer_start.is_(None)).scalar() or 0
        dash_paid_invoices = Document.query.filter(
            Document.doc_type == 'invoice', Document.status == 'paid',
            Document.payment_date >= dash_period_start
        ).all()
        dash_revenue = sum(d.total_net for d in dash_paid_invoices)
        dash_interactions = Interaction.query.filter(Interaction.datum >= dash_period_start).count()
    else:
        dash_new_contacts = None
        dash_won_value = None
        dash_won_count = None
        dash_lost_count = None
        dash_time_minutes = None
        dash_revenue = None
        dash_interactions = None

    overdue_reminders = Contact.query.filter(
        Contact.erinnerung < today,
        Contact.erinnerung.isnot(None)
    ).order_by(Contact.erinnerung).limit(5).all()

    today_reminders = Contact.query.filter(
        Contact.erinnerung == today
    ).all()

    week_later = today + timedelta(days=7)
    upcoming_reminders = Contact.query.filter(
        Contact.erinnerung > today,
        Contact.erinnerung <= week_later
    ).order_by(Contact.erinnerung).limit(5).all()

    open_tasks = Task.query.filter(Task.status != 'erledigt').order_by(Task.due_date).limit(5).all()

    recent_contacts = Contact.query.order_by(Contact.updated_at.desc()).limit(5).all()

    recent_activity = Interaction.query.order_by(
        Interaction.datum.desc(), Interaction.created_at.desc()
    ).limit(8).all()

    active_contacts = Contact.query.filter(Contact.status.notin_(['lost'])).count()

    active_deals = Contact.query.filter(Contact.status.notin_(['lost', 'won']), Contact.deal_value.isnot(None)).all()
    pipeline_value = sum(c.deal_value or 0 for c in active_deals)
    weighted_forecast = sum(c.weighted_value for c in active_deals)
    won_contacts = Contact.query.filter(Contact.status == 'won').all()
    won_value = sum(c.deal_value or 0 for c in won_contacts if c.deal_value)

    all_with_rate = Contact.query.filter(Contact.hourly_rate.isnot(None)).all()
    total_time_value = sum(c.time_value for c in all_with_rate)

    upcoming_payments = ExpectedPayment.query.filter(
        ExpectedPayment.status == 'pending'
    ).order_by(ExpectedPayment.expected_date).all()
    expected_payments_total = sum(p.amount for p in upcoming_payments)
    overdue_payments = [p for p in upcoming_payments if p.expected_date < today]
    overdue_payments_total = sum(p.amount for p in overdue_payments)

    lost_reasons_stats = db.session.query(
        Contact.lost_reason, func.count(Contact.id)
    ).filter(Contact.status == 'lost', Contact.lost_reason.isnot(None)).group_by(Contact.lost_reason).all()
    lost_reasons_dict = dict(lost_reasons_stats)

    quellen_stats = db.session.query(
        Contact.quelle, func.count(Contact.id)
    ).filter(Contact.quelle.isnot(None), Contact.quelle != '').group_by(Contact.quelle).all()
    quellen_dict = dict(quellen_stats)

    conversion_rates = _calc_conversion_rates()

    hour = datetime.now().hour


    today_time_minutes = db.session.query(func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)).filter(
        TimeEntry.datum == today, TimeEntry.timer_start.is_(None)
    ).scalar()
    active_timer = _get_active_timer()

    open_invoices = Document.query.filter(
        Document.doc_type == 'invoice',
        Document.status.in_(['sent', 'overdue'])
    ).all()
    open_invoice_count = len(open_invoices)
    open_invoice_sum = sum(d.total_gross for d in open_invoices)
    overdue_invoices = [d for d in open_invoices if d.status == 'overdue']
    overdue_invoice_count = len(overdue_invoices)
    overdue_invoice_sum = sum(d.total_gross for d in overdue_invoices)

    open_quotes = Document.query.filter(
        Document.doc_type == 'quote',
        Document.status.in_(['draft', 'sent'])
    ).all()
    open_quote_count = len(open_quotes)
    open_quote_sum = sum(d.total_net for d in open_quotes)
    expiring_quotes = [d for d in open_quotes
                       if d.valid_until and d.status == 'sent'
                       and 0 <= (d.valid_until - today).days <= 5]

    _generate_quote_followup_tasks()

    monthly_revenue = []
    for i in range(5, -1, -1):
        m_start = (today.replace(day=1) - timedelta(days=30 * i)).replace(day=1)
        if m_start.month == 12:
            m_end = m_start.replace(year=m_start.year + 1, month=1, day=1) - timedelta(days=1)
        else:
            m_end = m_start.replace(month=m_start.month + 1, day=1) - timedelta(days=1)
        paid = Document.query.filter(
            Document.doc_type == 'invoice',
            Document.status == 'paid',
            Document.payment_date >= m_start,
            Document.payment_date <= m_end,
        ).all()
        total = sum(d.total_net for d in paid)
        monthly_revenue.append({
            'month': m_start.strftime('%b %Y'),
            'month_short': m_start.strftime('%b'),
            'total': total,
            'count': len(paid),
        })
    max_revenue = max((m['total'] for m in monthly_revenue), default=1) or 1

    active_retainers = Retainer.query.filter_by(active=True).all()
    mrr_retainers = sum(r.monthly_fee for r in active_retainers)
    retainer_contact_ids = {r.contact_id for r in active_retainers}
    recurring_payments = [
        p for p in ExpectedPayment.query.filter(
            ExpectedPayment.status == 'pending',
            ExpectedPayment.payment_type == 'recurring'
        ).all()
        if p.contact_id not in retainer_contact_ids
    ]
    mrr_payments = sum(p.amount / (p.interval_months or 1) for p in recurring_payments if p.interval_months)
    mrr_total = round(mrr_retainers + mrr_payments, 2)
    mrr_client_count = len(set(
        [r.contact_id for r in active_retainers] +
        [p.contact_id for p in recurring_payments]
    ))

    stale_proposals = []
    sent_quotes = Document.query.filter(
        Document.doc_type == 'quote', Document.status == 'sent'
    ).all()
    for q in sent_quotes:
        sent_d = q.sent_date or q.date
        days_since = (today - sent_d).days
        if days_since > 14:
            stale_proposals.append({
                'doc': q,
                'contact': q.contact,
                'days': days_since,
                'value': q.total_net,
            })
    stale_proposals.sort(key=lambda x: x['days'], reverse=True)

    silent_clients = []
    threshold_date = today - timedelta(days=30)

    last_interactions = dict(db.session.query(
        Interaction.contact_id, func.max(Interaction.datum)
    ).join(Contact, Contact.id == Interaction.contact_id).filter(
        Contact.status == 'won'
    ).group_by(Interaction.contact_id).all())

    last_times = dict(db.session.query(
        TimeEntry.contact_id, func.max(TimeEntry.datum)
    ).join(Contact, Contact.id == TimeEntry.contact_id).filter(
        Contact.status == 'won', TimeEntry.timer_start.is_(None)
    ).group_by(TimeEntry.contact_id).all())

    won_contacts_all = Contact.query.filter(Contact.status == 'won').all()
    for c in won_contacts_all:
        li = last_interactions.get(c.id)
        lt = last_times.get(c.id)
        last_date = max(filter(None, [li, lt]), default=None)
        if last_date and last_date < threshold_date:
            silent_clients.append({
                'contact': c,
                'last_activity': last_date,
                'days_silent': (today - last_date).days,
            })
        elif not last_date and c.kontaktiert and c.kontaktiert < threshold_date:
            silent_clients.append({
                'contact': c,
                'last_activity': c.kontaktiert,
                'days_silent': (today - c.kontaktiert).days,
            })
    silent_clients.sort(key=lambda x: x['days_silent'], reverse=True)

    week_start = today - timedelta(days=today.weekday())
    week_end = week_start + timedelta(days=6)
    week_worked_min = db.session.query(
        func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
    ).filter(TimeEntry.datum >= week_start, TimeEntry.timer_start.is_(None)).scalar() or 0
    week_worked_hours = round(week_worked_min / 60, 1)
    weekly_target = float(CompanySettings.get('weekly_hours_target') or 40)
    retainer_hours_monthly = sum(r.monthly_hours for r in Retainer.query.filter_by(active=True).all())
    retainer_hours_weekly = round(retainer_hours_monthly / 4.33, 1)
    last_week_start = week_start - timedelta(days=7)
    last_week_min = db.session.query(
        func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
    ).filter(
        TimeEntry.datum >= last_week_start, TimeEntry.datum < week_start,
        TimeEntry.timer_start.is_(None)
    ).scalar() or 0
    last_week_hours = round(last_week_min / 60, 1)

    four_weeks_start = week_start - timedelta(days=28)
    four_weeks_min = db.session.query(
        func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
    ).filter(
        TimeEntry.datum >= four_weeks_start, TimeEntry.datum < week_start,
        TimeEntry.timer_start.is_(None)
    ).scalar() or 0
    avg_4w_hours = round(four_weeks_min / 60 / 4, 1)

    capacity = {
        'worked': week_worked_hours,
        'target': weekly_target,
        'pct': round(week_worked_hours / weekly_target * 100, 1) if weekly_target > 0 else 0,
        'retainer_reserved': retainer_hours_weekly,
        'free': round(weekly_target - week_worked_hours, 1),
        'last_week': last_week_hours,
        'avg_4w': avg_4w_hours,
    }

    capacity_projection = compute_capacity_projection(
        active_retainers, today,
        weekly_target_mandate=float(CompanySettings.get('mandate_kapazitaet_wochenstunden') or 20),
    )

    retainer_billing_due = []
    for r in Retainer.query.filter_by(active=True).all():
        last_billed = r.last_billed_date or r.start_date
        months_since = (today.year - last_billed.year) * 12 + today.month - last_billed.month
        if months_since >= r.billing_interval_months:
            retainer_billing_due.append({
                'retainer': r,
                'contact': r.contact,
                'months_since': months_since,
                'expected_amount': round(r.monthly_fee * r.billing_interval_months, 2),
            })

    hidden_widgets = [w.strip() for w in (CompanySettings.get('dashboard_hidden_widgets') or '').split(',') if w.strip()]

    setup_hints = []
    if total_contacts == 0:
        setup_hints.append({'text': 'Noch keine Kontakte angelegt', 'url': url_for('contact_new'), 'action': 'Kontakt erstellen'})
    if CompanySettings.query.count() == 0:
        setup_hints.append({'text': 'Firmendaten noch nicht angepasst', 'url': url_for('followup_settings'), 'action': 'Einstellungen'})
    if EmailTemplate.query.count() == 0:
        setup_hints.append({'text': 'Keine Email-Vorlagen vorhanden', 'url': url_for('email_templates'), 'action': 'Vorlagen erstellen'})

    cashflow = compute_cashflow_forecast(open_invoices, active_retainers)
    _all_invoices = Document.query.filter_by(doc_type='invoice').all()
    _invoiced_by_contact = defaultdict(float)
    for d in _all_invoices:
        if d.status in ('sent', 'overdue', 'paid'):
            _invoiced_by_contact[d.contact_id] += d.total_net
    _won_projects = [
        SimpleNamespace(name=c.organisation, deal_value=c.deal_value,
                        invoiced_net=_invoiced_by_contact.get(c.id, 0.0))
        for c in Contact.query.filter(Contact.status == 'won', Contact.deal_value > 0).all()
        if c.id not in retainer_contact_ids
    ]
    try:
        _plan = float(CompanySettings.get('umsatz_plan_jahr') or 0) or None
    except ValueError:
        _plan = None
    umsatz_prognose = compute_revenue_projection(
        _all_invoices, active_retainers, _won_projects, today, _plan,
        expected_recurring=recurring_payments,
        expected_oneoff=[p for p in upcoming_payments if p.payment_type != 'recurring'])
    pipeline = compute_pipeline(Document.query.filter_by(doc_type='quote').all(), today)
    _billing_rows = _gather_billing_rows(today, _get_default_hourly_rate())
    abrechnung_klaeren = sum(1 for r in _billing_rows if r['ampel'] == 'rot')
    abrechnung_faellig = sum(1 for r in _billing_rows if r['action'] in ('retainer', 'trueup'))

    return render_template('dashboard.html',
                           setup_hints=setup_hints,
                           stats=stats, total_contacts=total_contacts,
                           hidden_widgets=hidden_widgets,
                           overdue_reminders=overdue_reminders,
                           today_reminders=today_reminders,
                           upcoming_reminders=upcoming_reminders,
                           open_tasks=open_tasks,
                           recent_contacts=recent_contacts,
                           recent_activity=recent_activity,
                           active_contacts=active_contacts,
                           pipeline_value=pipeline_value,
                           weighted_forecast=weighted_forecast,
                           won_value=won_value,
                           lost_reasons_dict=lost_reasons_dict,
                           lost_reasons=LOST_REASONS,
                           quellen_dict=quellen_dict,
                           quellen_options=QUELLEN_OPTIONS,
                           conversion_rates=conversion_rates,
                           cashflow=cashflow,
                           umsatz_prognose=umsatz_prognose,
                           pipeline=pipeline,
                           abrechnung_klaeren=abrechnung_klaeren,
                           abrechnung_faellig=abrechnung_faellig,
                           followup_created=followup_created,
                           today_time_minutes=today_time_minutes,
                           active_timer=active_timer,
                           open_invoice_count=open_invoice_count,
                           open_invoice_sum=open_invoice_sum,
                           overdue_invoice_count=overdue_invoice_count,
                           overdue_invoice_sum=overdue_invoice_sum,
                           total_time_value=total_time_value,
                           upcoming_payments=upcoming_payments,
                           expected_payments_total=expected_payments_total,
                           overdue_payments=overdue_payments,
                           overdue_payments_total=overdue_payments_total,
                           open_quote_count=open_quote_count,
                           open_quote_sum=open_quote_sum,
                           expiring_quotes=expiring_quotes,
                           silent_clients=silent_clients,
                           capacity=capacity,
                           capacity_projection=capacity_projection,
                           retainer_billing_due=retainer_billing_due,
                           mrr_total=mrr_total,
                           mrr_client_count=mrr_client_count,
                           stale_proposals=stale_proposals,
                           monthly_revenue=monthly_revenue,
                           max_revenue=max_revenue,
                           format_minutes=_format_minutes,
                           hour=hour,
                           dash_periode=dash_periode,
                           dash_period_label=dash_period_label,
                           dash_new_contacts=dash_new_contacts,
                           dash_won_value=dash_won_value,
                           dash_won_count=dash_won_count,
                           dash_lost_count=dash_lost_count,
                           dash_time_minutes=dash_time_minutes,
                           dash_revenue=dash_revenue,
                           dash_interactions=dash_interactions,
                           stages=PIPELINE_STAGES, today=today)



def _render_document_html(doc, dunning_level=0, dunning_title='', dunning_text=''):
    logo_path = os.path.join(BASE_DIR, 'static', 'logo_base64.txt')
    logo_base64 = ''
    if os.path.exists(logo_path):
        with open(logo_path, 'r') as f:
            logo_base64 = f.read().strip()
    return render_template('pdf/document.html', doc=doc,
                           logo_base64=logo_base64,
                           company=CompanySettings.get_all(),
                           dunning_level=dunning_level,
                           dunning_title=dunning_title,
                           dunning_text=dunning_text)


def _fertiges_pdf_oder_none(pdf_path):
    if not os.path.exists(pdf_path):
        return None
    with open(pdf_path, 'rb') as f:
        daten = f.read()
    if daten[:4] == b'%PDF' and daten.rstrip()[-5:] == b'%%EOF':
        return daten
    return None


def _generate_pdf(html_content, company=None):
    with tempfile.NamedTemporaryFile(mode='w', suffix='.html', delete=False, encoding='utf-8') as f:
        f.write(html_content)
        html_path = f.name
    pdf_path = html_path.replace('.html', '.pdf')

    if company:
        c = company if isinstance(company, dict) else dict(company)
        uid_line = f'<br>UID: {c.get("uid","")}' if c.get("uid") else ''
        footer_html = (
            '<div style="width:100%;font-size:7px;color:#a0aec0;padding:0 20mm;'
            'border-top:1.5px solid #167986;display:flex;justify-content:space-between;gap:10px;line-height:1.5;">'
            '<div style="flex:1;">'
            '<span style="font-weight:700;color:#167986;text-transform:uppercase;letter-spacing:0.5px;">Firma</span><br>'
            f'{c.get("company_name","")}<br>'
            f'{c.get("street","")}, {c.get("zip_city_short","")}<br>'
            f'Steuernr: {c.get("tax_number","")} &middot; {c.get("company_register","")}{uid_line}'
            '</div>'
            '<div style="flex:1;">'
            '<span style="font-weight:700;color:#167986;text-transform:uppercase;letter-spacing:0.5px;">Kontakt</span><br>'
            f'{c.get("phone","")}<br>'
            f'{c.get("email","")}<br>'
            f'{c.get("website","")}'
            '</div>'
            '<div style="flex:1;">'
            '<span style="font-weight:700;color:#167986;text-transform:uppercase;letter-spacing:0.5px;">Bankverbindung</span><br>'
            f'{c.get("bank_name","")} &middot; {c.get("bank_account_holder","")}<br>'
            f'IBAN: {c.get("iban","")}<br>'
            f'BIC: {c.get("bic","")}'
            '</div>'
            '<div style="flex:0 0 auto;align-self:flex-end;text-align:right;">'
            'Seite <span class="pageNumber"></span> von <span class="totalPages"></span>'
            '</div>'
            '</div>'
        )
    else:
        footer_html = '<div style="width:100%;text-align:right;font-size:8px;color:#a0aec0;padding-right:20mm;">Seite <span class="pageNumber"></span> von <span class="totalPages"></span></div>'

    footer_js = footer_html.replace('\\', '\\\\').replace('`', '\\`').replace('$', '\\$')

    try:
        node_script = f"""
        const puppeteer = require('puppeteer');
        (async () => {{
            const browser = await puppeteer.launch({{ headless: true, args: ['--no-sandbox'] }});
            const page = await browser.newPage();
            await page.goto('file:///{html_path.replace(chr(92), "/")}', {{ waitUntil: 'networkidle0' }});
            await page.pdf({{
                path: '{pdf_path.replace(chr(92), "/")}',
                format: 'A4',
                printBackground: true,
                margin: {{ top: '25mm', bottom: '30mm', left: '20mm', right: '20mm' }},
                displayHeaderFooter: true,
                headerTemplate: '<span></span>',
                footerTemplate: `{footer_js}`
            }});
            await Promise.race([
                browser.close(),
                new Promise(resolve => setTimeout(resolve, 3000))
            ]);
            process.exit(0);
        }})();
        """
        process = subprocess.Popen(['node', '-e', node_script],
                                    stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                    text=True, errors='replace')
        try:
            stdout, stderr = process.communicate(timeout=60)
        except subprocess.TimeoutExpired:
            process.kill()
            process.communicate()
            gerettet = _fertiges_pdf_oder_none(pdf_path)
            if gerettet is not None:
                return gerettet, None
            return None, 'PDF-Generierung: Timeout nach 60 Sekunden.'

        if process.returncode != 0:
            return None, f'PDF-Fehler: {stderr[:200]}'

        with open(pdf_path, 'rb') as f:
            return f.read(), None
    finally:
        for path in [html_path, pdf_path]:
            if os.path.exists(path):
                try:
                    os.unlink(path)
                except OSError:
                    pass



def _get_custom_values(field_name, predefined_options):
    predefined_keys = {v for v, _ in predefined_options if v}
    column = getattr(Contact, field_name)
    db_values = db.session.query(column).filter(
        column.isnot(None), column != ''
    ).distinct().order_by(column).all()
    return [v[0] for v in db_values if v[0] not in predefined_keys]


def _get_custom_lost_reasons():
    predefined_keys = {v for v, _ in LOST_REASONS if v}
    db_values = db.session.query(Contact.lost_reason).filter(
        Contact.lost_reason.isnot(None), Contact.lost_reason != ''
    ).distinct().order_by(Contact.lost_reason).all()
    return [v[0] for v in db_values if v[0] not in predefined_keys]



SEARCHABLE_FIELDS = [
    ('organisation', 'Organisation'),
    ('organisationsform', 'Organisationsform'),
    ('website', 'Website'),
    ('mail_organisation', 'Mail Organisation'),
    ('telefon_organisation', 'Telefon Organisation'),
    ('in_a_nutshell', 'In a Nutshell'),
    ('steuernummer', 'Steuernummer'),
    ('gemeinnuetzig', 'Gemeinnützig'),
    ('land', 'Land'),
    ('bundesland', 'Bundesland'),
    ('plz', 'PLZ'),
    ('stadt', 'Stadt'),
    ('sprache_1', 'Sprache 1'),
    ('sprache_2', 'Sprache 2'),
    ('vorname', 'Vorname'),
    ('nachname', 'Nachname'),
    ('mail_person', 'Mail Person'),
    ('mail_rechnung', 'Mail Rechnung'),
    ('telefon_person', 'Telefon Person'),
    ('position', 'Position'),
    ('anrede', 'Anrede'),
    ('titel_vor', 'Titel Vor'),
    ('titel_nach', 'Titel Nach'),
    ('linkedin', 'LinkedIn'),
    ('wie', 'Wie kontaktiert'),
    ('kommentare', 'Kommentare'),
    ('erinnerungstext', 'Erinnerungstext'),
    ('quelle', 'Quelle'),
    ('branche', 'Branche'),
]


@app.route('/contacts')
def contacts():
    search = request.args.get('search', '')
    status_filter = request.args.get('status', '')
    land_filter = request.args.get('land', '')
    quelle_filter = request.args.get('quelle', '')
    branche_filter = request.args.get('branche', '')
    show_archived = request.args.get('archived') == '1'

    query = Contact.query

    if search:
        search_term = f'%{search}%'
        search_conditions = [
            getattr(Contact, field).ilike(search_term)
            for field, _ in SEARCHABLE_FIELDS
            if field not in BOOLEAN_FIELDS
        ]
        id_query = search.lstrip('#').strip()
        if id_query.isdigit():
            search_conditions.append(Contact.id == int(id_query))
        search_lower = search.lower()
        if search_lower in ('ja', 'gemeinnützig', 'gemeinnuetzig', 'nonprofit', 'npo'):
            search_conditions.append(Contact.gemeinnuetzig == True)
        elif search_lower in ('nein', 'nicht gemeinnützig'):
            search_conditions.append(Contact.gemeinnuetzig == False)
        query = query.filter(db.or_(*search_conditions))

    if status_filter:
        query = query.filter(Contact.status == status_filter)

    if land_filter:
        query = query.filter(Contact.land == land_filter)

    if quelle_filter:
        query = query.filter(Contact.quelle == quelle_filter)

    if branche_filter:
        query = query.filter(Contact.branche == branche_filter)

    if not show_archived:
        query = query.filter(Contact.archived == False)

    archived_count = Contact.query.filter(Contact.archived == True).count()

    countries = db.session.query(Contact.land).filter(
        Contact.land.isnot(None), Contact.land != ''
    ).distinct().order_by(Contact.land).all()
    countries = [c[0] for c in countries]

    today = date.today()
    view_mode = request.args.get('view', 'standard')
    sort_by = request.args.get('sort', '')
    sort_dir = request.args.get('dir', 'desc')
    page = request.args.get('page', 1, type=int)
    per_page = 50

    if view_mode == 'finanzen' and sort_by:
        contacts_list = query.order_by(Contact.updated_at.desc()).all()
        def _sort_key(c):
            if sort_by == 'deal':
                return c.deal_value or 0
            elif sort_by == 'received':
                return c.received_payments_total
            elif sort_by == 'zeitkosten':
                return c.time_value
            elif sort_by == 'profit':
                return c.profitability
            elif sort_by == 'eff_rate':
                total_h = sum(e.dauer_minuten for e in c.time_entries if e.timer_start is None) / 60
                return (c.received_payments_total / total_h) if total_h > 0 and c.received_payments_total > 0 else 0
            return 0
        contacts_list = sorted(contacts_list, key=_sort_key, reverse=(sort_dir == 'desc'))
        total_contacts_count = len(contacts_list)
        total_pages = max(1, (total_contacts_count + per_page - 1) // per_page)
        page = max(1, min(page, total_pages))
        contacts_paginated = contacts_list[(page - 1) * per_page:page * per_page]
    else:
        total_contacts_count = query.count()
        total_pages = max(1, (total_contacts_count + per_page - 1) // per_page)
        page = max(1, min(page, total_pages))

        sort_columns = {
            'name': Contact.organisation,
            'status': Contact.status,
            'branche': Contact.branche,
            'score': Contact.lead_score,
            'updated': Contact.updated_at,
            'created': Contact.created_at,
            'deal': Contact.deal_value,
            'kontaktiert': Contact.kontaktiert,
            'erinnerung': Contact.erinnerung,
        }
        if sort_by in sort_columns:
            col = sort_columns[sort_by]
            order = col.asc() if sort_dir == 'asc' else col.desc()
            if sort_dir == 'desc':
                order = db.case((col.is_(None), 0), else_=1).desc(), col.desc()
            else:
                order = db.case((col.is_(None), 1), else_=0).asc(), col.asc()
            contacts_paginated = query.order_by(*order).offset((page - 1) * per_page).limit(per_page).all()
        else:
            contacts_paginated = query.order_by(Contact.updated_at.desc()).offset((page - 1) * per_page).limit(per_page).all()

    return render_template('contacts.html',
                           contacts=contacts_paginated, search=search,
                           page=page, total_pages=total_pages,
                           total_contacts_count=total_contacts_count,
                           status_filter=status_filter, land_filter=land_filter,
                           quelle_filter=quelle_filter, branche_filter=branche_filter,
                           show_archived=show_archived, archived_count=archived_count,
                           stages=PIPELINE_STAGES, countries=countries,
                           quellen_options=QUELLEN_OPTIONS,
                           branchen_options=BRANCHEN_OPTIONS,
                           lost_reasons=LOST_REASONS,
                           kontakt_wie_options=KONTAKT_WIE_OPTIONS,
                           view_mode=view_mode,
                           sort_by=sort_by, sort_dir=sort_dir,
                           format_minutes=_format_minutes,
                           today=today,
                           tomorrow=today + timedelta(days=1),
                           next_week=today + timedelta(days=7))


ENUM_FIELD_OPTIONS = {
    'status': PIPELINE_STAGES,
    'prioritaet': PRIORITAET_OPTIONS,
    'quelle': QUELLEN_OPTIONS,
    'branche': BRANCHEN_OPTIONS,
    'wie': KONTAKT_WIE_OPTIONS,
    'anrede': ANREDE_OPTIONS,
    'anrede_en': ANREDE_EN_OPTIONS,
    'sprache_1': SPRACHE_OPTIONS,
    'sprache_2': SPRACHE_OPTIONS,
}
BOOLEAN_FIELDS = {'gemeinnuetzig'}


@app.route('/api/contacts/field-values')
def api_field_values():
    field = request.args.get('field', '')
    available_fields = dict(SEARCHABLE_FIELDS + [('status', 'Status'), ('prioritaet', 'Priorität')])

    if field not in available_fields:
        return jsonify({'error': 'Invalid field'}), 400

    if field in BOOLEAN_FIELDS:
        return jsonify({'type': 'boolean', 'values': [
            {'value': '1', 'label': 'Ja'},
            {'value': '0', 'label': 'Nein'},
        ]})

    if field in ENUM_FIELD_OPTIONS:
        options = [{'value': v, 'label': l} for v, l in ENUM_FIELD_OPTIONS[field] if v]
        return jsonify({'type': 'enum', 'values': options})

    column = getattr(Contact, field)
    rows = db.session.query(column).filter(
        column.isnot(None), column != ''
    ).distinct().order_by(column).all()
    values = [{'value': str(v[0]), 'label': str(v[0])} for v in rows]
    return jsonify({'type': 'text', 'values': values})


@app.route('/contacts/filter')
def contacts_filter():
    fields = request.args.getlist('field')
    values = request.args.getlist('value')
    show_archived = request.args.get('archived') == '1'

    active_filters = []
    available_fields = SEARCHABLE_FIELDS + [('status', 'Status'), ('prioritaet', 'Priorität')]
    valid_field_keys = dict(available_fields)

    for i in range(len(fields)):
        f = fields[i] if i < len(fields) else ''
        v = values[i] if i < len(values) else ''
        if f and v and f in valid_field_keys:
            active_filters.append((f, v))

    contacts_list = []
    if active_filters:
        query = Contact.query
        for field, value in active_filters:
            column = getattr(Contact, field)
            if field in BOOLEAN_FIELDS:
                query = query.filter(column == (value == '1'))
            else:
                query = query.filter(column == value)
        if not show_archived:
            query = query.filter(Contact.archived == False)
        contacts_list = query.order_by(Contact.updated_at.desc()).all()

    return render_template('contacts_filter.html',
                           contacts=contacts_list,
                           available_fields=available_fields,
                           active_filters=active_filters,
                           field_labels=valid_field_keys,
                           stages=PIPELINE_STAGES,
                           today=date.today())


@app.route('/contacts/new', methods=['GET', 'POST'])
def contact_new():
    if request.method == 'POST':
        contact = create_contact_from_form(request.form)
        db.session.add(contact)
        db.session.flush()

        sc = StatusChange(contact_id=contact.id, from_status=None, to_status=contact.status,
                          deal_value_at_change=contact.deal_value)
        db.session.add(sc)
        db.session.commit()

        flash(f'Kontakt "{contact.display_name}" erstellt', 'success')
        return redirect(url_for('contacts'))

    return render_template('contact_form.html', contact=None,
                           stages=PIPELINE_STAGES, anrede_options=ANREDE_OPTIONS,
                           anrede_en_options=ANREDE_EN_OPTIONS,
                           kontakt_wie_options=KONTAKT_WIE_OPTIONS,
                           sprache_options=SPRACHE_OPTIONS,
                           quellen_options=QUELLEN_OPTIONS,
                           branchen_options=BRANCHEN_OPTIONS,
                           prioritaet_options=PRIORITAET_OPTIONS,
                           tier_options=TIER_OPTIONS,
                           pipeline_phase_options=PIPELINE_PHASE_OPTIONS,
                           deal_probability_defaults=DEAL_PROBABILITY_DEFAULTS,
                           custom_quellen=_get_custom_values('quelle', QUELLEN_OPTIONS),
                           custom_branchen=_get_custom_values('branche', BRANCHEN_OPTIONS))


@app.route('/contacts/<int:id>')
def contact_view(id):
    contact = Contact.query.options(
        joinedload(Contact.interactions),
        joinedload(Contact.status_changes),
    ).get_or_404(id)
    tasks = Task.query.filter_by(contact_id=id).filter(Task.status != 'erledigt').order_by(Task.due_date).all()
    completed_tasks_count = Task.query.filter_by(contact_id=id).filter(Task.status == 'erledigt').count()
    templates = EmailTemplate.query.order_by(EmailTemplate.name).all()

    time_entries = TimeEntry.query.filter_by(contact_id=id).filter(
        TimeEntry.timer_start.is_(None)
    ).order_by(TimeEntry.datum.desc(), TimeEntry.created_at.desc()).all()
    total_minutes = sum(e.dauer_minuten for e in time_entries)
    active_timer = _get_active_timer()
    contact_timer = active_timer if active_timer and active_timer.contact_id == id else None

    documents = Document.query.filter_by(contact_id=id).order_by(Document.date.desc()).all()

    time_value = contact.time_value

    retainer_usage = []
    active_retainers = Retainer.query.filter_by(contact_id=id, active=True).all()
    today_date = date.today()
    month_start = today_date.replace(day=1)
    for ret in active_retainers:
        used_minutes = db.session.query(
            func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
        ).filter(
            TimeEntry.contact_id == id,
            TimeEntry.datum >= month_start,
            TimeEntry.timer_start.is_(None)
        ).scalar() or 0
        used_hours = round(used_minutes / 60, 1)
        pct = round(used_hours / ret.monthly_hours * 100, 1) if ret.monthly_hours > 0 else 0
        retainer_usage.append({
            'retainer': ret,
            'used_hours': used_hours,
            'included_hours': ret.monthly_hours,
            'pct': min(pct, 100),
            'overage': max(0, round(used_hours - ret.monthly_hours, 1)),
            'over_limit': used_hours > ret.monthly_hours,
        })

    unbilled_entries = [e for e in time_entries if e.billed_document_id is None]
    unbilled_minutes = sum(e.dauer_minuten for e in unbilled_entries)
    billing = compute_billing_status(
        contact, active_retainers[0] if active_retainers else None,
        unbilled_entries, today_date, default_rate=_get_default_hourly_rate())

    timeline = []
    for inter in contact.interactions:
        timeline.append({'date': inter.datum, 'type': 'interaction', 'icon': 'chat',
                         'title': inter.typ, 'detail': inter.notiz, 'color': 'cyan'})
    for sc in contact.status_changes:
        stage_labels = dict(PIPELINE_STAGES)
        from_l = stage_labels.get(sc.from_status, sc.from_status) if sc.from_status else '—'
        to_l = stage_labels.get(sc.to_status, sc.to_status)
        timeline.append({'date': sc.changed_at.date(), 'type': 'status', 'icon': 'arrow',
                         'title': f'{from_l} → {to_l}',
                         'detail': f'{sc.deal_value_at_change:,.0f} €'.replace(',', '.') if sc.deal_value_at_change else None,
                         'color': 'emerald' if sc.to_status == 'won' else ('rose' if sc.to_status == 'lost' else 'petrol')})
    for doc in documents:
        timeline.append({'date': doc.date, 'type': 'document', 'icon': 'doc',
                         'title': f'{doc.type_label} {doc.number}',
                         'detail': f'{doc.subject[:50]} — {doc.total_net:,.2f} €'.replace(',', 'X').replace('.', ',').replace('X', '.'),
                         'color': 'violet'})
    for te in time_entries[:20]:
        taetigkeit_labels = dict(TAETIGKEIT_KONTAKT_OPTIONS)
        timeline.append({'date': te.datum, 'type': 'time', 'icon': 'clock',
                         'title': taetigkeit_labels.get(te.taetigkeit, te.taetigkeit),
                         'detail': f'{te.dauer_display}' + (f' — {te.notiz}' if te.notiz else ''),
                         'color': 'sky'})
    timeline.sort(key=lambda x: x['date'], reverse=True)
    timeline = timeline[:30]

    accounts = CustomerAccount.query.filter_by(contact_id=id)\
        .join(AccountToolType).order_by(AccountToolType.sort_order).all()
    tool_types = AccountToolType.query.order_by(AccountToolType.sort_order).all()

    return render_template('contact_view.html', contact=contact, tasks=tasks,
                           completed_tasks_count=completed_tasks_count,
                           documents=documents, timeline=timeline,
                           stages=PIPELINE_STAGES, today=date.today(),
                           tomorrow=date.today() + timedelta(days=1),
                           next_week=date.today() + timedelta(days=7),
                           kontakt_wie_options=KONTAKT_WIE_OPTIONS,
                           lost_reasons=LOST_REASONS,
                           custom_lost_reasons=_get_custom_lost_reasons(),
                           templates=templates,
                           time_entries=time_entries,
                           total_minutes=total_minutes,
                           time_value=time_value,
                           format_minutes=_format_minutes,
                           contact_timer=contact_timer,
                           taetigkeit_kontakt=TAETIGKEIT_KONTAKT_OPTIONS,
                           custom_taetigkeiten_kontakt=_get_custom_taetigkeiten(True),
                           payment_types=PAYMENT_TYPES,
                           retainer_usage=retainer_usage,
                           billing=billing,
                           unbilled_minutes=unbilled_minutes,
                           tier_options=TIER_OPTIONS,
                           pipeline_phase_options=PIPELINE_PHASE_OPTIONS,
                           accounts=accounts,
                           tool_types=tool_types,
                           account_statuses=ACCOUNT_STATUSES,
                           account_categories=dict(ACCOUNT_CATEGORIES))


@app.route('/contacts/<int:id>/accounts/add', methods=['POST'])
def account_add(id):
    contact = Contact.query.get_or_404(id)
    account = CustomerAccount(
        contact_id=id,
        tool_type_id=int(request.form['tool_type_id']),
        account_id=request.form.get('account_id', '').strip() or None,
        login_email=request.form.get('login_email', '').strip() or None,
        login_password=request.form.get('login_password', '').strip() or None,
        client_login_email=request.form.get('client_login_email', '').strip() or None,
        client_login_password=request.form.get('client_login_password', '').strip() or None,
        url=request.form.get('url', '').strip() or None,
        status=request.form.get('status', 'pending'),
        notes=request.form.get('notes', '').strip() or None,
    )
    db.session.add(account)
    db.session.commit()
    flash('Konto hinzugefügt', 'success')
    return redirect(url_for('contact_view', id=id) + '#konten')


@app.route('/contacts/<int:contact_id>/accounts/<int:account_id>/edit', methods=['POST'])
def account_edit(contact_id, account_id):
    account = CustomerAccount.query.get_or_404(account_id)
    account.tool_type_id = int(request.form['tool_type_id'])
    account.account_id = request.form.get('account_id', '').strip() or None
    account.login_email = request.form.get('login_email', '').strip() or None
    account.login_password = request.form.get('login_password', '').strip() or None
    account.client_login_email = request.form.get('client_login_email', '').strip() or None
    account.client_login_password = request.form.get('client_login_password', '').strip() or None
    account.url = request.form.get('url', '').strip() or None
    account.status = request.form.get('status', 'pending')
    account.notes = request.form.get('notes', '').strip() or None
    db.session.commit()
    flash('Konto aktualisiert', 'success')
    return redirect(url_for('contact_view', id=contact_id) + '#konten')


@app.route('/contacts/<int:contact_id>/accounts/<int:account_id>/delete', methods=['POST'])
def account_delete(contact_id, account_id):
    account = CustomerAccount.query.get_or_404(account_id)
    db.session.delete(account)
    db.session.commit()
    flash('Konto gelöscht', 'success')
    return redirect(url_for('contact_view', id=contact_id) + '#konten')


@app.route('/contacts/<int:contact_id>/accounts/<int:account_id>/toggle-done', methods=['POST'])
def account_toggle_done(contact_id, account_id):
    account = CustomerAccount.query.get_or_404(account_id)
    account.done = not account.done
    db.session.commit()
    return jsonify({'done': account.done})


@app.route('/api/tool-types/add', methods=['POST'])
def tool_type_add():
    data = request.get_json()
    max_order = db.session.query(func.max(AccountToolType.sort_order)).scalar() or 0
    tool = AccountToolType(
        name=data['name'].strip(),
        short_name=data['short_name'].strip(),
        category=data.get('category', 'other'),
        default_url=data.get('default_url', '').strip() or None,
        sort_order=max_order + 10,
    )
    db.session.add(tool)
    db.session.commit()
    return jsonify({'id': tool.id, 'name': tool.name, 'short_name': tool.short_name,
                    'category': tool.category, 'default_url': tool.default_url or ''})


@app.route('/contacts/<int:id>/infrastruktur-pdf')
def infrastruktur_pdf(id):
    contact = Contact.query.get_or_404(id)
    accounts = CustomerAccount.query.filter_by(contact_id=id)\
        .join(AccountToolType).order_by(AccountToolType.sort_order).all()

    mode = request.args.get('mode', 'client')
    category_labels = dict(ACCOUNT_CATEGORIES)
    html_content = render_template('infrastruktur-pdf.html',
                                   contact=contact, accounts=accounts,
                                   category_labels=category_labels,
                                   mode=mode,
                                   now=datetime.now())

    tmp_html = os.path.join(tempfile.gettempdir(), f'infrastruktur-{contact.id}-{uuid.uuid4().hex[:8]}.html')
    tmp_pdf = tmp_html.replace('.html', '.pdf')
    try:
        with open(tmp_html, 'w', encoding='utf-8') as f:
            f.write(html_content)

        node_script = f"""
        const puppeteer = require('{os.path.join(BASE_DIR, 'node_modules', 'puppeteer').replace(os.sep, '/')}');
        (async () => {{
            const browser = await puppeteer.launch({{headless: 'new', args: ['--no-sandbox']}});
            const page = await browser.newPage();
            await page.goto('file:///{tmp_html.replace(os.sep, '/')}', {{waitUntil: 'networkidle0', timeout: 30000}});
            await page.pdf({{path: '{tmp_pdf.replace(os.sep, '/')}', format: 'A4', printBackground: true,
                            margin: {{top: '0', bottom: '0', left: '0', right: '0'}}}});
            await Promise.race([
                browser.close(),
                new Promise(resolve => setTimeout(resolve, 3000))
            ]);
            process.exit(0);
        }})();
        """
        safe_name = contact.display_name.replace(' ', '_').replace('/', '_')
        try:
            result = subprocess.run(['node', '-e', node_script], capture_output=True,
                                    text=True, errors='replace', timeout=60)
        except subprocess.TimeoutExpired:
            gerettet = _fertiges_pdf_oder_none(tmp_pdf)
            if gerettet is None:
                flash('PDF-Generierung: Timeout nach 60 Sekunden.', 'error')
                return redirect(url_for('contact_view', id=id))
            return send_file(io.BytesIO(gerettet), as_attachment=True,
                             download_name=f'Infrastruktur_{safe_name}.pdf',
                             mimetype='application/pdf')

        if result.returncode != 0:
            flash(f'PDF-Fehler: {result.stderr[:200]}', 'error')
            return redirect(url_for('contact_view', id=id))

        with open(tmp_pdf, 'rb') as f:
            fertiges_pdf = f.read()
        return send_file(io.BytesIO(fertiges_pdf), as_attachment=True,
                         download_name=f'Infrastruktur_{safe_name}.pdf',
                         mimetype='application/pdf')
    finally:
        for f in [tmp_html, tmp_pdf]:
            try:
                os.unlink(f)
            except OSError:
                pass


@app.route('/contacts/<int:id>/edit', methods=['GET', 'POST'])
def contact_edit(id):
    contact = Contact.query.get_or_404(id)

    if request.method == 'POST':
        update_contact_from_form(contact, request.form)
        db.session.commit()
        flash(f'Kontakt "{contact.display_name}" aktualisiert', 'success')
        return redirect(url_for('contact_view', id=id))

    return render_template('contact_form.html', contact=contact,
                           stages=PIPELINE_STAGES, anrede_options=ANREDE_OPTIONS,
                           anrede_en_options=ANREDE_EN_OPTIONS,
                           kontakt_wie_options=KONTAKT_WIE_OPTIONS,
                           sprache_options=SPRACHE_OPTIONS,
                           quellen_options=QUELLEN_OPTIONS,
                           branchen_options=BRANCHEN_OPTIONS,
                           prioritaet_options=PRIORITAET_OPTIONS,
                           tier_options=TIER_OPTIONS,
                           pipeline_phase_options=PIPELINE_PHASE_OPTIONS,
                           deal_probability_defaults=DEAL_PROBABILITY_DEFAULTS,
                           custom_quellen=_get_custom_values('quelle', QUELLEN_OPTIONS),
                           custom_branchen=_get_custom_values('branche', BRANCHEN_OPTIONS))


@app.route('/contacts/<int:id>/delete', methods=['POST'])
def contact_delete(id):
    contact = Contact.query.get_or_404(id)
    name = contact.display_name

    blockers = []
    if contact.incoming_invoices:
        blockers.append(f'{len(contact.incoming_invoices)} Eingangsrechnung(en)')
    if contact.documents:
        blockers.append(f'{len(contact.documents)} Rechnung(en)/Angebot(e)')
    if contact.retainers:
        blockers.append(f'{len(contact.retainers)} Retainer')
    if blockers:
        flash(f'Kontakt "{name}" kann nicht gelöscht werden — daran hängen noch: '
              f'{", ".join(blockers)}. Bitte zuerst diese Belege entfernen oder '
              f'per Merge auf einen anderen Kontakt umhängen.', 'error')
        return redirect(url_for('contact_view', id=id))

    db.session.delete(contact)
    db.session.commit()
    flash(f'Kontakt "{name}" gelöscht', 'success')
    return redirect(url_for('contacts'))


@app.route('/contacts/merge', methods=['POST'])
def contact_merge():
    data = request.get_json()
    source_id = data.get('source_id')
    target_id = data.get('target_id')
    if not source_id or not target_id or source_id == target_id:
        return jsonify({'success': False, 'error': 'Ungültige IDs'}), 400

    source = Contact.query.get_or_404(source_id)
    target = Contact.query.get_or_404(target_id)

    for inter in Interaction.query.filter_by(contact_id=source_id).all():
        inter.contact_id = target_id
    for sc in StatusChange.query.filter_by(contact_id=source_id).all():
        sc.contact_id = target_id
    for te in TimeEntry.query.filter_by(contact_id=source_id).all():
        te.contact_id = target_id
    for doc in Document.query.filter_by(contact_id=source_id).all():
        doc.contact_id = target_id
    for ep in ExpectedPayment.query.filter_by(contact_id=source_id).all():
        ep.contact_id = target_id
    for task in Task.query.filter_by(contact_id=source_id).all():
        task.contact_id = target_id
    for ret in Retainer.query.filter_by(contact_id=source_id).all():
        ret.contact_id = target_id

    for field in ['website', 'mail_organisation', 'telefon_organisation', 'in_a_nutshell',
                  'land', 'bundesland', 'plz', 'stadt', 'mail_person', 'mail_rechnung', 'telefon_person',
                  'position', 'linkedin', 'quelle', 'branche']:
        if not getattr(target, field) and getattr(source, field):
            setattr(target, field, getattr(source, field))

    if not target.deal_value and source.deal_value:
        target.deal_value = source.deal_value

    source_name = source.display_name
    db.session.delete(source)
    db.session.commit()

    return jsonify({'success': True, 'message': f'"{source_name}" in "{target.display_name}" zusammengeführt'})


@app.route('/contacts/<int:id>/status', methods=['POST'])
def contact_update_status(id):
    contact = Contact.query.get_or_404(id)
    data = request.json
    new_status = data.get('status')
    if new_status not in dict(PIPELINE_STAGES):
        return jsonify({'success': False, 'error': 'Invalid status'}), 400

    old_status = contact.status

    if new_status == 'lost':
        lost_reason = data.get('lost_reason', '')
        contact.lost_reason = lost_reason if lost_reason else None

    old_default = DEAL_PROBABILITY_DEFAULTS.get(old_status, 0)
    if contact.deal_probability is None or contact.deal_probability == old_default:
        contact.deal_probability = DEAL_PROBABILITY_DEFAULTS.get(new_status, 0)

    contact.status = new_status

    sc = StatusChange(contact_id=id, from_status=old_status, to_status=new_status,
                      deal_value_at_change=contact.deal_value)
    db.session.add(sc)

    if new_status == 'on_hold':
        open_system_tasks = Task.query.filter(
            Task.contact_id == id,
            Task.status != 'erledigt',
            Task.created_by == 'System',
            db.or_(Task.title.like('Follow-up:%'), Task.title.like('Nachhaken:%'))
        ).all()
        for t in open_system_tasks:
            t.status = 'erledigt'
            t.completed_at = datetime.now(timezone.utc).replace(tzinfo=None)

    db.session.commit()
    return jsonify({'success': True})


@app.route('/contacts/bulk-status', methods=['POST'])
def contacts_bulk_status():
    data = request.get_json()
    contact_ids = data.get('ids', [])
    new_status = data.get('status', '')
    lost_reason = data.get('lost_reason', '')

    if not contact_ids or new_status not in dict(PIPELINE_STAGES):
        return jsonify({'success': False, 'error': 'Ungültige Anfrage'}), 400

    count = 0
    for cid in contact_ids:
        contact = Contact.query.get(cid)
        if not contact:
            continue
        old_status = contact.status
        if old_status == new_status:
            continue

        old_default = DEAL_PROBABILITY_DEFAULTS.get(old_status, 0)
        if contact.deal_probability is None or contact.deal_probability == old_default:
            contact.deal_probability = DEAL_PROBABILITY_DEFAULTS.get(new_status, 0)

        if new_status == 'lost' and lost_reason:
            contact.lost_reason = lost_reason

        contact.status = new_status
        sc = StatusChange(contact_id=cid, from_status=old_status, to_status=new_status,
                          deal_value_at_change=contact.deal_value)
        db.session.add(sc)

        if new_status == 'on_hold':
            open_system_tasks = Task.query.filter(
                Task.contact_id == cid,
                Task.status != 'erledigt',
                Task.created_by == 'System',
                db.or_(Task.title.like('Follow-up:%'), Task.title.like('Nachhaken:%'))
            ).all()
            for t in open_system_tasks:
                t.status = 'erledigt'
                t.completed_at = datetime.now(timezone.utc).replace(tzinfo=None)

        count += 1

    db.session.commit()
    return jsonify({'success': True, 'count': count})


def _du_anrede(contact):
    if contact.vorname:
        return f'Hallo {contact.vorname},'
    if contact.organisation:
        return f'Liebes {contact.organisation} Team,'
    return 'Hallo zusammen,'


@app.route('/contacts/<int:id>/email')
def contact_email(id):
    contact = Contact.query.get_or_404(id)
    template_id = request.args.get('template_id', type=int)
    if not template_id:
        return jsonify({'betreff': '', 'text': ''})

    template = EmailTemplate.query.get_or_404(template_id)
    replacements = {
        '{anrede_du}': _du_anrede(contact),
        '{vorname}': contact.vorname or '',
        '{nachname}': contact.nachname or '',
        '{organisation}': contact.organisation or '',
        '{position}': contact.position or '',
        '{anrede}': contact.anrede or '',
        '{website}': contact.website or '',
        '{stadt}': contact.stadt or '',
        '{land}': contact.land or '',
    }

    latest_quote = Document.query.filter_by(
        contact_id=id, doc_type='quote'
    ).order_by(Document.date.desc()).first()
    latest_invoice = Document.query.filter_by(
        contact_id=id, doc_type='invoice'
    ).order_by(Document.date.desc()).first()
    replacements['{angebotsnummer}'] = latest_quote.number if latest_quote else ''
    replacements['{angebotsdatum}'] = latest_quote.date.strftime('%d.%m.%Y') if latest_quote else ''
    replacements['{angebotssumme}'] = f'{latest_quote.total_net:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.') if latest_quote else ''
    replacements['{rechnungsnummer}'] = latest_invoice.number if latest_invoice else ''
    replacements['{firma}'] = CompanySettings.get('company_name')
    owner_full = CompanySettings.get('owner_name') or ''
    replacements['{mein_name}'] = owner_full.split()[0] if owner_full else ''
    replacements['{signatur}'] = CompanySettings.get('email_signature')

    betreff = template.betreff
    text = template.text
    for placeholder, value in replacements.items():
        betreff = betreff.replace(placeholder, value)
        text = text.replace(placeholder, value)

    return jsonify({'betreff': betreff, 'text': text})


def parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except:
        return None


def parse_form_float(value):
    if not value or not value.strip():
        return None
    try:
        return float(value.strip().replace(',', '.'))
    except (ValueError, TypeError):
        return None


def parse_form_int(value):
    if not value or not value.strip():
        return None
    try:
        return int(value.strip())
    except (ValueError, TypeError):
        return None


def create_contact_from_form(form):
    return Contact(
        organisation=form.get('organisation', '').strip() or None,
        organisationsform=form.get('organisationsform', '').strip() or None,
        website=form.get('website', '').strip() or None,
        mail_organisation=form.get('mail_organisation', '').strip() or None,
        telefon_organisation=form.get('telefon_organisation', '').strip() or None,
        in_a_nutshell=form.get('in_a_nutshell', '').strip() or None,
        gruendungsdatum=parse_date(form.get('gruendungsdatum')),
        steuernummer=form.get('steuernummer', '').strip() or None,
        uid=form.get('uid', '').strip() or None,
        gemeinnuetzig=form.get('gemeinnuetzig') == 'on',
        has_social_mandate=form.get('has_social_mandate') == 'on',
        mandat_aktiv=form.get('mandat_aktiv') == 'on',
        land=form.get('land', '').strip() or None,
        bundesland=form.get('bundesland', '').strip() or None,
        plz=form.get('plz', '').strip() or None,
        stadt=form.get('stadt', '').strip() or None,
        strasse=form.get('strasse', '').strip() or None,
        sprache_1=form.get('sprache_1', '').strip() or None,
        sprache_2=form.get('sprache_2', '').strip() or None,
        vorname=form.get('vorname', '').strip() or None,
        nachname=form.get('nachname', '').strip() or None,
        mail_person=form.get('mail_person', '').strip() or None,
        mail_rechnung=form.get('mail_rechnung', '').strip() or None,
        rechnung_name=form.get('rechnung_name', '').strip() or None,
        rechnung_uid=form.get('rechnung_uid', '').strip() or None,
        rechnung_strasse=form.get('rechnung_strasse', '').strip() or None,
        rechnung_plz=form.get('rechnung_plz', '').strip() or None,
        rechnung_stadt=form.get('rechnung_stadt', '').strip() or None,
        rechnung_land=form.get('rechnung_land', '').strip() or None,
        telefon_person=form.get('telefon_person', '').strip() or None,
        position=form.get('position', '').strip() or None,
        anrede=form.get('anrede', '').strip() or None,
        anrede_en=form.get('anrede_en', '').strip() or None,
        titel_vor=form.get('titel_vor', '').strip() or None,
        titel_nach=form.get('titel_nach', '').strip() or None,
        linkedin=form.get('linkedin', '').strip() or None,
        kontaktiert=parse_date(form.get('kontaktiert')),
        wie=form.get('wie', '').strip() or None,
        kommentare=form.get('kommentare', '').strip() or None,
        erinnerung=parse_date(form.get('erinnerung')),
        erinnerungstext=form.get('erinnerungstext', '').strip() or None,
        status=form.get('status', 'new'),
        deal_value=parse_form_float(form.get('deal_value', '')),
        deal_probability=parse_form_int(form.get('deal_probability', '')),
        quelle=form.get('quelle', '').strip() or None,
        branche=form.get('branche', '').strip() or None,
        prioritaet=form.get('prioritaet', 'B').strip() or 'B',
        hourly_rate=parse_form_float(form.get('hourly_rate', '')),
        abrechnung_terminiert_bis=parse_date(form.get('abrechnung_terminiert_bis')),
        tier=form.get('tier', '').strip() or None,
        pipeline_phase=form.get('pipeline_phase', '').strip() or None,
    )


def update_contact_from_form(contact, form):
    contact.organisation = form.get('organisation', '').strip() or None
    contact.organisationsform = form.get('organisationsform', '').strip() or None
    contact.website = form.get('website', '').strip() or None
    contact.mail_organisation = form.get('mail_organisation', '').strip() or None
    contact.telefon_organisation = form.get('telefon_organisation', '').strip() or None
    contact.in_a_nutshell = form.get('in_a_nutshell', '').strip() or None
    contact.gruendungsdatum = parse_date(form.get('gruendungsdatum'))
    contact.steuernummer = form.get('steuernummer', '').strip() or None
    contact.uid = form.get('uid', '').strip() or None
    contact.gemeinnuetzig = form.get('gemeinnuetzig') == 'on'
    contact.has_social_mandate = form.get('has_social_mandate') == 'on'
    contact.mandat_aktiv = form.get('mandat_aktiv') == 'on'
    contact.land = form.get('land', '').strip() or None
    contact.bundesland = form.get('bundesland', '').strip() or None
    contact.plz = form.get('plz', '').strip() or None
    contact.stadt = form.get('stadt', '').strip() or None
    contact.strasse = form.get('strasse', '').strip() or None
    contact.sprache_1 = form.get('sprache_1', '').strip() or None
    contact.sprache_2 = form.get('sprache_2', '').strip() or None
    contact.vorname = form.get('vorname', '').strip() or None
    contact.nachname = form.get('nachname', '').strip() or None
    contact.mail_person = form.get('mail_person', '').strip() or None
    contact.mail_rechnung = form.get('mail_rechnung', '').strip() or None
    contact.rechnung_name = form.get('rechnung_name', '').strip() or None
    contact.rechnung_uid = form.get('rechnung_uid', '').strip() or None
    contact.rechnung_strasse = form.get('rechnung_strasse', '').strip() or None
    contact.rechnung_plz = form.get('rechnung_plz', '').strip() or None
    contact.rechnung_stadt = form.get('rechnung_stadt', '').strip() or None
    contact.rechnung_land = form.get('rechnung_land', '').strip() or None
    contact.telefon_person = form.get('telefon_person', '').strip() or None
    contact.position = form.get('position', '').strip() or None
    contact.anrede = form.get('anrede', '').strip() or None
    contact.anrede_en = form.get('anrede_en', '').strip() or None
    contact.titel_vor = form.get('titel_vor', '').strip() or None
    contact.titel_nach = form.get('titel_nach', '').strip() or None
    contact.linkedin = form.get('linkedin', '').strip() or None
    contact.kontaktiert = parse_date(form.get('kontaktiert'))
    contact.wie = form.get('wie', '').strip() or None
    contact.kommentare = form.get('kommentare', '').strip() or None
    contact.erinnerung = parse_date(form.get('erinnerung'))
    contact.erinnerungstext = form.get('erinnerungstext', '').strip() or None
    contact.status = form.get('status', contact.status)
    contact.deal_value = parse_form_float(form.get('deal_value', ''))
    contact.deal_probability = parse_form_int(form.get('deal_probability', ''))
    contact.quelle = form.get('quelle', '').strip() or None
    contact.branche = form.get('branche', '').strip() or None
    contact.prioritaet = form.get('prioritaet', 'B').strip() or 'B'
    contact.hourly_rate = parse_form_float(form.get('hourly_rate', ''))
    contact.abrechnung_terminiert_bis = parse_date(form.get('abrechnung_terminiert_bis'))
    contact.tier = form.get('tier', '').strip() or None
    contact.pipeline_phase = form.get('pipeline_phase', '').strip() or None



@app.route('/contacts/<int:id>/gemeinnuetzig', methods=['POST'])
def contact_toggle_gemeinnuetzig(id):
    contact = Contact.query.get_or_404(id)
    contact.gemeinnuetzig = not contact.gemeinnuetzig
    db.session.commit()
    return jsonify({'success': True, 'gemeinnuetzig': contact.gemeinnuetzig})


@app.route('/contacts/<int:id>/archive', methods=['POST'])
def contact_toggle_archive(id):
    contact = Contact.query.get_or_404(id)
    contact.archived = not contact.archived
    db.session.commit()
    return jsonify({'success': True, 'archived': contact.archived})



def _get_active_timer():
    return TimeEntry.query.filter(TimeEntry.timer_start.isnot(None)).first()


def _get_custom_taetigkeiten(is_kontakt=True):
    options = TAETIGKEIT_KONTAKT_OPTIONS if is_kontakt else TAETIGKEIT_ALLGEMEIN_OPTIONS
    predefined = {v for v, _ in options}
    query = TimeEntry.query
    if is_kontakt:
        query = query.filter(TimeEntry.contact_id.isnot(None))
    else:
        query = query.filter(TimeEntry.contact_id.is_(None))
    db_values = query.with_entities(TimeEntry.taetigkeit).filter(
        TimeEntry.taetigkeit.isnot(None), TimeEntry.taetigkeit != ''
    ).distinct().order_by(TimeEntry.taetigkeit).all()
    return [v[0] for v in db_values if v[0] not in predefined]


def _format_minutes(minutes):
    h = minutes // 60
    m = minutes % 60
    return f'{h}:{m:02d}'



def _eigenarbeit_clause():
    return db.or_(
        TimeEntry.contact_id.is_(None),
        TimeEntry.taetigkeit.in_([k for k, _ in TAETIGKEIT_ALLGEMEIN_OPTIONS]),
        TimeEntry.contact_id.in_(
            db.session.query(Contact.id).filter(Contact.abrechnungsmodell == 'intern')
        ),
    )


def _strang_split(von, bis=None):
    def _period(query, col):
        query = query.filter(col >= von)
        if bis is not None:
            query = query.filter(col < bis)
        return query

    ist = _period(
        db.session.query(func.coalesce(func.sum(WorkDay.minuten), 0)), WorkDay.datum
    ).scalar()

    verbucht_q = db.session.query(
        func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
    ).filter(TimeEntry.timer_start.is_(None))
    verbucht = _period(verbucht_q, TimeEntry.datum).scalar()

    eigen_q = db.session.query(
        func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
    ).filter(TimeEntry.timer_start.is_(None), _eigenarbeit_clause())
    eigen = _period(eigen_q, TimeEntry.datum).scalar()

    return ist, verbucht, eigen


def _strang_kennzahlen(von, bis=None):
    ist, verbucht, eigen = _strang_split(von, bis)
    return {
        'ist': ist,
        'verbucht': verbucht,
        'eigenarbeit': eigen,
        'kunde': verbucht - eigen,
        'quote': round(verbucht / ist * 100) if ist else None,
        'kundenquote': round((verbucht - eigen) / ist * 100) if ist else None,
    }


@app.route('/timetracking')
def timetracking():
    today = date.today()
    week_start = today - timedelta(days=today.weekday())
    month_start = today.replace(day=1)

    active_timer = _get_active_timer()

    today_minutes = db.session.query(func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)).filter(
        TimeEntry.datum == today, TimeEntry.timer_start.is_(None)
    ).scalar()
    week_minutes = db.session.query(func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)).filter(
        TimeEntry.datum >= week_start, TimeEntry.timer_start.is_(None)
    ).scalar()
    month_minutes = db.session.query(func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)).filter(
        TimeEntry.datum >= month_start, TimeEntry.timer_start.is_(None)
    ).scalar()

    workday_today = WorkDay.query.get(today)
    strang_heute = _strang_kennzahlen(today, today + timedelta(days=1))
    strang_woche = _strang_kennzahlen(week_start, week_start + timedelta(days=7))
    strang_monat = _strang_kennzahlen(month_start)

    erfasste_tage = {
        w.datum for w in WorkDay.query.filter(WorkDay.datum >= month_start).all()
    }
    workday_luecken = [
        month_start + timedelta(days=i)
        for i in range((today - month_start).days + 1)
        if (month_start + timedelta(days=i)) not in erfasste_tage
        and (month_start + timedelta(days=i)).weekday() < 5
    ]

    filter_zeitraum = request.args.get('zeitraum', 'woche')
    filter_contact = request.args.get('contact_id', '')
    filter_taetigkeit = request.args.get('taetigkeit', '')

    query = TimeEntry.query.filter(TimeEntry.timer_start.is_(None))
    if filter_zeitraum == 'heute':
        query = query.filter(TimeEntry.datum == today)
    elif filter_zeitraum == 'woche':
        query = query.filter(TimeEntry.datum >= week_start)
    elif filter_zeitraum == 'monat':
        query = query.filter(TimeEntry.datum >= month_start)

    if filter_contact:
        query = query.filter(TimeEntry.contact_id == int(filter_contact))
    if filter_taetigkeit:
        query = query.filter(TimeEntry.taetigkeit == filter_taetigkeit)

    entries = query.order_by(TimeEntry.datum.desc(), TimeEntry.created_at.desc()).all()
    contacts_list = Contact.query.order_by(Contact.organisation, Contact.nachname).all()

    taetigkeit_breakdown = db.session.query(
        TimeEntry.taetigkeit, func.sum(TimeEntry.dauer_minuten)
    ).filter(
        TimeEntry.datum >= month_start, TimeEntry.timer_start.is_(None)
    ).group_by(TimeEntry.taetigkeit).order_by(func.sum(TimeEntry.dauer_minuten).desc()).all()

    taetigkeit_labels = dict(TAETIGKEIT_KONTAKT_OPTIONS + TAETIGKEIT_ALLGEMEIN_OPTIONS)
    taetigkeit_stats = []
    total_breakdown = sum(m for _, m in taetigkeit_breakdown) or 1
    for tae, minutes in taetigkeit_breakdown:
        taetigkeit_stats.append({
            'key': tae,
            'label': taetigkeit_labels.get(tae, tae),
            'minuten': minutes,
            'pct': round(minutes / total_breakdown * 100, 1),
        })

    week_grid_days = []
    for d in range(5):
        day = week_start + timedelta(days=d)
        week_grid_days.append({'date': day, 'label': day.strftime('%a %d.%m.')})

    week_entries = TimeEntry.query.filter(
        TimeEntry.datum >= week_start,
        TimeEntry.datum < week_start + timedelta(days=7),
        TimeEntry.timer_start.is_(None)
    ).all()

    taetigkeit_labels_all = dict(TAETIGKEIT_KONTAKT_OPTIONS + TAETIGKEIT_ALLGEMEIN_OPTIONS)
    week_matrix = {}
    for e in week_entries:
        day_idx = (e.datum - week_start).days
        if 0 <= day_idx < 5:
            key = e.taetigkeit
            if key not in week_matrix:
                week_matrix[key] = {'label': taetigkeit_labels_all.get(key, key), 'days': [0]*5, 'total': 0}
            week_matrix[key]['days'][day_idx] += e.dauer_minuten
            week_matrix[key]['total'] += e.dauer_minuten

    week_grid_rows = sorted(week_matrix.values(), key=lambda x: x['total'], reverse=True)
    week_day_totals = [sum(r['days'][d] for r in week_grid_rows) for d in range(5)]

    return render_template('timetracking.html',
                           today=today,
                           active_timer=active_timer,
                           today_minutes=today_minutes,
                           week_minutes=week_minutes,
                           month_minutes=month_minutes,
                           format_minutes=_format_minutes,
                           entries=entries,
                           contacts=contacts_list,
                           filter_zeitraum=filter_zeitraum,
                           filter_contact=filter_contact,
                           filter_taetigkeit=filter_taetigkeit,
                           taetigkeit_kontakt=TAETIGKEIT_KONTAKT_OPTIONS,
                           taetigkeit_allgemein=TAETIGKEIT_ALLGEMEIN_OPTIONS,
                           taetigkeit_stats=taetigkeit_stats,
                           week_grid_days=week_grid_days,
                           week_grid_rows=week_grid_rows,
                           week_day_totals=week_day_totals,
                           workday_today=workday_today,
                           workday_luecken=workday_luecken,
                           strang_heute=strang_heute,
                           strang_woche=strang_woche,
                           strang_monat=strang_monat,
                           custom_taetigkeiten_kontakt=_get_custom_taetigkeiten(True),
                           custom_taetigkeiten_allgemein=_get_custom_taetigkeiten(False),
                           stages=PIPELINE_STAGES)


def _parse_arbeitszeit(raw):
    if raw is None:
        return None
    s = str(raw).strip().lower().replace(',', '.')
    if not s:
        return None
    if s.endswith('m'):
        s = s[:-1].strip()
        try:
            return int(round(float(s)))
        except ValueError:
            return None
    if ':' in s:
        h, _, m = s.partition(':')
        try:
            return int(h or 0) * 60 + int(m or 0)
        except ValueError:
            return None
    try:
        val = float(s)
    except ValueError:
        return None
    return int(round(val * 60)) if val <= 24 else int(round(val))


@app.route('/workday', methods=['POST'])
def workday_save():
    datum = parse_date(request.form.get('datum')) or date.today()
    minuten = _parse_arbeitszeit(request.form.get('arbeitszeit'))
    notiz = (request.form.get('notiz') or '').strip() or None

    existing = WorkDay.query.get(datum)
    if not minuten:
        if existing:
            db.session.delete(existing)
            db.session.commit()
        return redirect(url_for('timetracking'))

    if existing:
        existing.minuten = minuten
        existing.notiz = notiz
    else:
        db.session.add(WorkDay(datum=datum, minuten=minuten, notiz=notiz))
    db.session.commit()
    return redirect(url_for('timetracking'))


@app.route('/timetracking/export')
def timetracking_export():
    today = date.today()
    week_start = today - timedelta(days=today.weekday())
    month_start = today.replace(day=1)

    filter_zeitraum = request.args.get('zeitraum', 'woche')
    filter_contact = request.args.get('contact_id', '')
    filter_taetigkeit = request.args.get('taetigkeit', '')

    query = TimeEntry.query.filter(TimeEntry.timer_start.is_(None))
    if filter_zeitraum == 'heute':
        query = query.filter(TimeEntry.datum == today)
    elif filter_zeitraum == 'woche':
        query = query.filter(TimeEntry.datum >= week_start)
    elif filter_zeitraum == 'monat':
        query = query.filter(TimeEntry.datum >= month_start)

    if filter_contact:
        query = query.filter(TimeEntry.contact_id == int(filter_contact))
    if filter_taetigkeit:
        query = query.filter(TimeEntry.taetigkeit == filter_taetigkeit)

    entries = query.order_by(TimeEntry.datum.desc(), TimeEntry.created_at.desc()).all()

    taetigkeit_labels = dict(TAETIGKEIT_KONTAKT_OPTIONS + TAETIGKEIT_ALLGEMEIN_OPTIONS)

    import csv
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(['Datum', 'Kontakt', 'Tätigkeit', 'Dauer (h:mm)', 'Dauer (min)', 'Notiz'])
    for e in entries:
        contact_name = e.contact.display_name if e.contact else 'Allgemein'
        tae_label = taetigkeit_labels.get(e.taetigkeit, e.taetigkeit)
        hours = e.dauer_minuten // 60
        mins = e.dauer_minuten % 60
        writer.writerow([
            e.datum.strftime('%d.%m.%Y'),
            contact_name,
            tae_label,
            f'{hours}:{mins:02d}',
            e.dauer_minuten,
            e.notiz or '',
        ])

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'text/csv; charset=utf-8'
    response.headers['Content-Disposition'] = f'attachment; filename=zeiterfassung_{filter_zeitraum}_{today.isoformat()}.csv'
    return response


@app.route('/timetracking/copy-last-week', methods=['POST'])
def timetracking_copy_last_week():
    today = date.today()
    this_week_start = today - timedelta(days=today.weekday())
    last_week_start = this_week_start - timedelta(days=7)
    last_week_end = this_week_start - timedelta(days=1)

    last_week_entries = TimeEntry.query.filter(
        TimeEntry.datum >= last_week_start,
        TimeEntry.datum <= last_week_end,
        TimeEntry.timer_start.is_(None)
    ).all()

    if not last_week_entries:
        flash('Keine Einträge letzte Woche gefunden', 'error')
        return redirect(url_for('timetracking'))

    created = 0
    for e in last_week_entries:
        day_offset = (e.datum - last_week_start).days
        new_date = this_week_start + timedelta(days=day_offset)
        exists = TimeEntry.query.filter_by(
            contact_id=e.contact_id, datum=new_date,
            taetigkeit=e.taetigkeit, dauer_minuten=e.dauer_minuten
        ).first()
        if not exists:
            new_entry = TimeEntry(
                contact_id=e.contact_id, datum=new_date,
                dauer_minuten=e.dauer_minuten, taetigkeit=e.taetigkeit,
                notiz=f'[Kopie] {e.notiz or ""}'.strip()
            )
            db.session.add(new_entry)
            created += 1

    db.session.commit()
    flash(f'{created} Einträge aus letzter Woche kopiert', 'success')
    return redirect(url_for('timetracking'))


@app.route('/timetracking/add', methods=['POST'])
def timetracking_add():
    data = request.form
    contact_id = data.get('contact_id') or None

    dauer_raw = data.get('dauer', '').strip()
    if ':' in dauer_raw:
        parts = dauer_raw.split(':')
        dauer_minuten = int(parts[0]) * 60 + int(parts[1])
    else:
        try:
            dauer_minuten = int(float(dauer_raw.replace(',', '.')) * 60) if '.' in dauer_raw or ',' in dauer_raw else int(dauer_raw)
        except (ValueError, TypeError):
            dauer_minuten = 0

    if dauer_minuten <= 0:
        flash('Bitte eine gültige Dauer eingeben', 'error')
        if contact_id:
            return redirect(url_for('contact_view', id=contact_id))
        return redirect(url_for('timetracking'))

    entry = TimeEntry(
        contact_id=int(contact_id) if contact_id else None,
        datum=parse_date(data.get('datum')) or date.today(),
        dauer_minuten=dauer_minuten,
        taetigkeit=data.get('taetigkeit', '').strip() or 'sonstiges',
        notiz=data.get('notiz', '').strip() or None,
    )
    db.session.add(entry)
    db.session.commit()
    flash(f'{entry.dauer_display} erfasst', 'success')

    if contact_id:
        return redirect(url_for('contact_view', id=contact_id))
    return redirect(url_for('timetracking'))


@app.route('/timetracking/<int:id>/edit', methods=['POST'])
def timetracking_edit(id):
    entry = TimeEntry.query.get_or_404(id)
    data = request.form

    dauer_raw = data.get('dauer', '').strip()
    if dauer_raw:
        if ':' in dauer_raw:
            parts = dauer_raw.split(':')
            dauer_minuten = int(parts[0]) * 60 + int(parts[1])
        else:
            try:
                dauer_minuten = int(float(dauer_raw.replace(',', '.')) * 60) if '.' in dauer_raw or ',' in dauer_raw else int(dauer_raw)
            except (ValueError, TypeError):
                dauer_minuten = 0
        if dauer_minuten > 0:
            entry.dauer_minuten = dauer_minuten

    datum = data.get('datum')
    if datum:
        parsed = parse_date(datum)
        if parsed:
            entry.datum = parsed

    taetigkeit = data.get('taetigkeit', '').strip()
    if taetigkeit:
        entry.taetigkeit = taetigkeit

    entry.notiz = data.get('notiz', '').strip() or None

    contact_id_raw = data.get('contact_id', '').strip()
    if contact_id_raw == '':
        entry.contact_id = None
    elif contact_id_raw:
        entry.contact_id = int(contact_id_raw)

    db.session.commit()
    flash('Zeiteintrag aktualisiert', 'success')

    redirect_to = data.get('redirect', '')
    if redirect_to == 'contact' and entry.contact_id:
        return redirect(url_for('contact_view', id=entry.contact_id))
    return redirect(url_for('timetracking'))


@app.route('/timetracking/<int:id>/delete', methods=['POST'])
def timetracking_delete(id):
    entry = TimeEntry.query.get_or_404(id)
    contact_id = entry.contact_id
    db.session.delete(entry)
    db.session.commit()
    flash('Zeiteintrag gelöscht', 'success')

    redirect_to = request.form.get('redirect', '')
    if redirect_to == 'contact' and contact_id:
        return redirect(url_for('contact_view', id=contact_id))
    return redirect(url_for('timetracking'))


@app.route('/timetracking/timer/start', methods=['POST'])
def timer_start():
    active = _get_active_timer()
    if active:
        elapsed = (datetime.now(timezone.utc).replace(tzinfo=None) - active.timer_start).total_seconds()
        active.dauer_minuten += max(1, int(elapsed / 60))
        active.timer_start = None

    data = request.json
    contact_id = data.get('contact_id')
    taetigkeit = data.get('taetigkeit', 'sonstiges')

    entry = TimeEntry(
        contact_id=int(contact_id) if contact_id else None,
        datum=date.today(),
        dauer_minuten=0,
        taetigkeit=taetigkeit,
        notiz=data.get('notiz', ''),
        timer_start=datetime.now(timezone.utc).replace(tzinfo=None)
    )
    db.session.add(entry)
    db.session.commit()
    return jsonify({'success': True, 'id': entry.id, 'timer_start': entry.timer_start.isoformat()})


@app.route('/timetracking/timer/stop', methods=['POST'])
def timer_stop():
    active = _get_active_timer()
    if not active:
        return jsonify({'success': False, 'error': 'Kein Timer aktiv'}), 404

    elapsed = (datetime.now(timezone.utc).replace(tzinfo=None) - active.timer_start).total_seconds()
    active.dauer_minuten += max(1, int(elapsed / 60))
    active.timer_start = None
    db.session.commit()

    return jsonify({
        'success': True,
        'id': active.id,
        'dauer_minuten': active.dauer_minuten,
        'dauer_display': active.dauer_display
    })


@app.route('/timetracking/timer/status')
def timer_status():
    active = _get_active_timer()
    if not active:
        return jsonify({'active': False})

    contact_name = active.contact.display_name if active.contact else None
    return jsonify({
        'active': True,
        'id': active.id,
        'contact_name': contact_name,
        'taetigkeit': active.taetigkeit,
        'timer_start': active.timer_start.isoformat(),
    })



@app.route('/api/retainers/<int:contact_id>')
def api_retainers(contact_id):
    retainers = Retainer.query.filter_by(contact_id=contact_id, active=True).all()
    return jsonify([{
        'id': r.id, 'name': r.name, 'monthly_fee': r.monthly_fee,
        'monthly_hours': r.monthly_hours, 'billing_interval_months': r.billing_interval_months,
        'hourly_rate': r.effective_hourly_rate, 'start_date': r.start_date.isoformat(),
        'last_billed_date': r.last_billed_date.isoformat() if r.last_billed_date else None,
    } for r in retainers])


@app.route('/api/retainers', methods=['POST'])
def api_retainer_create():
    data = request.get_json()
    r = Retainer(
        contact_id=data['contact_id'],
        name=data.get('name', ''),
        monthly_fee=data.get('monthly_fee', 0),
        monthly_hours=data.get('monthly_hours', 0),
        billing_interval_months=data.get('billing_interval_months', 6),
        hourly_rate=data.get('hourly_rate'),
        start_date=datetime.strptime(data.get('start_date', date.today().isoformat()), '%Y-%m-%d').date(),
        end_date=datetime.strptime(data['end_date'], '%Y-%m-%d').date() if data.get('end_date') else None,
    )
    db.session.add(r)
    db.session.commit()
    return jsonify({'success': True, 'id': r.id})


@app.route('/api/retainers/<int:id>/bill', methods=['POST'])
def api_retainer_bill(id):
    retainer = Retainer.query.get_or_404(id)
    contact = retainer.contact

    period_start = retainer.last_billed_date or retainer.start_date
    period_end = date.today()

    entries = TimeEntry.query.filter(
        TimeEntry.contact_id == contact.id,
        TimeEntry.datum >= period_start,
        TimeEntry.datum <= period_end,
        TimeEntry.timer_start.is_(None),
        TimeEntry.billed_document_id.is_(None),
    ).order_by(TimeEntry.datum).all()

    orphans = TimeEntry.query.filter(
        TimeEntry.contact_id == contact.id,
        TimeEntry.datum < period_start,
        TimeEntry.timer_start.is_(None),
        TimeEntry.billed_document_id.is_(None),
    ).all()
    orphan_hours = round(sum(e.dauer_minuten for e in orphans) / 60, 2)

    total_minutes = sum(e.dauer_minuten for e in entries)
    total_hours = round(total_minutes / 60, 2)
    rate = retainer.effective_hourly_rate

    months_since = ((period_end.year - period_start.year) * 12
                    + period_end.month - period_start.month)
    budget_months = months_since + 1
    included_hours = retainer.monthly_hours * budget_months

    covered_periods = {
        p for p in (
            service_period_to_yyyymm(d.service_period)
            for d in Document.query.filter(
                Document.contact_id == contact.id,
                Document.doc_type == 'invoice',
                Document.status != 'cancelled',
            ).all()
        ) if p
    }
    fee_months = plan_trueup_fee_months(retainer, period_start, period_end, covered_periods)

    _ensure_customer_number(contact)
    inv = Document(
        contact_id=contact.id,
        doc_type='invoice',
        number=f'_PENDING_{uuid.uuid4().hex[:8]}',
        subject=f'{retainer.name} — Abrechnung {period_start.strftime("%d.%m.%Y")} bis {period_end.strftime("%d.%m.%Y")}',
        status='draft',
        date=date.today(),
        payment_term_days=14,
        service_period=f'{period_start.strftime("%d.%m.%Y")} – {period_end.strftime("%d.%m.%Y")}',
        conditions_text=f'Abrechnung für {retainer.name}. Zeitraum: {period_start.strftime("%d.%m.%Y")} – {period_end.strftime("%d.%m.%Y")}.',
    )
    db.session.add(inv)
    db.session.flush()

    group_pos = 0

    if fee_months:
        g1 = LineItemGroup(document_id=inv.id, position=group_pos,
                           title=f'Pauschale ({len(fee_months)} Monat{"e" if len(fee_months) != 1 else ""})')
        db.session.add(g1)
        db.session.flush()
        db.session.add(LineItem(
            group_id=g1.id, position=0,
            title=f'{retainer.name} — Monatspauschale',
            description=f'{retainer.monthly_hours}h Guthaben/Monat · '
                        + ', '.join(period_label_de(p) for p in fee_months),
            quantity=len(fee_months), unit='Monat',
            unit_price=retainer.monthly_fee, vat_rate=20, billing_type='recurring',
        ))
        group_pos += 1

    if entries:
        g2 = LineItemGroup(document_id=inv.id, position=group_pos,
                           title=f'Geleistete Stunden ({total_hours}h)')
        db.session.add(g2)
        db.session.flush()
        for i, e in enumerate(entries):
            hours = round(e.dauer_minuten / 60, 2)
            db.session.add(LineItem(
                group_id=g2.id, position=i,
                title=e.taetigkeit,
                description=f'{e.datum.strftime("%d.%m.%Y")} — {e.dauer_display}' + (f' ({e.notiz})' if e.notiz else ''),
                quantity=hours, unit='Stunde',
                unit_price=0, vat_rate=20, billing_type='onetime',
            ))
        group_pos += 1

    if total_hours > included_hours:
        overage = round(total_hours - included_hours, 2)
        g3 = LineItemGroup(document_id=inv.id, position=group_pos,
                           title='Mehraufwand')
        db.session.add(g3)
        db.session.flush()
        db.session.add(LineItem(
            group_id=g3.id, position=0,
            title=f'Mehraufwand über Guthaben ({included_hours}h)',
            description=f'{total_hours}h geleistet – {included_hours}h inkludiert = {overage}h Mehraufwand',
            quantity=overage, unit='Stunde',
            unit_price=rate, vat_rate=20, billing_type='onetime',
        ))

    seq = NumberSequence.query.filter_by(prefix='R').first()
    inv.number = seq.generate()

    for e in entries:
        e.billed_document_id = inv.id

    if fee_months:
        newest_fee_month = max(fee_months)
        if not retainer.last_fee_period or newest_fee_month > retainer.last_fee_period:
            retainer.last_fee_period = newest_fee_month

    retainer.last_billed_date = period_end
    db.session.commit()

    if orphans:
        flash(f'Hinweis: {orphan_hours} h unverrechnete Zeit liegen VOR dem Abrechnungszeitraum '
              f'({period_start.strftime("%d.%m.%Y")}) und sind NICHT auf dieser Rechnung. '
              f'{len(orphans)} Eintrag/Eintraege — bei Bedarf separat abrechnen.', 'warning')

    return jsonify({
        'success': True, 'invoice_id': inv.id, 'number': inv.number,
        'total_hours': total_hours, 'included_hours': included_hours,
        'budget_months': budget_months, 'fee_months': fee_months,
        'orphan_hours': orphan_hours,
    })


@app.route('/api/invoice-from-time/<int:contact_id>', methods=['POST'])
def invoice_from_time(contact_id):
    contact = Contact.query.get_or_404(contact_id)
    entries = TimeEntry.query.filter(
        TimeEntry.contact_id == contact_id,
        TimeEntry.timer_start.is_(None),
        TimeEntry.billed_document_id.is_(None)
    ).order_by(TimeEntry.datum).all()

    if not entries:
        return jsonify({'success': False, 'error': 'Keine offenen Zeiteinträge'}), 400

    rate = contact.hourly_rate or _get_default_hourly_rate()
    total_minutes = sum(e.dauer_minuten for e in entries)
    total_hours = round(total_minutes / 60, 2)

    _ensure_customer_number(contact)
    inv = Document(
        contact_id=contact_id, doc_type='invoice',
        number=f'_PENDING_{uuid.uuid4().hex[:8]}',
        subject=f'Leistungsabrechnung {entries[0].datum.strftime("%d.%m.")} – {entries[-1].datum.strftime("%d.%m.%Y")}',
        status='draft', date=date.today(), payment_term_days=14,
    )
    db.session.add(inv)
    db.session.flush()

    taetigkeit_labels = dict(TAETIGKEIT_KONTAKT_OPTIONS + TAETIGKEIT_ALLGEMEIN_OPTIONS)
    g = LineItemGroup(document_id=inv.id, position=0, title=f'Geleistete Stunden ({total_hours}h)')
    db.session.add(g)
    db.session.flush()

    for i, e in enumerate(entries):
        hours = round(e.dauer_minuten / 60, 2)
        db.session.add(LineItem(
            group_id=g.id, position=i,
            title=taetigkeit_labels.get(e.taetigkeit, e.taetigkeit),
            description=f'{e.datum.strftime("%d.%m.%Y")} — {e.dauer_display}' + (f' ({e.notiz})' if e.notiz else ''),
            quantity=hours, unit='Stunde', unit_price=rate, vat_rate=20, billing_type='onetime',
        ))

    seq = NumberSequence.query.filter_by(prefix='R').first()
    inv.number = seq.generate()

    for e in entries:
        e.billed_document_id = inv.id

    db.session.commit()
    return jsonify({'success': True, 'invoice_id': inv.id, 'number': inv.number, 'hours': total_hours})



@app.route('/contacts/<int:id>/payments/add', methods=['POST'])
def payment_add(id):
    contact = Contact.query.get_or_404(id)
    data = request.form
    payment = ExpectedPayment(
        contact_id=id,
        label=data.get('label', '').strip(),
        amount=parse_form_float(data.get('amount', '')) or 0,
        expected_date=parse_date(data.get('expected_date')) or date.today(),
        payment_type=data.get('payment_type', 'setup'),
        interval_months=parse_form_int(data.get('interval_months', '')) if data.get('payment_type') == 'recurring' else None,
        status='pending',
    )
    db.session.add(payment)
    db.session.commit()
    flash('Erwartete Zahlung hinzugefügt', 'success')
    return redirect(url_for('contact_view', id=id))


@app.route('/payments/<int:id>/status', methods=['POST'])
def payment_update_status(id):
    payment = ExpectedPayment.query.get_or_404(id)
    data = request.json or request.form
    new_status = data.get('status', payment.status)
    if new_status in dict(PAYMENT_STATUSES):
        payment.status = new_status

        if new_status == 'received' and payment.payment_type == 'recurring' and payment.interval_months:
            from dateutil.relativedelta import relativedelta
            next_date = payment.expected_date + relativedelta(months=payment.interval_months)
            existing = ExpectedPayment.query.filter_by(
                contact_id=payment.contact_id,
                label=payment.label,
                expected_date=next_date,
                status='pending'
            ).first()
            if not existing:
                next_payment = ExpectedPayment(
                    contact_id=payment.contact_id,
                    label=payment.label,
                    amount=payment.amount,
                    expected_date=next_date,
                    payment_type='recurring',
                    interval_months=payment.interval_months,
                    status='pending',
                )
                db.session.add(next_payment)
                flash(f'Nächste Zahlung am {next_date.strftime("%d.%m.%Y")} automatisch angelegt', 'success')

        db.session.commit()
    if request.is_json:
        return jsonify({'success': True})
    return redirect(url_for('contact_view', id=payment.contact_id))


@app.route('/payments/<int:id>/delete', methods=['POST'])
def payment_delete(id):
    payment = ExpectedPayment.query.get_or_404(id)
    contact_id = payment.contact_id
    db.session.delete(payment)
    db.session.commit()
    flash('Zahlung gelöscht', 'success')
    return redirect(url_for('contact_view', id=contact_id))



@app.route('/contacts/<int:id>/reminder', methods=['POST'])
def reminder_set(id):
    contact = Contact.query.get_or_404(id)
    data = request.json

    reminder_date = parse_date(data.get('date', ''))
    reminder_text = data.get('text', '').strip() or None

    contact.erinnerung = reminder_date
    contact.erinnerungstext = reminder_text

    if reminder_date:
        open_system_tasks = Task.query.filter(
            Task.contact_id == id,
            Task.status != 'erledigt',
            Task.created_by == 'System',
            db.or_(Task.title.like('Follow-up:%'), Task.title.like('Nachhaken:%'))
        ).all()
        for t in open_system_tasks:
            t.status = 'erledigt'
            t.completed_at = datetime.now(timezone.utc).replace(tzinfo=None)

    db.session.commit()

    return jsonify({
        'success': True,
        'date': reminder_date.strftime('%d.%m.%Y') if reminder_date else None,
        'text': reminder_text
    })


@app.route('/contacts/<int:id>/reminder/complete', methods=['POST'])
def reminder_complete(id):
    contact = Contact.query.get_or_404(id)

    if not contact.erinnerung:
        flash('Keine Erinnerung vorhanden', 'error')
        return redirect(request.referrer or url_for('dashboard'))

    reminder_text = contact.erinnerungstext or 'Erinnerung'
    reminder_date = contact.erinnerung

    interaction = Interaction(
        contact_id=id,
        datum=date.today(),
        typ='Erinnerung',
        notiz=f'Erledigt ({reminder_date.strftime("%d.%m.%Y")}): {reminder_text}'
    )
    db.session.add(interaction)

    contact.erinnerung = None
    contact.erinnerungstext = None

    db.session.commit()
    flash(f'Erinnerung für "{contact.display_name}" erledigt', 'success')

    return redirect(request.referrer or url_for('dashboard'))



@app.route('/contacts/<int:id>/interactions/quick', methods=['POST'])
def interaction_quick(id):
    contact = Contact.query.get_or_404(id)
    data = request.get_json()
    datum = parse_date(data.get('datum'))
    typ = (data.get('typ') or '').strip()
    notiz = (data.get('notiz') or '').strip() or None
    if not datum or not typ:
        return jsonify({'success': False, 'error': 'Datum und Typ erforderlich'}), 400
    interaction = Interaction(contact_id=id, datum=datum, typ=typ, notiz=notiz)
    db.session.add(interaction)
    if not contact.kontaktiert or datum >= contact.kontaktiert:
        contact.kontaktiert = datum
        contact.wie = typ
    db.session.commit()
    return jsonify({'success': True, 'name': contact.display_name})


@app.route('/contacts/<int:id>/interactions', methods=['POST'])
def interaction_add(id):
    contact = Contact.query.get_or_404(id)

    datum = parse_date(request.form.get('datum'))
    typ = request.form.get('typ', '').strip()
    notiz = request.form.get('notiz', '').strip() or None

    if not datum or not typ:
        flash('Datum und Typ sind erforderlich', 'error')
        return redirect(url_for('contact_view', id=id))

    interaction = Interaction(contact_id=id, datum=datum, typ=typ, notiz=notiz)
    db.session.add(interaction)

    if not contact.kontaktiert or datum >= contact.kontaktiert:
        contact.kontaktiert = datum
        contact.wie = typ

    db.session.commit()
    flash('Interaktion hinzugefügt', 'success')
    return redirect(url_for('contact_view', id=id))


@app.route('/interactions/<int:id>/delete', methods=['POST'])
def interaction_delete(id):
    interaction = Interaction.query.get_or_404(id)
    contact_id = interaction.contact_id
    db.session.delete(interaction)
    db.session.commit()
    flash('Interaktion gelöscht', 'success')
    return redirect(url_for('contact_view', id=contact_id))



@app.route('/pipeline')
def pipeline():
    NEW_STAGE_LIMIT = 20
    view = request.args.get('view', 'kanban')
    include_dead = request.args.get('include_dead') == '1'

    if view == 'reminders':
        return _pipeline_reminders_view(include_dead)

    visible_stages = [(k, n) for k, n in PIPELINE_STAGES
                      if include_dead or k not in DEAD_PIPELINE_STAGES]
    contacts_by_stage = {}
    stage_values = {}
    stats = get_pipeline_stats()
    for stage_key, stage_name in visible_stages:
        base_query = Contact.query.filter_by(status=stage_key).options(
            joinedload(Contact.tasks),
            joinedload(Contact.status_changes)
        )
        if stage_key == 'new':
            stage_contacts = base_query.order_by(Contact.created_at.desc()).limit(NEW_STAGE_LIMIT).all()
            row = db.session.query(func.coalesce(func.sum(Contact.deal_value), 0)).filter_by(status='new').first()
            stage_values[stage_key] = row[0] if row else 0
        else:
            stage_contacts = base_query.all()
            stage_values[stage_key] = sum(c.deal_value or 0 for c in stage_contacts if c.deal_value)
        contacts_by_stage[stage_key] = stage_contacts
    stage_weighted = {}
    for stage_key, stage_name in visible_stages:
        stage_weighted[stage_key] = round(sum(c.weighted_value for c in contacts_by_stage[stage_key]), 2)
    new_total = stats.get('new', 0)
    return render_template('pipeline.html', contacts_by_stage=contacts_by_stage,
                           stages=visible_stages, stats=stats,
                           stage_values=stage_values, stage_weighted=stage_weighted,
                           lost_reasons=LOST_REASONS,
                           custom_lost_reasons=_get_custom_lost_reasons(),
                           today=date.today(),
                           new_stage_limit=NEW_STAGE_LIMIT,
                           new_total=new_total,
                           include_dead=include_dead,
                           dead_stages=DEAD_PIPELINE_STAGES,
                           view='kanban')


def _pipeline_reminders_view(include_dead):
    today = date.today()
    filter_mode = request.args.get('filter', 'default')

    q = Contact.query.filter(Contact.erinnerung.isnot(None))
    if not include_dead:
        q = q.filter(Contact.status.notin_(DEAD_PIPELINE_STAGES))

    if filter_mode == 'ueberfaellig':
        q = q.filter(Contact.erinnerung < today)
    elif filter_mode == 'heute':
        q = q.filter(Contact.erinnerung == today)
    elif filter_mode == 'diese_woche':
        q = q.filter(Contact.erinnerung <= today + timedelta(days=7))
    elif filter_mode == 'dieser_monat':
        q = q.filter(Contact.erinnerung <= today + timedelta(days=30))

    reminders = q.order_by(Contact.erinnerung.asc()).all()

    counts = {
        'default': Contact.query.filter(
            Contact.erinnerung.isnot(None),
            Contact.status.notin_(DEAD_PIPELINE_STAGES)
        ).count(),
        'ueberfaellig': Contact.query.filter(
            Contact.erinnerung < today,
            Contact.status.notin_(DEAD_PIPELINE_STAGES)
        ).count(),
        'heute': Contact.query.filter(
            Contact.erinnerung == today,
            Contact.status.notin_(DEAD_PIPELINE_STAGES)
        ).count(),
        'diese_woche': Contact.query.filter(
            Contact.erinnerung.isnot(None),
            Contact.erinnerung <= today + timedelta(days=7),
            Contact.status.notin_(DEAD_PIPELINE_STAGES)
        ).count(),
        'dieser_monat': Contact.query.filter(
            Contact.erinnerung.isnot(None),
            Contact.erinnerung <= today + timedelta(days=30),
            Contact.status.notin_(DEAD_PIPELINE_STAGES)
        ).count(),
    }

    stage_labels = dict(PIPELINE_STAGES)

    return render_template('pipeline.html',
                           view='reminders',
                           reminders=reminders,
                           reminder_filter=filter_mode,
                           reminder_counts=counts,
                           stage_labels=stage_labels,
                           include_dead=include_dead,
                           dead_stages=DEAD_PIPELINE_STAGES,
                           today=today,
                           stages=PIPELINE_STAGES,
                           contacts_by_stage={},
                           stats={},
                           stage_values={},
                           stage_weighted={},
                           lost_reasons=LOST_REASONS,
                           custom_lost_reasons=_get_custom_lost_reasons(),
                           new_stage_limit=20,
                           new_total=0)



@app.route('/tasks')
def tasks():
    return redirect(url_for('board'))


@app.route('/tasks/new', methods=['GET', 'POST'])
def task_new():
    contact_id = request.args.get('contact_id')

    if request.method == 'POST':
        task = Task(
            contact_id=request.form.get('contact_id'),
            title=request.form.get('title', '').strip(),
            due_date=parse_date(request.form.get('due_date'))
        )

        if not task.title or not task.contact_id:
            flash('Titel und Kontakt sind erforderlich', 'error')
            contacts_list = Contact.query.order_by(Contact.organisation, Contact.nachname).all()
            return render_template('task_form.html', task=None, contacts=contacts_list, selected_contact_id=contact_id)

        db.session.add(task)
        db.session.commit()
        flash('Aufgabe erstellt', 'success')

        if contact_id:
            return redirect(url_for('contact_view', id=contact_id))
        return redirect(url_for('tasks'))

    contacts_list = Contact.query.order_by(Contact.organisation, Contact.nachname).all()
    return render_template('task_form.html', task=None, contacts=contacts_list, selected_contact_id=contact_id)


@app.route('/tasks/<int:id>/toggle', methods=['POST'])
def task_toggle(id):
    task = Task.query.get_or_404(id)
    task.completed = not task.completed
    db.session.commit()
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({'success': True, 'completed': task.completed})
    return redirect(request.referrer or url_for('tasks'))


@app.route('/tasks/<int:id>/delete', methods=['POST'])
def task_delete(id):
    task = Task.query.get_or_404(id)
    contact_id = task.contact_id
    db.session.delete(task)
    db.session.commit()
    flash('Aufgabe gelöscht', 'success')
    if request.referrer and 'contact' in request.referrer:
        return redirect(url_for('contact_view', id=contact_id))
    return redirect(url_for('tasks'))



@app.route('/templates')
def email_templates():
    templates = EmailTemplate.query.order_by(EmailTemplate.created_at.desc()).all()
    return render_template('email_templates.html', templates=templates)


@app.route('/templates/new', methods=['GET', 'POST'])
def email_template_new():
    if request.method == 'POST':
        t = EmailTemplate(
            name=request.form.get('name', '').strip(),
            betreff=request.form.get('betreff', '').strip(),
            text=request.form.get('text', '').strip()
        )
        if not t.name or not t.betreff or not t.text:
            flash('Alle Felder sind erforderlich', 'error')
            return render_template('email_templates.html',
                                   templates=EmailTemplate.query.all(),
                                   edit_template=None, is_new=True,
                                   form_name=request.form.get('name', ''),
                                   form_betreff=request.form.get('betreff', ''),
                                   form_text=request.form.get('text', ''))
        db.session.add(t)
        db.session.commit()
        flash(f'Template "{t.name}" erstellt', 'success')
        return redirect(url_for('email_templates'))

    return render_template('email_templates.html',
                           templates=EmailTemplate.query.all(),
                           edit_template=None, is_new=True,
                           form_name='', form_betreff='', form_text='')


@app.route('/templates/<int:id>/edit', methods=['GET', 'POST'])
def email_template_edit(id):
    template = EmailTemplate.query.get_or_404(id)
    if request.method == 'POST':
        template.name = request.form.get('name', '').strip()
        template.betreff = request.form.get('betreff', '').strip()
        template.text = request.form.get('text', '').strip()
        if not template.name or not template.betreff or not template.text:
            flash('Alle Felder sind erforderlich', 'error')
        else:
            db.session.commit()
            flash(f'Template "{template.name}" aktualisiert', 'success')
            return redirect(url_for('email_templates'))

    return render_template('email_templates.html',
                           templates=EmailTemplate.query.all(),
                           edit_template=template, is_new=False,
                           form_name=template.name,
                           form_betreff=template.betreff,
                           form_text=template.text)


@app.route('/templates/<int:id>/delete', methods=['POST'])
def email_template_delete(id):
    template = EmailTemplate.query.get_or_404(id)
    name = template.name
    db.session.delete(template)
    db.session.commit()
    flash(f'Template "{name}" gelöscht', 'success')
    return redirect(url_for('email_templates'))



def _calc_conversion_rates():
    stage_order = ['new', 'contacted', 'qualified', 'proposal', 'negotiation', 'won']
    rates = []
    for i in range(len(stage_order) - 1):
        from_stage = stage_order[i]
        to_stage = stage_order[i + 1]
        from_count = StatusChange.query.filter(StatusChange.from_status == from_stage).count()
        to_count = StatusChange.query.filter(
            StatusChange.from_status == from_stage,
            StatusChange.to_status == to_stage
        ).count()
        rate = round(to_count / from_count * 100) if from_count > 0 else 0
        stage_labels = dict(PIPELINE_STAGES)
        rates.append({
            'from': from_stage, 'to': to_stage,
            'from_label': stage_labels.get(from_stage, from_stage),
            'to_label': stage_labels.get(to_stage, to_stage),
            'rate': rate, 'from_count': from_count, 'to_count': to_count
        })
    return rates


def _gather_billing_rows(today, default_rate):
    company_name = CompanySettings.get('company_name')
    retainer_by_contact = {r.contact_id: r
                           for r in Retainer.query.filter_by(active=True).all()}

    unbilled = TimeEntry.query.filter(
        TimeEntry.timer_start.is_(None),
        TimeEntry.billed_document_id.is_(None),
        TimeEntry.contact_id.isnot(None),
    ).all()
    entries_by_contact = defaultdict(list)
    for e in unbilled:
        entries_by_contact[e.contact_id].append(e)

    cids = set(retainer_by_contact) | set(entries_by_contact)
    contacts_by_id = {c.id: c for c in Contact.query.filter(Contact.id.in_(cids)).all()} if cids else {}

    rows = []
    for cid in cids:
        contact = contacts_by_id.get(cid)
        if contact is None:
            continue
        if (company_name and contact.organisation == company_name) \
                or (contact.abrechnungsmodell or '').strip() == 'intern' \
                or contact.status in DEAD_STATUSES:
            continue
        ret = retainer_by_contact.get(cid)
        status = compute_billing_status(
            contact, ret, entries_by_contact.get(cid, []), today, default_rate=default_rate)
        if status['ampel'] == 'aus':
            continue
        rows.append({'contact': contact, 'retainer_id': ret.id if ret else None, **status})
    return rows



BANK_GROUP_ORDER = ('sicher', 'unsicher', 'kein_match', 'beleg_fehlt')


def _gather_bank_rows():
    transactions = BankTransaction.query.filter(
        BankTransaction.match_status.in_(('unmatched', 'suggested'))
    ).order_by(BankTransaction.booking_date.desc()).all()

    invoices = Document.query.filter(Document.doc_type == 'invoice').all()
    incoming = IncomingInvoice.query.all()

    cids = {i.contact_id for i in invoices} | {i.supplier_id for i in incoming}
    contacts_by_id = ({c.id: c for c in Contact.query.filter(Contact.id.in_(cids)).all()}
                      if cids else {})
    docs_by_id = {d.id: d for d in invoices}
    incoming_by_id = {i.id: i for i in incoming}

    groups = {key: [] for key in BANK_GROUP_ORDER}
    for tx in transactions:
        result = match_bank_transaction(tx, invoices, incoming, contacts_by_id)
        for cand in result['candidates']:
            beleg = (docs_by_id.get(cand['id']) if cand['kind'] == 'invoice'
                     else incoming_by_id.get(cand['id']))
            cand['beleg'] = beleg
            if beleg is not None:
                cid = getattr(beleg, 'contact_id', None) or getattr(beleg, 'supplier_id', None)
                cand['kunde'] = contacts_by_id.get(cid)

        row = {'tx': tx, **result}
        if result['direction'] == 'out' and result['confidence'] == 'kein_match':
            groups['beleg_fehlt'].append(row)
        else:
            groups[result['confidence']].append(row)
    return groups


@app.route('/bank')
def bank():
    groups = _gather_bank_rows()
    kpi = {
        'sicher': len(groups['sicher']),
        'unsicher': len(groups['unsicher']),
        'kein_match': len(groups['kein_match']),
        'beleg_fehlt': len(groups['beleg_fehlt']),
        'beleg_fehlt_eur': round(sum(r['tx'].abs_amount
                                     for r in groups['beleg_fehlt']), 2),
    }
    imports = BankImport.query.order_by(BankImport.imported_at.desc()).limit(10).all()
    open_invoices = Document.query.filter(
        Document.doc_type == 'invoice',
        Document.status.in_(OPEN_INVOICE_STATUSES),
    ).order_by(Document.date).all()
    return render_template('bank.html', groups=groups, kpi=kpi, imports=imports,
                           open_invoices=open_invoices, today=date.today())


@app.route('/bank/import', methods=['POST'])
def bank_import_upload():
    upload = request.files.get('bank_file')
    if upload is None or not upload.filename:
        flash('Keine Datei ausgewählt.', 'error')
        return redirect(url_for('bank'))

    if not upload.filename.lower().endswith('.csv'):
        flash('Nur CSV-Dateien. N26 exportiert Umsätze als CSV, kein CAMT.', 'error')
        return redirect(url_for('bank'))
    filename = secure_filename(upload.filename) or 'bank-export.csv'

    try:
        text = upload.read().decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            upload.seek(0)
            text = upload.read().decode('cp1252')
        except (UnicodeDecodeError, OSError):
            flash('Die Datei ist weder UTF-8 noch Windows-1252 lesbar.', 'error')
            return redirect(url_for('bank'))

    try:
        rows, schema_name = parse_csv(text)
    except BankFormatError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('bank'))

    if not rows:
        flash('Die Datei enthält keine verwertbaren Buchungszeilen.', 'error')
        return redirect(url_for('bank'))

    result = stellwerk_write.import_bank_transactions(
        rows, filename, schema_name=schema_name, commit=True)
    if not result.ok:
        flash('Import fehlgeschlagen: ' + ' · '.join(result.errors), 'error')
        return redirect(url_for('bank'))

    meta = result.meta
    flash(f'{meta["row_count"]} Zeilen gelesen, {meta["new"]} neu, '
          f'{meta["duplicates"]} Duplikate übersprungen. '
          f'Zeitraum {meta["period_from"]} bis {meta["period_to"]} '
          f'(Schema {schema_name}).', 'success')
    for warning in result.warnings:
        flash(warning, 'info')
    return redirect(url_for('bank'))


@app.route('/bank/confirm/<int:tx_id>', methods=['POST'])
def bank_confirm(tx_id):
    target = request.form.get('target') or ''
    if ':' not in target:
        flash('Kein Beleg ausgewählt.', 'error')
        return redirect(url_for('bank'))
    kind, _, raw_id = target.partition(':')
    try:
        beleg_id = int(raw_id)
    except ValueError:
        flash('Ungültige Beleg-Auswahl.', 'error')
        return redirect(url_for('bank'))

    kwargs = ({'document_id': beleg_id} if kind == 'invoice'
              else {'incoming_invoice_id': beleg_id})
    result = stellwerk_write.confirm_bank_match(tx_id, commit=True, **kwargs)
    if result.ok:
        flash(result.summary.replace('\n', ' · '), 'success')
    else:
        flash(' · '.join(result.errors), 'error')
    return redirect(url_for('bank'))


@app.route('/bank/ignore/<int:tx_id>', methods=['POST'])
def bank_ignore(tx_id):
    note = (request.form.get('note') or '').strip() or None
    result = stellwerk_write.ignore_bank_transaction(tx_id, note=note, commit=True)
    if result.ok:
        flash(result.summary, 'success')
    else:
        flash(' · '.join(result.errors), 'error')
    return redirect(url_for('bank'))


@app.route('/abrechnung')
def abrechnung():
    today = date.today()
    rows = _gather_billing_rows(today, _get_default_hourly_rate())

    order = {'rot': 0, 'gelb': 1, 'terminiert': 2, 'grau': 3, 'gruen': 4}
    rows.sort(key=lambda r: (order.get(r['ampel'], 9), -(r['delivered_h'] or 0)))

    kpi = {
        'rot': sum(1 for r in rows if r['ampel'] == 'rot'),
        'gelb': sum(1 for r in rows if r['ampel'] == 'gelb'),
        'value_open': round(sum(r['value_eur'] or 0
                                for r in rows if r['action'] == 'time'), 2),
    }
    return render_template('abrechnung.html', rows=rows, kpi=kpi, today=today)


@app.route('/abrechnung/run-monthly', methods=['POST'])
def abrechnung_run_monthly():
    target_period = request.form.get('target_period') or previous_period(date.today())
    retainers = Retainer.query.filter_by(active=True).all()
    created, skipped = [], []
    for r in retainers:
        existing = {
            p for p in (
                service_period_to_yyyymm(d.service_period)
                for d in Document.query.filter(
                    Document.contact_id == r.contact_id,
                    Document.doc_type == 'invoice',
                    Document.status != 'cancelled',
                ).all()
            ) if p
        }
        plan = plan_monthly_fee_draft(r, target_period, existing)
        if plan['action'] == 'create' and (not r.monthly_fee or r.monthly_fee <= 0):
            plan = {'action': 'skip', 'reason': 'keine Pauschale (0 EUR)', 'amount': 0.0}
        if plan['action'] == 'skip':
            skipped.append((r.name, plan['reason']))
            continue

        contact = r.contact
        _ensure_customer_number(contact)
        label = period_label_de(target_period)
        inv = Document(
            contact_id=contact.id, doc_type='invoice',
            number=f'_PENDING_{uuid.uuid4().hex[:8]}',
            subject=f'Monatspauschale {label}',
            status='draft', date=date.today(), payment_term_days=14,
            service_period=period_service_string(target_period),
        )
        db.session.add(inv)
        db.session.flush()
        g = LineItemGroup(document_id=inv.id, position=0, title=f'Monatspauschale {label}')
        db.session.add(g)
        db.session.flush()
        db.session.add(LineItem(
            group_id=g.id, position=0,
            title=f'Monatspauschale {label}',
            description=r.name,
            quantity=1, unit='Monat',
            unit_price=r.monthly_fee, vat_rate=20, billing_type='recurring',
        ))
        seq = NumberSequence.query.filter_by(prefix='R').first()
        inv.number = seq.generate()
        r.last_fee_period = target_period
        created.append((r.name, inv.number))

    db.session.commit()
    msgs = []
    if created:
        msgs.append(f'{len(created)} Vormonats-Entwurf/-Entwuerfe fuer {period_label_de(target_period)} '
                    f'angelegt: {", ".join(n for n, _ in created)}.')
    if skipped:
        msgs.append('Uebersprungen: '
                    + '; '.join(f'{name} ({reason})' for name, reason in skipped) + '.')
    if not created and not skipped:
        msgs.append(f'Keine aktiven Retainer fuer {period_label_de(target_period)}.')
    flash(' '.join(msgs), 'success' if created else 'info')
    return redirect(url_for('abrechnung'))


def _find_partneragentur_parent():
    return (Contact.query
            .filter(Contact.abrechnungsmodell == 'retainer',
                    func.lower(Contact.organisation).like('%partneragentur%'))
            .first())


@app.route('/abrechnung/profitabilitaet')
def abrechnung_profitabilitaet():
    today = date.today()
    zeitraum = request.args.get('zeitraum', '12m')
    if zeitraum not in ('12m', 'gesamt'):
        zeitraum = '12m'
    window_start = None if zeitraum == 'gesamt' else _months_ago(today, 12)
    default_rate = _get_default_hourly_rate()
    company_name = CompanySettings.get('company_name')

    paid = Document.query.filter(
        Document.doc_type == 'invoice',
        Document.status == 'paid',
        Document.contact_id.isnot(None),
    ).all()
    entries = TimeEntry.query.filter(
        TimeEntry.timer_start.is_(None),
        TimeEntry.contact_id.isnot(None),
    ).all()

    inv_by_contact = defaultdict(list)
    for d in paid:
        inv_by_contact[d.contact_id].append(d)
    ent_by_contact = defaultdict(list)
    for e in entries:
        ent_by_contact[e.contact_id].append(e)

    cids = set(inv_by_contact) | set(ent_by_contact)
    contacts_by_id = ({c.id: c for c in Contact.query.filter(Contact.id.in_(cids)).all()}
                      if cids else {})

    partneragentur_parent = _find_partneragentur_parent()
    all_partner_ids = {cid for cid, c in contacts_by_id.items()
                       if (c.abrechnungsmodell or '').strip() == 'agentur_partner'}
    active_children = [contacts_by_id[cid] for cid in all_partner_ids
                       if contacts_by_id[cid].status not in DEAD_STATUSES]
    rollup_active = partneragentur_parent is not None
    skip_ids = (all_partner_ids | {partneragentur_parent.id}) if rollup_active else set()

    segments = {'regulaer': [], 'pauschale': [], 'partner': [],
                'akquise': [], 'pro_bono': []}
    for cid in cids:
        if cid in skip_ids:
            continue
        contact = contacts_by_id.get(cid)
        if contact is None:
            continue
        if (company_name and contact.organisation == company_name) \
                or (contact.abrechnungsmodell or '').strip() == 'intern':
            continue
        p = compute_profitability(
            contact, inv_by_contact.get(cid, []), ent_by_contact.get(cid, []),
            window_start=window_start, default_rate=default_rate)
        if p['kind'] == 'leer':
            continue
        segments[p['segment']].append({'contact': contact, **p})

    if rollup_active:
        roll = compute_partneragentur_rollup(
            partneragentur_parent, active_children, inv_by_contact, ent_by_contact,
            window_start=window_start, default_rate=default_rate)
        if roll['kind'] != 'leer':
            segments['partner'].append({'contact': partneragentur_parent,
                                        'label': 'partneragentur (Kooperation)', **roll})

    kind_order = {'voll': 0, 'nur_zeit': 1, 'nur_umsatz': 2}
    for rows in segments.values():
        rows.sort(key=lambda r: (kind_order.get(r['kind'], 9),
                                 r['effective_rate'] if r['effective_rate'] is not None else 0,
                                 -(r['hours'] or 0)))

    def _seg_summary(rows):
        rev = round(sum(r['revenue'] for r in rows), 2)
        hrs = round(sum(r['hours'] for r in rows), 2)
        cost = round(sum(r.get('notional_cost') or 0 for r in rows), 2)
        return {'count': len(rows), 'revenue': rev, 'hours': hrs,
                'blended_rate': round(rev / hrs, 2) if hrs > 0 else None,
                'notional_cost': cost}

    summary = {seg: _seg_summary(rows) for seg, rows in segments.items()}

    return render_template('abrechnung_profitabilitaet.html',
                           segments=segments, summary=summary,
                           kosten_segmente=PROFIT_KOSTEN_SEGMENTE,
                           zeitraum=zeitraum, today=today, default_rate=default_rate)


def _saldo_ist_verrechenbar(contact):
    return (contact.abrechnungsmodell or '').strip() == 'retainer_trueup'


@app.route('/abrechnung/trend')
def abrechnung_trend():
    today = date.today()
    fenster = request.args.get('fenster', '12')
    if fenster not in ('12', '24'):
        fenster = '12'
    months = month_range(_months_ago(today, int(fenster) - 1), today)
    window_first = date(int(months[0][:4]), int(months[0][5:7]), 1)

    entries = TimeEntry.query.filter(
        TimeEntry.timer_start.is_(None),
        TimeEntry.contact_id.isnot(None),
        TimeEntry.datum >= window_first,
    ).all()
    entries_by_contact = defaultdict(list)
    for e in entries:
        entries_by_contact[e.contact_id].append(e)

    retainer_by_contact = {r.contact_id: r
                           for r in Retainer.query.filter_by(active=True).all()}

    paid = Document.query.filter(
        Document.doc_type == 'invoice',
        Document.status == 'paid',
        Document.contact_id.isnot(None),
        Document.payment_date >= window_first,
    ).all()
    revenue_by_contact = defaultdict(float)
    invoices_by_contact = defaultdict(int)
    for d in paid:
        revenue_by_contact[d.contact_id] += d.total_net
        invoices_by_contact[d.contact_id] += 1

    company_name = CompanySettings.get('company_name')
    candidates = Contact.query.filter(
        Contact.abrechnungsmodell.in_(sorted(TREND_BILLING_MODELS))
    ).all()

    partneragentur_parent = _find_partneragentur_parent()
    rollup_active = partneragentur_parent is not None
    children = (Contact.query.filter(
        Contact.abrechnungsmodell == 'agentur_partner',
        Contact.status.notin_(sorted(DEAD_STATUSES)),
    ).all() if rollup_active else [])

    rows = []
    for contact in candidates:
        if rollup_active and contact.id == partneragentur_parent.id:
            continue
        if (company_name and contact.organisation == company_name) \
                or contact.status in DEAD_STATUSES:
            continue
        trend = compute_utilization_trend(
            retainer_by_contact.get(contact.id), entries_by_contact.get(contact.id, []), months)
        if trend['hours_total'] == 0 and trend['monthly_hours'] is None:
            continue
        rows.append({'contact': contact, 'label': contact.display_name,
                     'retainer': retainer_by_contact.get(contact.id),
                     'contributors': [],
                     'revenue_window': round(revenue_by_contact.get(contact.id, 0.0), 2),
                     'invoice_count': invoices_by_contact.get(contact.id, 0),
                     'saldo_verrechenbar': _saldo_ist_verrechenbar(contact),
                     **trend})

    if rollup_active:
        parent_ret = retainer_by_contact.get(partneragentur_parent.id)
        agg_entries = [e for cid in [partneragentur_parent.id] + [c.id for c in children]
                       for e in entries_by_contact.get(cid, [])]
        roll = compute_utilization_trend(parent_ret, agg_entries, months)
        contributors = []
        for child in children:
            ct = compute_utilization_trend(None, entries_by_contact.get(child.id, []), months)
            if ct['hours_total'] > 0:
                contributors.append({'name': child.display_name, **ct})
        contributors.sort(key=lambda c: -c['hours_total'])
        if roll['hours_total'] > 0 or roll['monthly_hours'] is not None:
            agg_ids = [partneragentur_parent.id] + [c.id for c in children]
            rows.append({'contact': partneragentur_parent, 'label': 'partneragentur (Kooperation)',
                         'retainer': parent_ret, 'contributors': contributors,
                         'revenue_window': round(sum(revenue_by_contact.get(i, 0.0)
                                                     for i in agg_ids), 2),
                         'invoice_count': sum(invoices_by_contact.get(i, 0) for i in agg_ids),
                         'saldo_verrechenbar': _saldo_ist_verrechenbar(partneragentur_parent),
                         **roll})

    rows.sort(key=lambda r: (0 if r['saldo_total'] is not None else 1,
                             r['saldo_total'] if r['saldo_total'] is not None else 0,
                             -r['hours_total']))

    kpi = {
        'mandate': len(rows),
        'hours_total': round(sum(r['hours_total'] for r in rows), 2),
        'over_budget': sum(1 for r in rows
                           if r['saldo_total'] is not None and r['saldo_total'] < 0),
        'saldo_total': round(sum(r['saldo_total'] for r in rows
                                 if r['saldo_total'] is not None), 2),
        'saldo_verrechenbar': round(sum(r['saldo_total'] for r in rows
                                        if r['saldo_total'] is not None
                                        and r['saldo_verrechenbar']), 2),
    }
    return render_template('abrechnung_trend.html', rows=rows, months=months,
                           kpi=kpi, fenster=fenster, today=today)


@app.route('/analytics')
def analytics():
    conversion_rates = _calc_conversion_rates()
    hourly_rate = _get_default_hourly_rate()

    stage_durations = []
    stage_labels = dict(PIPELINE_STAGES)
    target_stages = ['new', 'contacted', 'qualified', 'proposal', 'negotiation']

    all_changes = StatusChange.query.filter(
        db.or_(
            StatusChange.from_status.in_(target_stages),
            StatusChange.to_status.in_(target_stages)
        )
    ).order_by(StatusChange.changed_at).all()

    entries_by_contact_stage = {}
    exits_by_stage = {s: [] for s in target_stages}
    for sc in all_changes:
        if sc.to_status in target_stages:
            key = (sc.contact_id, sc.to_status)
            entries_by_contact_stage[key] = sc.changed_at
        if sc.from_status in target_stages:
            exits_by_stage[sc.from_status].append(sc)

    for stage_key in target_stages:
        durations = []
        changes_out = exits_by_stage[stage_key]
        for sc in changes_out:
            entry_time = entries_by_contact_stage.get((sc.contact_id, stage_key))
            if entry_time and entry_time < sc.changed_at:
                days = (sc.changed_at - entry_time).days
                if days >= 0:
                    durations.append(days)
        avg_days = round(sum(durations) / len(durations), 1) if durations else 0
        stage_durations.append({
            'stage': stage_key,
            'label': stage_labels.get(stage_key, stage_key),
            'avg_days': avg_days,
            'count': len(changes_out)
        })

    stuck_contacts = []
    now = datetime.now()
    active_stages = [sd['stage'] for sd in stage_durations if sd['avg_days'] > 0]
    if active_stages:
        from sqlalchemy import and_
        latest_entry_sub = db.session.query(
            StatusChange.contact_id,
            func.max(StatusChange.changed_at).label('last_entered')
        ).join(Contact, and_(
            Contact.id == StatusChange.contact_id,
            Contact.status == StatusChange.to_status,
            Contact.status.in_(active_stages)
        )).group_by(StatusChange.contact_id).subquery()

        stuck_rows = db.session.query(Contact, latest_entry_sub.c.last_entered).join(
            latest_entry_sub, Contact.id == latest_entry_sub.c.contact_id
        ).filter(Contact.status.in_(active_stages)).all()

        stage_thresholds = {sd['stage']: sd['avg_days'] * 2 for sd in stage_durations if sd['avg_days'] > 0}
        for c, last_entered in stuck_rows:
            if last_entered:
                days_in = (now - last_entered).days
                threshold = stage_thresholds.get(c.status, 0)
                if days_in > threshold:
                    stuck_contacts.append({
                        'contact': c,
                        'stage_label': stage_labels.get(c.status, c.status),
                        'days_in': days_in,
                        'avg': threshold / 2
                    })

    quellen_periode = request.args.get('quellen_periode', 'monat')
    quellen_labels = dict(QUELLEN_OPTIONS)
    today = date.today()

    def _period_bounds(ref_date, periode):
        if periode == 'jahr':
            return ref_date.replace(month=1, day=1), str(ref_date.year)
        elif periode == 'quartal':
            q_month = ((ref_date.month - 1) // 3) * 3 + 1
            q_num = (ref_date.month - 1) // 3 + 1
            return ref_date.replace(month=q_month, day=1), f'Q{q_num} {ref_date.year}'
        else:
            return ref_date.replace(day=1), ref_date.strftime('%b %Y')

    def _prev_period_ref(p_start, periode):
        return p_start - timedelta(days=1)

    period_start, period_label = _period_bounds(today, quellen_periode)

    def _quellen_for_period(p_start):
        p_start_dt = datetime.combine(p_start, datetime.min.time())

        kontakte_q = db.session.query(
            Contact.quelle, func.count(Contact.id)
        ).filter(
            Contact.created_at >= p_start_dt,
            Contact.quelle.isnot(None), Contact.quelle != ''
        ).group_by(Contact.quelle).all()

        won_q = []
        won_ch = StatusChange.query.filter(
            StatusChange.to_status == 'won',
            StatusChange.changed_at >= p_start_dt
        ).all()
        won_ids = [sc.contact_id for sc in won_ch]
        if won_ids:
            won_q = db.session.query(
                Contact.quelle, func.count(Contact.id), func.sum(Contact.deal_value)
            ).filter(
                Contact.id.in_(won_ids),
                Contact.quelle.isnot(None), Contact.quelle != ''
            ).group_by(Contact.quelle).all()

        zeit_q = db.session.query(
            Contact.quelle, func.sum(TimeEntry.dauer_minuten)
        ).join(Contact, TimeEntry.contact_id == Contact.id).filter(
            TimeEntry.datum >= p_start,
            TimeEntry.timer_start.is_(None),
            Contact.quelle.isnot(None), Contact.quelle != ''
        ).group_by(Contact.quelle).all()

        zeit_allgemein = db.session.query(
            func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
        ).filter(
            TimeEntry.datum >= p_start,
            TimeEntry.timer_start.is_(None),
            TimeEntry.contact_id.is_(None)
        ).scalar() or 0

        kontakte_d = dict(kontakte_q)
        won_d = {q: (cnt, val or 0) for q, cnt, val in won_q}
        zeit_d = dict(zeit_q)

        all_q = sorted(set(
            [q for q, _ in kontakte_q] +
            [q for q, _, _ in won_q] +
            [q for q, _ in zeit_q]
        ))

        result = []
        for q in all_q:
            w_cnt, w_val = won_d.get(q, (0, 0))
            minuten = zeit_d.get(q, 0) or 0
            zeitkosten = round(minuten / 60 * hourly_rate, 2)
            result.append({
                'key': q,
                'label': quellen_labels.get(q, q),
                'kontakte': kontakte_d.get(q, 0),
                'won': w_cnt,
                'won_value': w_val,
                'minuten': minuten,
                'zeitkosten': zeitkosten,
                'roi': round(w_val - zeitkosten, 2) if w_val else None,
            })

        if zeit_allgemein > 0:
            result.append({
                'key': '_allgemein',
                'label': 'Allgemein (ohne Kontakt)',
                'kontakte': 0,
                'won': 0,
                'won_value': 0,
                'minuten': zeit_allgemein,
                'zeitkosten': round(zeit_allgemein / 60 * hourly_rate, 2),
                'roi': None,
            })

        return result

    quellen_stats = _quellen_for_period(period_start)

    def _collect_history(periode, num_periods):
        history = []
        ref = today
        for _ in range(num_periods):
            p_start, p_label = _period_bounds(ref, periode)
            p_start_dt = datetime.combine(p_start, datetime.min.time())

            kontakte_cnt = Contact.query.filter(
                Contact.created_at >= p_start_dt,
                Contact.quelle.isnot(None), Contact.quelle != ''
            ).count()

            won_ch = StatusChange.query.filter(
                StatusChange.to_status == 'won',
                StatusChange.changed_at >= p_start_dt
            ).all()
            won_ids = [sc.contact_id for sc in won_ch]
            won_cnt = 0
            won_val = 0
            if won_ids:
                row = db.session.query(
                    func.count(Contact.id), func.coalesce(func.sum(Contact.deal_value), 0)
                ).filter(Contact.id.in_(won_ids)).first()
                if row:
                    won_cnt, won_val = row

            zeit_total = db.session.query(
                func.coalesce(func.sum(TimeEntry.dauer_minuten), 0)
            ).filter(
                TimeEntry.datum >= p_start,
                TimeEntry.timer_start.is_(None)
            ).scalar() or 0

            history.append({
                'label': p_label,
                'kontakte': kontakte_cnt,
                'won': won_cnt,
                'won_value': float(won_val),
                'minuten': zeit_total,
                'zeitkosten': round(zeit_total / 60 * hourly_rate, 2),
            })

            ref = _prev_period_ref(p_start, periode)

        history.reverse()
        return history

    if quellen_periode == 'monat':
        quellen_history = _collect_history('monat', 6)
    elif quellen_periode == 'quartal':
        quellen_history = _collect_history('quartal', 4)
    else:
        quellen_history = _collect_history('jahr', 3)

    doc_durations = []
    accepted_quotes = Document.query.filter(
        Document.doc_type == 'quote', Document.status.in_(QUOTE_WON_STATUSES)
    ).all()
    if accepted_quotes:
        q_days = []
        for d in accepted_quotes:
            accept_log = DocumentStatusLog.query.filter(
                DocumentStatusLog.document_id == d.id,
                DocumentStatusLog.to_status == 'accepted'
            ).first()
            if accept_log:
                q_days.append((accept_log.changed_at.date() - d.date).days)
        doc_durations.append({
            'label': 'Angebot → Angenommen',
            'avg_days': round(sum(q_days) / len(q_days), 1) if q_days else 0,
            'count': len(q_days),
        })

    paid_invoices = Document.query.filter(
        Document.doc_type == 'invoice', Document.status == 'paid',
        Document.sent_date.isnot(None), Document.payment_date.isnot(None)
    ).all()
    if paid_invoices:
        inv_days = [(d.payment_date - d.sent_date).days for d in paid_invoices if d.payment_date >= d.sent_date]
        doc_durations.append({
            'label': 'Rechnung gesendet → Bezahlt',
            'avg_days': round(sum(inv_days) / len(inv_days), 1) if inv_days else 0,
            'count': len(inv_days),
        })

    slow_payers = []
    if paid_invoices:
        payer_days = {}
        for d in paid_invoices:
            if d.payment_date >= d.sent_date:
                days = (d.payment_date - d.sent_date).days
                payer_days.setdefault(d.contact_id, []).append(days)
        for cid, days_list in payer_days.items():
            avg = round(sum(days_list) / len(days_list), 1)
            if avg > 0:
                c = Contact.query.get(cid)
                if c:
                    slow_payers.append({'contact': c, 'avg_days': avg, 'count': len(days_list)})
        slow_payers.sort(key=lambda x: x['avg_days'], reverse=True)
        slow_payers = slow_payers[:5]

    quote_conversion = compute_quote_conversion(
        Document.query.filter(Document.doc_type == 'quote').all(),
        Document.query.filter(Document.doc_type == 'invoice').all(),
    )

    interaction_stats = db.session.query(
        Interaction.typ, func.count(Interaction.id)
    ).group_by(Interaction.typ).order_by(func.count(Interaction.id).desc()).all()

    interaction_monthly = []
    for i in range(5, -1, -1):
        m_start = (today.replace(day=1) - timedelta(days=30 * i)).replace(day=1)
        if m_start.month == 12:
            m_end = m_start.replace(year=m_start.year + 1, month=1, day=1) - timedelta(days=1)
        else:
            m_end = m_start.replace(month=m_start.month + 1, day=1) - timedelta(days=1)
        cnt = Interaction.query.filter(
            Interaction.datum >= m_start, Interaction.datum <= m_end
        ).count()
        interaction_monthly.append({
            'label': m_start.strftime('%b %Y'),
            'label_short': m_start.strftime('%b'),
            'count': cnt,
        })
    max_interactions = max((m['count'] for m in interaction_monthly), default=1) or 1

    interaction_win_rates = []
    for typ, total_cnt in interaction_stats:
        contact_ids_with_type = db.session.query(Interaction.contact_id).filter(
            Interaction.typ == typ
        ).distinct()
        contacts_with = Contact.query.filter(Contact.id.in_(contact_ids_with_type.subquery().select())).count()
        won_with = Contact.query.filter(
            Contact.id.in_(contact_ids_with_type.subquery().select()), Contact.status == 'won'
        ).count()
        if contacts_with > 0:
            interaction_win_rates.append({
                'typ': typ,
                'contacts': contacts_with,
                'won': won_with,
                'rate': round(won_with / contacts_with * 100) if contacts_with > 0 else 0,
            })
    interaction_win_rates.sort(key=lambda x: x['rate'], reverse=True)

    return render_template('analytics.html',
                           conversion_rates=conversion_rates,
                           stage_durations=stage_durations,
                           stuck_contacts=stuck_contacts,
                           quellen_stats=quellen_stats,
                           quellen_historie=quellen_history,
                           quellen_periode=quellen_periode,
                           quellen_period_label=period_label,
                           default_hourly_rate=hourly_rate,
                           doc_durations=doc_durations,
                           slow_payers=slow_payers,
                           interaction_stats=interaction_stats,
                           interaction_monthly=interaction_monthly,
                           max_interactions=max_interactions,
                           interaction_win_rates=interaction_win_rates,
                           quote_conversion=quote_conversion,
                           format_minutes=_format_minutes,
                           stages=PIPELINE_STAGES)



@app.route('/export', methods=['GET', 'POST'])
def export_data():
    if request.method == 'POST' or any(request.args.get(k) for k in ('status', 'land', 'quelle', 'branche', 'prioritaet')):
        query = Contact.query

        status_filter = request.form.get('status', '') or request.args.get('status', '')
        land_filter = request.form.get('land', '') or request.args.get('land', '')
        quelle_filter = request.form.get('quelle', '') or request.args.get('quelle', '')
        branche_filter = request.form.get('branche', '') or request.args.get('branche', '')
        prioritaet_filter = request.form.get('prioritaet', '') or request.args.get('prioritaet', '')

        if status_filter:
            query = query.filter(Contact.status == status_filter)
        if land_filter:
            query = query.filter(Contact.land == land_filter)
        if quelle_filter:
            query = query.filter(Contact.quelle == quelle_filter)
        if branche_filter:
            query = query.filter(Contact.branche == branche_filter)
        if prioritaet_filter:
            query = query.filter(Contact.prioritaet == prioritaet_filter)

        contacts_list = query.order_by(Contact.organisation, Contact.nachname).all()

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = 'Kontakte'

        headers = [
            'Organisation', 'Organisationsform', 'Website', 'Mail Organisation',
            'Telefon Organisation', 'In a Nutshell', 'Steuernummer',
            'Land', 'Bundesland', 'PLZ', 'Stadt',
            'Sprache 1', 'Sprache 2', 'Vorname', 'Nachname', 'Mail Person',
            'Telefon Person', 'Position', 'Anrede', 'Anrede EN', 'Titel Vor', 'Titel Nach',
            'LinkedIn', 'Kontaktiert', 'Wie', 'Kommentare', 'Erinnerung', 'Erinnerungstext',
            'Status', 'Deal Wert', 'Wahrscheinlichkeit %', 'Gewichteter Wert',
            'Quelle', 'Branche', 'Priorität', 'Lost Reason', 'Erstellt', 'Aktualisiert'
        ]

        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        for row_idx, c in enumerate(contacts_list, 2):
            ws.cell(row=row_idx, column=1, value=c.organisation)
            ws.cell(row=row_idx, column=2, value=c.organisationsform)
            ws.cell(row=row_idx, column=3, value=c.website)
            ws.cell(row=row_idx, column=4, value=c.mail_organisation)
            ws.cell(row=row_idx, column=5, value=c.telefon_organisation)
            ws.cell(row=row_idx, column=6, value=c.in_a_nutshell)
            ws.cell(row=row_idx, column=7, value=c.steuernummer)
            ws.cell(row=row_idx, column=8, value=c.land)
            ws.cell(row=row_idx, column=9, value=c.bundesland)
            ws.cell(row=row_idx, column=10, value=c.plz)
            ws.cell(row=row_idx, column=11, value=c.stadt)
            ws.cell(row=row_idx, column=12, value=c.sprache_1)
            ws.cell(row=row_idx, column=13, value=c.sprache_2)
            ws.cell(row=row_idx, column=14, value=c.vorname)
            ws.cell(row=row_idx, column=15, value=c.nachname)
            ws.cell(row=row_idx, column=16, value=c.mail_person)
            ws.cell(row=row_idx, column=17, value=c.telefon_person)
            ws.cell(row=row_idx, column=18, value=c.position)
            ws.cell(row=row_idx, column=19, value=c.anrede)
            ws.cell(row=row_idx, column=20, value=c.anrede_en)
            ws.cell(row=row_idx, column=21, value=c.titel_vor)
            ws.cell(row=row_idx, column=22, value=c.titel_nach)
            ws.cell(row=row_idx, column=23, value=c.linkedin)
            ws.cell(row=row_idx, column=24, value=c.kontaktiert.isoformat() if c.kontaktiert else None)
            ws.cell(row=row_idx, column=25, value=c.wie)
            ws.cell(row=row_idx, column=26, value=c.kommentare)
            ws.cell(row=row_idx, column=27, value=c.erinnerung.isoformat() if c.erinnerung else None)
            ws.cell(row=row_idx, column=28, value=c.erinnerungstext)
            ws.cell(row=row_idx, column=29, value=c.status)
            ws.cell(row=row_idx, column=30, value=c.deal_value)
            ws.cell(row=row_idx, column=31, value=c.deal_probability)
            ws.cell(row=row_idx, column=32, value=round(c.weighted_value, 2) if c.weighted_value else None)
            ws.cell(row=row_idx, column=33, value=c.quelle)
            ws.cell(row=row_idx, column=34, value=c.branche)
            ws.cell(row=row_idx, column=35, value=c.prioritaet)
            ws.cell(row=row_idx, column=36, value=c.lost_reason)
            ws.cell(row=row_idx, column=37, value=c.created_at.strftime('%Y-%m-%d %H:%M') if c.created_at else None)
            ws.cell(row=row_idx, column=38, value=c.updated_at.strftime('%Y-%m-%d %H:%M') if c.updated_at else None)

        for col in ws.columns:
            max_length = 0
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

        ws2 = wb.create_sheet('Zeiterfassung')
        time_headers = ['Datum', 'Kontakt', 'Tätigkeit', 'Dauer (h:mm)', 'Dauer (Minuten)', 'Notiz']
        for col_idx, header in enumerate(time_headers, 1):
            cell = ws2.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill

        all_time_entries = TimeEntry.query.filter(
            TimeEntry.timer_start.is_(None)
        ).order_by(TimeEntry.datum.desc()).all()

        for row_idx, te in enumerate(all_time_entries, 2):
            ws2.cell(row=row_idx, column=1, value=te.datum.isoformat() if te.datum else None)
            ws2.cell(row=row_idx, column=2, value=te.contact.display_name if te.contact else 'Allgemein')
            ws2.cell(row=row_idx, column=3, value=te.taetigkeit)
            ws2.cell(row=row_idx, column=4, value=te.dauer_display)
            ws2.cell(row=row_idx, column=5, value=te.dauer_minuten)
            ws2.cell(row=row_idx, column=6, value=te.notiz)

        for col in ws2.columns:
            max_length = 0
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws2.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

        ws3 = wb.create_sheet('Finanz-Report')
        fin_headers = [
            'Kontakt', 'Status', 'Quelle', 'Deal-Wert', 'Gewichteter Wert',
            'Eingegangene Zahlungen', 'Offene Zahlungen',
            'Erfasste Stunden', 'Stundensatz', 'Zeitkosten',
            'Profit', 'Eff. Stundensatz'
        ]
        for col_idx, header in enumerate(fin_headers, 1):
            cell = ws3.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill

        fin_contacts = contacts_list if contacts_list else Contact.query.order_by(Contact.organisation).all()
        for row_idx, c in enumerate(fin_contacts, 2):
            total_min = sum(e.dauer_minuten for e in c.time_entries if e.timer_start is None)
            total_hours = round(total_min / 60, 2)
            time_cost = round(total_hours * (c.hourly_rate or _get_default_hourly_rate()), 2)
            received = c.received_payments_total
            profit = round(received - time_cost, 2)
            eff_rate = round(received / total_hours, 2) if total_hours > 0 and received > 0 else None

            ws3.cell(row=row_idx, column=1, value=c.display_name)
            ws3.cell(row=row_idx, column=2, value=c.status)
            ws3.cell(row=row_idx, column=3, value=c.quelle)
            ws3.cell(row=row_idx, column=4, value=c.deal_value)
            ws3.cell(row=row_idx, column=5, value=round(c.weighted_value, 2) if c.weighted_value else None)
            ws3.cell(row=row_idx, column=6, value=received if received > 0 else None)
            ws3.cell(row=row_idx, column=7, value=c.expected_payments_total if c.expected_payments_total > 0 else None)
            ws3.cell(row=row_idx, column=8, value=total_hours if total_hours > 0 else None)
            ws3.cell(row=row_idx, column=9, value=c.hourly_rate or _get_default_hourly_rate())
            ws3.cell(row=row_idx, column=10, value=time_cost if time_cost > 0 else None)
            ws3.cell(row=row_idx, column=11, value=profit if received > 0 or time_cost > 0 else None)
            ws3.cell(row=row_idx, column=12, value=eff_rate)

        for col in ws3.columns:
            max_length = 0
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws3.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'crm_export_{date.today().isoformat()}.xlsx'
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=filename)

    countries = db.session.query(Contact.land).filter(
        Contact.land.isnot(None), Contact.land != ''
    ).distinct().order_by(Contact.land).all()
    countries = [c[0] for c in countries]

    return render_template('export.html',
                           stages=PIPELINE_STAGES, countries=countries,
                           quellen_options=QUELLEN_OPTIONS,
                           branchen_options=BRANCHEN_OPTIONS,
                           prioritaet_options=PRIORITAET_OPTIONS)


def _ea_resolve_period(args):
    von_raw = (args.get('von') or '').strip()
    bis_raw = (args.get('bis') or '').strip()
    if von_raw and bis_raw:
        try:
            von = date.fromisoformat(von_raw)
            bis = date.fromisoformat(bis_raw)
        except ValueError:
            return None, None, None
        if von > bis:
            return None, None, None
        return von, bis, f'{von.isoformat()}_{bis.isoformat()}'

    jahr_raw = (args.get('jahr') or '').strip()
    if not jahr_raw:
        return None, None, None
    try:
        jahr = int(jahr_raw)
    except ValueError:
        return None, None, None

    quartal_raw = (args.get('quartal') or '').strip()
    if quartal_raw:
        try:
            q = int(quartal_raw)
        except ValueError:
            return None, None, None
        if q not in (1, 2, 3, 4):
            return None, None, None
        start_month = (q - 1) * 3 + 1
        end_month = start_month + 2
        von = date(jahr, start_month, 1)
        last_day = calendar.monthrange(jahr, end_month)[1]
        bis = date(jahr, end_month, last_day)
        return von, bis, f'Q{q}_{jahr}'

    return date(jahr, 1, 1), date(jahr, 12, 31), str(jahr)


def _ea_income_rows(von, bis):
    docs = Document.query.options(joinedload(Document.groups)).filter(
        Document.doc_type == 'invoice',
        Document.status == 'paid',
        Document.payment_date.isnot(None),
        Document.payment_date >= von,
        Document.payment_date <= bis,
    ).all()
    rows = []
    for d in docs:
        by_rate = {}
        for g in d.groups:
            for item in g.items:
                rate = float(item.vat_rate)
                acc = by_rate.setdefault(rate, {'net': 0.0, 'vat': 0.0})
                acc['net'] += item.line_total
                acc['vat'] += item.line_vat
        rows.append(ea_export.EaIncomeRow(
            number=d.number,
            date=d.payment_date,
            contact=d.contact.display_name if d.contact else '',
            subject=d.subject or '',
            net=d.total_net,
            vat=d.total_vat,
            gross=d.total_gross,
            vat_by_rate=by_rate,
        ))
    return rows


def _ea_expense_rows(von, bis):
    invs = IncomingInvoice.query.filter(
        IncomingInvoice.status == 'paid',
        IncomingInvoice.payment_date.isnot(None),
        IncomingInvoice.payment_date >= von,
        IncomingInvoice.payment_date <= bis,
    ).all()
    rows = []
    for inv in invs:
        cat = inv.category
        vat = float(inv.vat_amount or 0)
        rows.append(ea_export.EaExpenseRow(
            number=inv.number,
            date=inv.payment_date,
            supplier=inv.supplier.display_name if inv.supplier else '',
            subject=inv.subject or '',
            category_code=cat.code if cat else '(ohne)',
            category_name=cat.name if cat else 'Ohne Kategorie',
            ear_group=cat.ear_group if (cat and cat.ear_group) else '(ohne)',
            net=float(inv.net_amount or 0),
            vat=vat,
            gross=float(inv.gross_amount or 0),
            vat_rate=float(inv.vat_rate or 0),
            vat_deductible=vat if inv.is_vat_deductible else 0.0,
        ))
    return rows


@app.route('/export/ea', methods=['GET'])
def export_ea():
    von, bis, label = _ea_resolve_period(request.args)

    if von and bis:
        income_rows = _ea_income_rows(von, bis)
        expense_rows = _ea_expense_rows(von, bis)
        summary = ea_export.compute_ea_summary(income_rows, expense_rows)
        wb = ea_export.build_ea_workbook(income_rows, expense_rows, summary, von, bis)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f'EA-Export_{label}.xlsx'
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=filename,
        )

    today = date.today()
    current_q = (today.month - 1) // 3 + 1
    years = list(range(today.year, today.year - 5, -1))
    return render_template('export_ea.html',
                           years=years,
                           current_year=today.year,
                           current_quarter=current_q,
                           invalid=bool(request.args))



def compute_goal_progress(invoices, goal):
    achieved = sum(
        d.total_net for d in invoices
        if d.status == 'paid' and d.payment_date
        and goal.period_start <= d.payment_date <= goal.period_end
    )
    open_net = sum(
        d.total_net for d in invoices
        if d.status in ('sent', 'overdue') and d.date
        and goal.period_start <= d.date <= goal.period_end
    )
    return {'achieved': round(achieved, 2), 'open_net': round(open_net, 2)}


def compute_cashflow_forecast(open_invoices, active_retainers, vat_rate=0.2):
    open_gross = round(sum(d.total_gross for d in open_invoices), 2)
    recurring_gross = round(sum(r.monthly_fee for r in active_retainers) * (1 + vat_rate), 2)
    return {'open_gross': open_gross, 'recurring_gross': recurring_gross,
            'total': round(open_gross + recurring_gross, 2)}


def compute_revenue_projection(invoices, retainers, won_projects, today, plan_value=None,
                               expected_recurring=(), expected_oneoff=()):
    year = str(today.year)
    ist = 0.0
    for d in invoices:
        if d.status not in ('sent', 'overdue', 'paid'):
            continue
        period = service_period_to_yyyymm(d.service_period) or d.date.strftime('%Y-%m')
        if period[:4] == year:
            ist += d.total_net

    retainer_detail = []
    for r in retainers:
        start_p = r.start_date.strftime('%Y-%m')
        end_p = r.end_date.strftime('%Y-%m') if r.end_date else f'{year}-12'
        months = [p for p in (f'{year}-{m:02d}' for m in range(1, 13))
                  if start_p <= p <= end_p
                  and not (r.last_fee_period and p <= r.last_fee_period)]
        if months and r.monthly_fee:
            kunde = getattr(getattr(r, 'contact', None), 'organisation', None)
            retainer_detail.append({'name': kunde or r.name, 'months': len(months),
                                    'amount': round(len(months) * r.monthly_fee, 2)})

    projekt_detail = []
    for p in won_projects:
        rest = (p.deal_value or 0) - (p.invoiced_net or 0)
        if rest > 0:
            projekt_detail.append({'name': p.name, 'amount': round(rest, 2)})

    erwartet_detail = []
    now_p = today.strftime('%Y-%m')
    for e in expected_recurring:
        step = e.interval_months or 1
        start = e.expected_date
        months = [p for i, p in enumerate(month_range(start, date(today.year, 12, 1)))
                  if i % step == 0 and p[:4] == year and p >= now_p]
        if months and e.amount:
            erwartet_detail.append({'name': e.label, 'months': len(months),
                                    'amount': round(len(months) * e.amount, 2)})
    for e in expected_oneoff:
        if e.amount and e.amount > 0 and str(e.expected_date.year) == year:
            erwartet_detail.append({'name': e.label, 'months': None,
                                    'amount': round(e.amount, 2)})

    retainer_rest = round(sum(x['amount'] for x in retainer_detail), 2)
    projekte_rest = round(sum(x['amount'] for x in projekt_detail), 2)
    erwartet_rest = round(sum(x['amount'] for x in erwartet_detail), 2)
    total = round(ist + retainer_rest + projekte_rest + erwartet_rest, 2)
    return {
        'year': today.year,
        'ist': round(ist, 2),
        'retainer_rest': retainer_rest,
        'projekte_rest': projekte_rest,
        'erwartet_rest': erwartet_rest,
        'total': total,
        'plan': plan_value,
        'delta': round(total - plan_value, 2) if plan_value else None,
        'retainer_detail': retainer_detail,
        'projekt_detail': sorted(projekt_detail, key=lambda x: -x['amount']),
        'erwartet_detail': erwartet_detail,
    }


def compute_quote_conversion(quotes, invoices):
    quote_ids = {q.id for q in quotes}
    total = len(quotes)
    won = [q for q in quotes if q.status in QUOTE_WON_STATUSES]
    rejected = sum(1 for q in quotes if q.status == 'rejected')
    converted = sum(1 for q in quotes if q.status == 'converted')
    from_quote = [i for i in invoices if i.source_document_id in quote_ids]

    avg_quote_net = round(sum(q.total_net for q in won) / len(won), 2) if won else 0
    avg_invoice_net = (round(sum(i.total_net for i in from_quote) / len(from_quote), 2)
                       if from_quote else 0)
    return {
        'total': total,
        'accepted': len(won),
        'rejected': rejected,
        'converted': converted,
        'accept_rate': round(len(won) / total * 100) if total > 0 else 0,
        'convert_rate': round(converted / total * 100) if total > 0 else 0,
        'avg_quote_net': avg_quote_net,
        'avg_invoice_net': avg_invoice_net,
        'value_retention': round(avg_invoice_net / avg_quote_net * 100) if avg_quote_net > 0 else 0,
    }


def compute_pipeline(quotes, today):
    live = [q for q in quotes
            if q.status == 'sent' and (q.valid_until is None or q.valid_until >= today)]
    return {'count': len(live),
            'value_net': round(sum(q.total_net for q in live), 2),
            'quotes': live}


def compute_capacity_projection(retainers, today, months=12,
                                weekly_target_mandate=20.0, weeks_per_month=4.33):
    soll_h = round(weekly_target_mandate * weeks_per_month, 1)
    base = today.year * 12 + (today.month - 1)
    out = []
    for i in range(months):
        y, m = divmod(base + i, 12)
        m += 1
        month_first = date(y, m, 1)
        month_last = date(y, m, calendar.monthrange(y, m)[1])
        items = []
        planned = 0.0
        for r in retainers:
            hrs = r.monthly_hours or 0
            if hrs <= 0:
                continue
            if r.start_date and r.start_date > month_last:
                continue
            if r.end_date and r.end_date < month_first:
                continue
            planned += hrs
            items.append({'name': r.name, 'hours': round(hrs, 1)})
        planned = round(planned, 1)
        out.append({
            'label': month_first.strftime('%b %y'),
            'label_short': month_first.strftime('%b'),
            'ym': month_first.strftime('%Y-%m'),
            'planned_h': planned,
            'free_h': round(soll_h - planned, 1),
            'pct': round(planned / soll_h * 100, 1) if soll_h > 0 else 0.0,
            'over': planned > soll_h,
            'items': sorted(items, key=lambda x: -x['hours']),
        })

    current_free = out[0]['free_h'] if out else soll_h
    min_free = min((mo['free_h'] for mo in out), default=soll_h)
    first_relief = None
    for mo in out[1:]:
        if mo['free_h'] - current_free >= 4:
            first_relief = mo['label']
            break

    return {
        'soll_h': soll_h,
        'weekly_target': round(weekly_target_mandate, 1),
        'months': out,
        'current_free_h': current_free,
        'min_free_h': round(min_free, 1),
        'first_relief_month': first_relief,
    }


GRAU_BILLING_MODELS = {'lead_quote', 'agentur_partner', 'aufwand_extern',
                       'lost', 'paused'}
DEAD_STATUSES = {'lost', 'closed_no_response', 'bounced'}

PROFIT_PARTNER_MODELS = {'agentur_partner', 'aufwand_extern'}
PROFIT_PAUSCHAL_MODELS = {'projekt_fix', 'retainer'}

PROFIT_AKQUISE_MODELS = {'lost'}
PROFIT_PRO_BONO_MODELS = {'pro_bono'}
PROFIT_KOSTEN_SEGMENTE = ('akquise', 'pro_bono')


TREND_BILLING_MODELS = {'retainer', 'retainer_trueup'}

ABRECHNUNGSMODELL_GUARD_SETS = {
    'GRAU_BILLING_MODELS': GRAU_BILLING_MODELS,
    'PROFIT_PARTNER_MODELS': PROFIT_PARTNER_MODELS,
    'PROFIT_PAUSCHAL_MODELS': PROFIT_PAUSCHAL_MODELS,
    'PROFIT_AKQUISE_MODELS': PROFIT_AKQUISE_MODELS,
    'PROFIT_PRO_BONO_MODELS': PROFIT_PRO_BONO_MODELS,
    'TREND_BILLING_MODELS': TREND_BILLING_MODELS,
}

ABRECHNUNGSMODELL_KLASSIFIKATOREN = {
    'ampel (_classify_by_model)': {
        'explizit': GRAU_BILLING_MODELS | {'pro_bono', 'intern'},
        'else_bewusst': {
            'retainer': "regulaer abrechenbar. Ohne aktive Retainer-Zeile ist eine "
                        "Pauschale mit offenen Stunden ein echter Handlungsbedarf.",
            'retainer_trueup': "regulaer abrechenbar; die Stunden gehen in die True-Up.",
            'stundenbasis': "regulaer abrechenbar, das ist der Normalfall des Modells.",
            'projekt_fix': "regulaer abrechenbar, solange keine Retainer-Zeile greift.",
        },
    },
    'profitabilitaet (compute_profitability)': {
        'explizit': (PROFIT_AKQUISE_MODELS | PROFIT_PRO_BONO_MODELS
                     | PROFIT_PARTNER_MODELS | PROFIT_PAUSCHAL_MODELS),
        'else_bewusst': {
            'retainer_trueup': "regulaer, weil die Stunden ueber die True-Up-Rechnung "
                               "real abgerechnet werden (Satz ist aussagekraeftig).",
            'stundenbasis': "regulaer; noch nicht fakturierte Stunden sind ein "
                            "Timing-Effekt, der Umsatz kommt (Beispiel).",
            'lead_quote': "regulaer; loest sich selbst auf — gewonnen bringt Umsatz, "
                          "verloren wird zu 'lost' und damit zu Akquise.",
            'paused': "regulaer; ein pausiertes Mandat WAR ein Mandat, Umsatz und "
                      "Stunden waren real (Beispiel).",
            'intern': "erreicht die Funktion nicht — die Route filtert "
                      "abrechnungsmodell == 'intern' vorher heraus.",
        },
    },
}


def _classify_by_model(mode, delivered_h):
    if mode == 'pro_bono':
        return 'gruen', None
    elif mode in GRAU_BILLING_MODELS:
        return 'grau', None
    elif mode == 'intern':
        return 'aus', None
    elif delivered_h > 0:
        return 'rot', 'time'
    else:
        return 'gruen', None


def _apply_billing_deferral(status, contact, today):
    bis = getattr(contact, 'abrechnung_terminiert_bis', None)
    status['terminiert_bis'] = bis
    if bis and today <= bis:
        status['ampel'] = 'terminiert'
        status['action'] = None
    return status


def compute_billing_status(contact, retainer, entries, today, default_rate=DEFAULT_HOURLY_RATE):
    mode = (contact.abrechnungsmodell or '').strip() or None

    if retainer is not None:
        period_start = retainer.last_billed_date or retainer.start_date
        months_since = ((today.year - period_start.year) * 12
                        + today.month - period_start.month)
        elapsed_months = months_since + 1
        period_min = sum(e.dauer_minuten for e in entries if e.datum >= period_start)
        delivered_h = round(period_min / 60, 2)
        budget_h = round(retainer.monthly_hours * elapsed_months, 2) if retainer.monthly_hours else None
        rate = retainer.effective_hourly_rate
        value_eur = round(delivered_h * rate, 2) if rate else None
        period_due = months_since >= retainer.billing_interval_months
        over_budget = budget_h is not None and delivered_h > budget_h
        ampel = 'gelb' if (over_budget or period_due) else 'gruen'
        action = ('trueup' if mode == 'retainer_trueup' else 'retainer') if period_due else None
        return _apply_billing_deferral({
            'mode': mode, 'delivered_h': delivered_h, 'budget_h': budget_h,
            'value_eur': value_eur, 'ampel': ampel,
            'saldo_h': round(budget_h - delivered_h, 2) if budget_h is not None else None,
            'action': action, 'period_start': period_start, 'period_due': period_due,
        }, contact, today)

    delivered_h = round(sum(e.dauer_minuten for e in entries) / 60, 2)
    rate = contact.hourly_rate or default_rate
    value_eur = round(delivered_h * rate, 2) if rate else None
    ampel, action = _classify_by_model(mode, delivered_h)
    return _apply_billing_deferral({
        'mode': mode, 'delivered_h': delivered_h, 'budget_h': None,
        'value_eur': value_eur, 'ampel': ampel, 'saldo_h': None,
        'action': action, 'period_start': None, 'period_due': False,
    }, contact, today)


def _months_ago(d, months):
    total = (d.year * 12 + (d.month - 1)) - months
    y, m = divmod(total, 12)
    m += 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def compute_profitability(contact, invoices, entries, window_start=None,
                          default_rate=DEFAULT_HOURLY_RATE):
    def _in_window(d):
        return d is not None and (window_start is None or d >= window_start)

    revenue = round(sum(
        inv.total_net for inv in invoices
        if inv.status == 'paid' and _in_window(inv.payment_date)
    ), 2)
    hours = round(sum(
        e.dauer_minuten for e in entries if _in_window(e.datum)
    ) / 60, 2)

    effective_rate = round(revenue / hours, 2) if hours > 0 else None
    nominal_rate = contact.hourly_rate or default_rate
    delta = round(effective_rate - nominal_rate, 2) if effective_rate is not None else None

    if hours > 0 and revenue > 0:
        kind = 'voll'
    elif hours > 0:
        kind = 'nur_zeit'
    elif revenue > 0:
        kind = 'nur_umsatz'
    else:
        kind = 'leer'

    mode = (contact.abrechnungsmodell or '').strip() or None

    war_nie_mandat = not any(inv.status == 'paid' for inv in invoices)

    if mode in PROFIT_AKQUISE_MODELS and war_nie_mandat:
        segment = 'akquise'
    elif mode in PROFIT_PRO_BONO_MODELS:
        segment = 'pro_bono'
    elif mode in PROFIT_PARTNER_MODELS:
        segment = 'partner'
    elif mode in PROFIT_PAUSCHAL_MODELS:
        segment = 'pauschale'
    else:
        segment = 'regulaer'

    notional_cost = None
    if segment in PROFIT_KOSTEN_SEGMENTE:
        effective_rate = None
        delta = None
        notional_cost = round(hours * nominal_rate, 2) if nominal_rate else None

    return {
        'mode': mode, 'revenue': revenue, 'hours': hours,
        'effective_rate': effective_rate, 'nominal_rate': nominal_rate,
        'delta': delta, 'kind': kind, 'segment': segment,
        'self_tracking': segment == 'pauschale',
        'notional_cost': notional_cost,
    }


def compute_partneragentur_rollup(parent, children, invoices_by_cid, entries_by_cid,
                             window_start=None, default_rate=DEFAULT_HOURLY_RATE):
    ids = [parent.id] + [c.id for c in children]
    agg_invoices = [inv for cid in ids for inv in invoices_by_cid.get(cid, [])]
    agg_entries = [e for cid in ids for e in entries_by_cid.get(cid, [])]
    agg = compute_profitability(parent, agg_invoices, agg_entries,
                                window_start=window_start, default_rate=default_rate)
    agg['segment'] = 'partner'
    agg['self_tracking'] = False

    contributors = []
    for c in children:
        ch = compute_profitability(c, invoices_by_cid.get(c.id, []), entries_by_cid.get(c.id, []),
                                   window_start=window_start, default_rate=default_rate)
        if ch['hours'] > 0 or ch['revenue'] > 0:
            contributors.append({'name': c.display_name, 'hours': ch['hours'],
                                 'revenue': ch['revenue'], 'status': c.status})
    contributors.sort(key=lambda x: -x['hours'])
    agg['contributors'] = contributors
    return agg


def compute_utilization_trend(retainer, entries, months):
    monthly_h = (retainer.monthly_hours or 0) if retainer is not None else 0
    budget_per_month = monthly_h if monthly_h > 0 else None

    start_period = retainer.start_date.strftime('%Y-%m') if retainer is not None else None
    end_date = getattr(retainer, 'end_date', None) if retainer is not None else None
    end_period = end_date.strftime('%Y-%m') if end_date else None

    minutes_by_period = defaultdict(int)
    for e in entries:
        if e.datum is None:
            continue
        minutes_by_period[e.datum.strftime('%Y-%m')] += e.dauer_minuten or 0

    rows = []
    for period in months:
        hours = round(minutes_by_period.get(period, 0) / 60, 2)
        under_contract = (start_period is not None and period >= start_period
                          and (end_period is None or period <= end_period))
        budget = budget_per_month if under_contract else None
        rows.append({
            'period': period,
            'label': period_label_de(period),
            'hours': hours,
            'budget_h': budget,
            'saldo_h': round(budget - hours, 2) if budget is not None else None,
            'under_contract': under_contract,
        })

    covered = [m for m in rows if m['budget_h'] is not None]
    covered_periods = {m['period'] for m in covered}
    outside_periods = {m['period'] for m in rows if not m['under_contract']}
    total_min = sum(minutes_by_period.get(p, 0) for p in months)
    covered_min = sum(minutes_by_period.get(p, 0) for p in covered_periods)
    outside_min = sum(minutes_by_period.get(p, 0) for p in outside_periods)
    hours_total = round(total_min / 60, 2)
    hours_covered = round(covered_min / 60, 2)
    budget_total = round(sum(m['budget_h'] for m in covered), 2) if covered else None
    peak = max((max(m['hours'], m['budget_h'] or 0) for m in rows), default=0)
    busiest = max(rows, key=lambda m: m['hours'])['period'] if hours_total > 0 else None

    return {
        'months': rows,
        'monthly_hours': budget_per_month,
        'hours_total': hours_total,
        'hours_covered': hours_covered,
        'hours_outside': round(outside_min / 60, 2),
        'budget_total': budget_total,
        'saldo_total': (round(budget_total - hours_covered, 2)
                        if budget_total is not None else None),
        'covered_months': len(covered),
        'peak_h': round(peak, 2),
        'busiest': busiest,
    }


def previous_period(today):
    first_of_month = today.replace(day=1)
    last_of_prev = first_of_month - timedelta(days=1)
    return last_of_prev.strftime('%Y-%m')


def plan_monthly_fee_draft(retainer, target_period, existing_periods):
    if not retainer.active:
        return {'action': 'skip', 'reason': 'inaktiv', 'amount': 0.0}
    start_period = retainer.start_date.strftime('%Y-%m')
    if target_period < start_period:
        return {'action': 'skip', 'reason': 'im Zielmonat noch nicht aktiv', 'amount': 0.0}
    if retainer.last_fee_period and retainer.last_fee_period >= target_period:
        return {'action': 'skip', 'reason': 'bereits fakturiert', 'amount': 0.0}
    if target_period in existing_periods:
        return {'action': 'skip', 'reason': 'Rechnung mit diesem Zeitraum existiert', 'amount': 0.0}
    return {'action': 'create', 'reason': 'faellig', 'amount': retainer.monthly_fee}


def month_range(start, end):
    out = []
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        out.append(f'{y:04d}-{m:02d}')
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def plan_trueup_fee_months(retainer, period_start, period_end, existing_periods):
    start_period = retainer.start_date.strftime('%Y-%m')
    last_complete = previous_period(period_end)
    covered = set(existing_periods or ())
    return [p for p in month_range(period_start, period_end)
            if start_period <= p <= last_complete
            and not (retainer.last_fee_period and p <= retainer.last_fee_period)
            and p not in covered]



OPEN_INVOICE_STATUSES = ('sent', 'overdue')

BANK_AMOUNT_TOLERANCE = 0.01

BANK_MATCH_WINDOW_DAYS = 120

_ORG_SUFFIX_TOKENS = {
    'gmbh', 'mbh', 'co', 'kg', 'og', 'ag', 'ohg', 'gbr', 'se', 'ug', 'ev',
    'eu', 'ku', 'elp', 'kgaa', 'gesmbh', 'ges',
}

_DOC_NUMBER_RE = re.compile(r'\b([RH])\s*-?\s*0*(\d{1,5})\b', re.IGNORECASE)

_MIN_EXTERNAL_NUMBER_LEN = 4


def _extract_doc_numbers(reference):
    if not reference:
        return []
    found = []
    for prefix, digits in _DOC_NUMBER_RE.findall(reference):
        number = f'{prefix.upper()}-{int(digits):05d}'
        if number not in found:
            found.append(number)
    return found


def _normalize_org_name(name):
    if not name:
        return ''
    cleaned = re.sub(r'[^\w\s]', ' ', name.lower().replace('.', ''))
    tokens = cleaned.split()
    while tokens and tokens[-1] in _ORG_SUFFIX_TOKENS:
        tokens.pop()
    return ' '.join(tokens)


def _partner_matches(tx, contact):
    if contact is None:
        return False
    partner = _normalize_org_name(getattr(tx, 'partner_name', None))
    org = _normalize_org_name(getattr(contact, 'organisation', None))
    if not partner or not org:
        return False
    return partner == org or partner in org or org in partner


def _in_match_window(booking_date, beleg_date, window_days):
    if booking_date is None or beleg_date is None:
        return False
    if booking_date < beleg_date:
        return False
    return (booking_date - beleg_date).days <= window_days


def classify_match(candidates):
    if not candidates:
        return 'kein_match'
    if len(candidates) == 1:
        signals = candidates[0].get('signals', ())
        if all(s in signals for s in ('number_exact', 'amount_exact', 'in_window')):
            return 'sicher'
    return 'unsicher'


def _hints_for_known_numbers(numbers, invoices_by_number):
    hints = []
    for number in numbers:
        inv = invoices_by_number.get(number)
        if inv is None:
            hints.append(f'Nummer {number} im Verwendungszweck, aber kein Beleg mit dieser Nummer.')
        elif inv.status == 'paid':
            pay = getattr(inv, 'payment_date', None)
            when = f' (Zahldatum {pay.strftime("%d.%m.%Y")})' if pay else ''
            hints.append(f'{number} ist bereits als bezahlt gebucht{when}.')
        elif inv.status == 'cancelled':
            hints.append(f'{number} ist storniert.')
        elif inv.status == 'draft':
            hints.append(f'{number} ist noch ein Entwurf und nicht versendet.')
    return hints


def _hint_for_settled_amount(abs_amount, invoices, limit=3):
    matches = [inv for inv in invoices
               if inv.status in ('paid', 'cancelled')
               and abs(abs_amount - round(float(inv.total_gross or 0), 2))
               <= BANK_AMOUNT_TOLERANCE]
    if not matches:
        return None
    shown = ', '.join(f'{inv.number} ({inv.status_label if hasattr(inv, "status_label") else inv.status})'
                      for inv in matches[:limit])
    rest = f' und {len(matches) - limit} weitere' if len(matches) > limit else ''
    return (f'Kein offener Beleg über {abs_amount:.2f} €. Derselbe Betrag steht auf '
            f'{shown}{rest} — vermutlich eine schon gebuchte Zahlung.')


def match_bank_transaction(tx, invoices, incoming_invoices, contacts_by_id=None, *,
                           window_days=BANK_MATCH_WINDOW_DAYS):
    contacts_by_id = contacts_by_id or {}
    amount = float(tx.amount or 0)
    direction = 'in' if amount >= 0 else 'out'
    abs_amount = round(abs(amount), 2)
    numbers = _extract_doc_numbers(getattr(tx, 'reference', None))
    hints = []
    candidates = []

    def _result():
        return {'direction': direction, 'confidence': classify_match(candidates),
                'candidates': candidates, 'hints': hints}

    foreign_currency = (getattr(tx, 'foreign_currency', None) or '').upper()
    if foreign_currency and foreign_currency != 'EUR':
        foreign_amount = getattr(tx, 'foreign_amount', None)
        betrag_teil = f'{float(foreign_amount):.2f} {foreign_currency}' if foreign_amount \
            else foreign_currency
        hints.append(f'Aus {betrag_teil} umgerechnet: der Betrag kann um '
                     f'Kursdifferenzen von einer Rechnung abweichen.')

    if abs_amount == 0:
        return _result()

    if direction == 'in':
        invoices_by_number = {}
        for inv in invoices:
            invoices_by_number.setdefault(inv.number, inv)
        hints.extend(_hints_for_known_numbers(numbers, invoices_by_number))
        settled_hint = _hint_for_settled_amount(abs_amount, invoices)

        named_existing = {n for n in numbers if n in invoices_by_number}

        for inv in invoices:
            if named_existing and inv.number not in named_existing:
                continue
            if inv.status not in OPEN_INVOICE_STATUSES:
                continue
            if not _in_match_window(tx.booking_date, inv.date, window_days):
                continue
            gross = round(float(inv.total_gross or 0), 2)
            signals = ['in_window']
            if inv.number in numbers:
                signals.append('number_exact')
            if abs(abs_amount - gross) <= BANK_AMOUNT_TOLERANCE:
                signals.append('amount_exact')
            elif abs_amount < gross:
                signals.append('amount_partial')
            if _partner_matches(tx, contacts_by_id.get(inv.contact_id)):
                signals.append('partner_match')

            if 'number_exact' not in signals and 'amount_exact' not in signals:
                continue

            if 'amount_partial' in signals:
                hints.append(f'Teilzahlung? Betrag {abs_amount:.2f} liegt unter '
                             f'Rechnung {gross:.2f} ({inv.number}).')
            candidates.append({'kind': 'invoice', 'id': inv.id, 'number': inv.number,
                               'signals': signals,
                               'amount_delta': round(abs_amount - gross, 2)})

        if not candidates and settled_hint and not named_existing:
            hints.append(settled_hint)
        return _result()

    for inc in incoming_invoices:
        if inc.status == 'paid':
            continue
        if not _in_match_window(tx.booking_date, inc.invoice_date, window_days):
            continue
        gross = round(float(inc.gross_amount or 0), 2)
        signals = ['in_window']
        external = (inc.external_invoice_number or '').strip().lower()
        reference = (getattr(tx, 'reference', None) or '').lower()
        if len(external) >= _MIN_EXTERNAL_NUMBER_LEN and external in reference:
            signals.append('number_exact')
        if abs(abs_amount - gross) <= BANK_AMOUNT_TOLERANCE:
            signals.append('amount_exact')
        elif abs_amount < gross:
            signals.append('amount_partial')
        if _partner_matches(tx, contacts_by_id.get(inc.supplier_id)):
            signals.append('partner_match')

        if 'number_exact' not in signals and 'amount_exact' not in signals:
            continue
        candidates.append({'kind': 'incoming', 'id': inc.id, 'number': inc.number,
                           'signals': signals,
                           'amount_delta': round(abs_amount - gross, 2)})
    return _result()


MONTHS_DE = ['', 'Januar', 'Februar', 'März', 'April', 'Mai', 'Juni',
             'Juli', 'August', 'September', 'Oktober', 'November', 'Dezember']


def period_label_de(target_period):
    y, m = target_period.split('-')
    return f'{MONTHS_DE[int(m)]} {y}'


def period_service_string(target_period):
    y, m = int(target_period[:4]), int(target_period[5:7])
    last = calendar.monthrange(y, m)[1]
    return f'01.{m:02d}.{y} – {last:02d}.{m:02d}.{y}'


def service_period_to_yyyymm(s):
    if not s:
        return None
    match = re.match(r'\s*\d{2}\.(\d{2})\.(\d{4})', s)
    if not match:
        return None
    mm, yyyy = match.group(1), match.group(2)
    return f'{yyyy}-{mm}'


@app.route('/goals', methods=['GET', 'POST'])
def sales_goals():
    if request.method == 'POST':
        period_type = request.form.get('period_type', 'monthly')
        target_amount = parse_form_float(request.form.get('target_amount', ''))
        period_start_str = request.form.get('period_start', '')

        if not target_amount or not period_start_str:
            flash('Ziel-Betrag und Startdatum sind erforderlich', 'error')
        else:
            period_start = parse_date(period_start_str)
            if period_start:
                period_start = period_start.replace(day=1)
                if period_type == 'quarterly':
                    quarter_month = ((period_start.month - 1) // 3) * 3 + 1
                    period_start = period_start.replace(month=quarter_month)

                goal = SalesGoal(
                    period_type=period_type,
                    period_start=period_start,
                    target_amount=target_amount
                )
                db.session.add(goal)
                db.session.commit()
                flash(f'Ziel "{goal.label}: {target_amount:,.0f} EUR" erstellt', 'success')
                return redirect(url_for('sales_goals'))

    goals = SalesGoal.query.order_by(SalesGoal.period_start.desc()).all()

    all_invoices = Document.query.filter_by(doc_type='invoice').all()
    goal_data = []
    for goal in goals:
        p = compute_goal_progress(all_invoices, goal)
        achieved = p['achieved']
        open_net = p['open_net']
        pct = round(achieved / goal.target_amount * 100, 1) if goal.target_amount > 0 else 0
        paid_w = min(pct, 100)
        open_pct = round(open_net / goal.target_amount * 100, 1) if goal.target_amount > 0 else 0
        open_w = min(open_pct, 100 - paid_w)
        goal_data.append({
            'goal': goal,
            'achieved': achieved,
            'open_net': open_net,
            'pct': pct,
            'paid_w': paid_w,
            'open_w': open_w,
        })

    return render_template('goals.html', goal_data=goal_data,
                           period_types=GOAL_PERIOD_TYPES)


@app.route('/goals/<int:id>/delete', methods=['POST'])
def sales_goal_delete(id):
    goal = SalesGoal.query.get_or_404(id)
    label = goal.label
    db.session.delete(goal)
    db.session.commit()
    flash(f'Ziel "{label}" gelöscht', 'success')
    return redirect(url_for('sales_goals'))



@app.route('/settings/followup', methods=['GET', 'POST'])
def followup_settings():
    if request.method == 'POST':
        for stage_key, stage_name in PIPELINE_STAGES:
            if stage_key in ('won', 'lost', 'on_hold'):
                continue
            days = parse_form_int(request.form.get(f'days_{stage_key}', ''))
            enabled = request.form.get(f'enabled_{stage_key}') == 'on'

            rule = FollowUpRule.query.filter_by(stage=stage_key).first()
            if rule:
                rule.days_threshold = days or FOLLOWUP_DEFAULTS.get(stage_key, 7)
                rule.enabled = enabled
            else:
                rule = FollowUpRule(
                    stage=stage_key,
                    days_threshold=days or FOLLOWUP_DEFAULTS.get(stage_key, 7),
                    enabled=enabled
                )
                db.session.add(rule)

        db.session.commit()
        flash('Follow-Up Regeln gespeichert', 'success')
        return redirect(url_for('followup_settings'))

    rules = {}
    for stage_key, stage_name in PIPELINE_STAGES:
        if stage_key in ('won', 'lost', 'on_hold'):
            continue
        rule = FollowUpRule.query.filter_by(stage=stage_key).first()
        if rule:
            rules[stage_key] = {'days': rule.days_threshold, 'enabled': rule.enabled}
        else:
            rules[stage_key] = {'days': FOLLOWUP_DEFAULTS.get(stage_key, 7), 'enabled': False}

    active_stages = [(k, v) for k, v in PIPELINE_STAGES if k not in ('won', 'lost', 'on_hold')]
    sequences = NumberSequence.query.order_by(NumberSequence.prefix).all()
    company_settings = CompanySettings.get_all()
    return render_template('settings.html', rules=rules, stages=active_stages,
                           sequences=sequences, company_settings=company_settings)


@app.route('/settings/company', methods=['POST'])
def company_settings_save():
    for key in CompanySettings.DEFAULTS:
        value = request.form.get(key, '').strip()
        if value:
            row = CompanySettings.query.get(key)
            if row:
                row.value = value
            else:
                db.session.add(CompanySettings(key=key, value=value))
    db.session.commit()
    flash('Firmendaten gespeichert.', 'success')
    return redirect(url_for('followup_settings'))


@app.route('/settings/documents', methods=['POST'])
def document_settings():
    sequences = NumberSequence.query.all()
    for seq in sequences:
        new_next = request.form.get(f'next_{seq.prefix}', type=int)
        if new_next and new_next >= seq.next_number:
            seq.next_number = new_next
    db.session.commit()
    flash('Nummernkreise aktualisiert.', 'success')
    return redirect(url_for('followup_settings'))



@app.route('/import/preview', methods=['POST'])
def import_preview():
    filepath = request.form.get('filepath', '').strip()
    if not filepath:
        return jsonify({'error': 'Kein Dateipfad'})

    import os
    if not os.path.exists(filepath):
        return jsonify({'error': f'Datei nicht gefunden: {filepath}'})

    try:
        from openpyxl import load_workbook
        from database import COLUMN_MAPPING, _build_lookup_dicts, _parse_row_data, _find_existing_contact
        wb = load_workbook(filepath, read_only=True)
        ws = wb.active

        headers = [str(cell.value).strip().lower() if cell.value else '' for cell in ws[1]]
        field_indices = {}
        for field, possible_names in COLUMN_MAPPING.items():
            for i, header in enumerate(headers):
                if header in possible_names:
                    field_indices[field] = i
                    break

        email_lookup, name_lookup, org_lookup = _build_lookup_dicts()
        preview = {'new': [], 'update': [], 'skip': 0, 'total': 0}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(cell is None for cell in row):
                continue
            preview['total'] += 1
            data = _parse_row_data(row, field_indices)
            if not data:
                preview['skip'] += 1
                continue
            existing = _find_existing_contact(data, email_lookup, name_lookup, org_lookup)
            name = data.get('organisation') or f"{data.get('vorname', '')} {data.get('nachname', '')}".strip() or '?'
            email = data.get('mail_person') or data.get('mail_organisation') or ''
            if existing:
                preview['update'].append({'name': name, 'email': email, 'existing_name': existing.display_name, 'existing_id': existing.id})
            else:
                preview['new'].append({'name': name, 'email': email})

        wb.close()
        return jsonify(preview)
    except Exception as e:
        return jsonify({'error': str(e)})


@app.route('/import', methods=['GET', 'POST'])
def import_data():
    if request.method == 'POST':
        filepath = request.form.get('filepath', '').strip()
        if not filepath:
            flash('Bitte Dateipfad eingeben', 'error')
            return render_template('import.html')

        result = import_from_excel(filepath)
        if result['error']:
            flash(f'Import Fehler: {result["error"]}', 'error')
        else:
            parts = []
            if result['created']:
                parts.append(f'{result["created"]} erstellt')
            if result['updated']:
                parts.append(f'{result["updated"]} aktualisiert')
            if result['skipped']:
                parts.append(f'{result["skipped"]} übersprungen')
            flash(f'Import abgeschlossen: {", ".join(parts) or "keine Änderungen"}', 'success')
            return redirect(url_for('contacts'))

    default_path = ''
    return render_template('import.html', default_path=default_path)



def _update_document_statuses():
    today = date.today()
    changed = 0
    expired_quotes = Document.query.filter(
        Document.doc_type == 'quote', Document.status == 'sent',
        Document.valid_until < today, Document.valid_until.isnot(None)
    ).all()
    for q in expired_quotes:
        q.status = 'expired'
        changed += 1
    overdue_invoices = Document.query.filter(
        Document.doc_type == 'invoice', Document.status == 'sent',
        Document.due_date < today, Document.due_date.isnot(None)
    ).all()
    for inv in overdue_invoices:
        inv.status = 'overdue'
        changed += 1
    if changed:
        db.session.commit()
    return changed


def _generate_quote_followup_tasks():
    today = date.today()
    threshold = today - timedelta(days=7)
    stale_quotes = Document.query.filter(
        Document.doc_type == 'quote',
        Document.status == 'sent',
        Document.date <= threshold,
    ).all()

    created = 0
    for q in stale_quotes:
        if q.contact and q.contact.status == 'on_hold':
            continue

        if q.contact and q.contact.erinnerung and q.contact.erinnerung >= today:
            continue

        existing_task = Task.query.filter(
            Task.title.like(f'%{q.number}%'),
        ).first()
        if existing_task:
            continue

        task = Task(
            title=f'Nachhaken: {q.number} bei {q.contact.display_name}',
            description=f'Angebot "{q.subject}" vom {q.date.strftime("%d.%m.%Y")} ist seit {(today - q.date).days} Tagen unbeantwortet.',
            status='todo',
            priority='hoch',
            contact_id=q.contact_id,
            created_by='System',
            due_date=today,
        )
        db.session.add(task)
        created += 1

    if created:
        db.session.commit()
    return created


def _create_stammdaten_check_task(contact, doc):
    try:
        existing = Task.query.filter(
            Task.contact_id == contact.id,
            Task.title.like('Stammdaten prüfen:%'),
            Task.status != 'erledigt',
        ).first()
        if existing:
            return
        task = Task(
            title=f'Stammdaten prüfen: {contact.display_name}',
            description=(
                f'Angebot {doc.number} erstellt — Kontakt-Stammdaten gegen Impressum/Firmenbuch prüfen: '
                f'Firmenwortlaut (organisation), organisationsform, firmenbuchnummer, gruendungsdatum, uid, Sitz. '
                f'Danach contact.stammdaten_verifiziert setzen, damit kein weiterer Task entsteht.'
            ),
            status='todo',
            priority='mittel',
            contact_id=contact.id,
            created_by='System',
            due_date=date.today() + timedelta(days=3),
        )
        db.session.add(task)
        db.session.commit()
    except Exception:
        db.session.rollback()


@app.route('/documents')
def documents():
    _update_document_statuses()
    tab = request.args.get('tab', 'quotes')
    status_filter = request.args.get('status', 'all')

    if tab == 'invoices':
        query = Document.query.options(joinedload(Document.contact)).filter_by(doc_type='invoice')
    else:
        query = Document.query.options(joinedload(Document.contact)).filter_by(doc_type='quote')

    if status_filter == 'open':
        if tab == 'invoices':
            query = query.filter(Document.status.in_(['sent', 'overdue']))
        else:
            query = query.filter(Document.status.in_(['draft', 'sent']))
    elif status_filter == 'done':
        if tab == 'invoices':
            query = query.filter_by(status='paid')
        else:
            query = query.filter(Document.status.in_(QUOTE_WON_STATUSES))
    elif status_filter == 'overdue':
        query = query.filter_by(status='overdue')

    contact_filter = request.args.get('contact_id', type=int)
    if contact_filter:
        query = query.filter_by(contact_id=contact_filter)

    search = request.args.get('q', '').strip()
    if search:
        query = query.filter(db.or_(
            Document.number.ilike(f'%{search}%'),
            Document.subject.ilike(f'%{search}%'),
        ))

    sort = request.args.get('sort', 'date_desc')
    if sort == 'date_asc':
        query = query.order_by(Document.date.asc())
    elif sort == 'amount_desc':
        query = query.order_by(Document.id.desc())
    else:
        query = query.order_by(Document.date.desc())

    docs = query.all()

    if sort == 'amount_desc':
        docs.sort(key=lambda d: d.total_net, reverse=True)
    elif sort == 'amount_asc':
        docs.sort(key=lambda d: d.total_net)

    page = request.args.get('page', 1, type=int)
    per_page = 50
    total_docs = len(docs)
    total_pages = max(1, (total_docs + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    docs_paginated = docs[(page - 1) * per_page:page * per_page]

    finance_summary = None
    if tab == 'invoices':
        inv_query = Document.query.filter_by(doc_type='invoice')
        if contact_filter:
            inv_query = inv_query.filter_by(contact_id=contact_filter)
        all_invoices = inv_query.all()
        this_year = date.today().year
        this_month = date.today().month
        paid_invoices = [d for d in all_invoices if d.status == 'paid' and d.payment_date]
        finance_summary = {
            'open_net': sum(d.total_net for d in all_invoices if d.status in ('sent', 'overdue')),
            'overdue_net': sum(d.total_net for d in all_invoices if d.status == 'overdue'),
            'paid_this_month': sum(d.total_net for d in paid_invoices
                                   if d.payment_date.year == this_year
                                   and d.payment_date.month == this_month),
            'paid_this_year': sum(d.total_net for d in paid_invoices
                                  if d.payment_date.year == this_year),
            'open_count': sum(1 for d in all_invoices if d.status in ('sent', 'overdue')),
            'overdue_count': sum(1 for d in all_invoices if d.status == 'overdue'),
            'monthly': [{
                'month': m,
                'total': sum(d.total_net for d in paid_invoices
                            if d.payment_date.year == this_year and d.payment_date.month == m)
            } for m in range(1, this_month + 1)],
        }

    quote_summary = None
    if tab == 'quotes':
        qt_query = Document.query.filter_by(doc_type='quote')
        if contact_filter:
            qt_query = qt_query.filter_by(contact_id=contact_filter)
        all_quotes = qt_query.all()
        total = len(all_quotes)
        accepted = sum(1 for q in all_quotes if q.status in QUOTE_WON_STATUSES)
        rejected = sum(1 for q in all_quotes if q.status in ('rejected', 'expired'))
        decided = accepted + rejected
        this_year = date.today().year
        this_q = ((date.today().month - 1) // 3) + 1
        quotes_this_q = [q for q in all_quotes if q.date.year == this_year and ((q.date.month - 1) // 3) + 1 == this_q]
        prev_q_month = ((this_q - 2) % 4) * 3 + 1
        prev_q_year = this_year if this_q > 1 else this_year - 1
        prev_q_num = this_q - 1 if this_q > 1 else 4
        quotes_prev_q = [q for q in all_quotes if q.date.year == prev_q_year and ((q.date.month - 1) // 3) + 1 == prev_q_num]
        quote_summary = {
            'open_net': sum(q.total_net for q in all_quotes if q.status in ('draft', 'sent')),
            'open_count': sum(1 for q in all_quotes if q.status in ('draft', 'sent')),
            'accepted_count': accepted,
            'conversion_rate': round(accepted / decided * 100) if decided > 0 else 0,
            'avg_value': round(sum(q.total_net for q in all_quotes if q.status in QUOTE_WON_STATUSES) / accepted, 2) if accepted > 0 else 0,
            'this_quarter_count': len(quotes_this_q),
            'prev_quarter_count': len(quotes_prev_q),
        }

    yoy_comparison = None
    if tab == 'invoices':
        this_year = date.today().year
        last_year = this_year - 1
        all_inv = Document.query.filter_by(doc_type='invoice').all()
        paid_this_year = [d for d in all_inv if d.status == 'paid' and d.payment_date and d.payment_date.year == this_year]
        paid_last_year = [d for d in all_inv if d.status == 'paid' and d.payment_date and d.payment_date.year == last_year]
        this_month = date.today().month
        paid_last_ytd = [d for d in paid_last_year if d.payment_date.month <= this_month]
        rev_this = sum(d.total_net for d in paid_this_year)
        rev_last_ytd = sum(d.total_net for d in paid_last_ytd)
        yoy_comparison = {
            'this_year': this_year,
            'last_year': last_year,
            'rev_this': rev_this,
            'rev_last_ytd': rev_last_ytd,
            'count_this': len(paid_this_year),
            'count_last_ytd': len(paid_last_ytd),
            'change_pct': round((rev_this - rev_last_ytd) / rev_last_ytd * 100) if rev_last_ytd > 0 else None,
        }

    filter_contact = Contact.query.get(contact_filter) if contact_filter else None
    return render_template('documents.html', documents=docs_paginated, tab=tab,
                           status_filter=status_filter, search=search, sort=sort,
                           today=date.today(), finance_summary=finance_summary,
                           quote_summary=quote_summary,
                           contact_filter=filter_contact,
                           yoy_comparison=yoy_comparison,
                           page=page, total_pages=total_pages, total_docs=total_docs)


@app.route('/documents/export')
def documents_export():
    year = request.args.get('year', date.today().year, type=int)
    doc_type = request.args.get('type', 'invoice')

    invoices = Document.query.filter(
        Document.doc_type == doc_type,
        db.extract('year', Document.date) == year,
        Document.status != 'draft',
    ).order_by(Document.date).all()

    import csv
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow([
        'Nummer', 'Typ', 'Datum', 'Kontakt', 'Kundennummer', 'Betreff',
        'Netto', 'USt', 'Brutto', 'Status', 'Bezahlt am', 'Versand am',
        'Zahlungsziel (Tage)',
    ])
    for d in invoices:
        writer.writerow([
            d.number,
            d.type_label,
            d.date.strftime('%d.%m.%Y'),
            d.contact.display_name,
            d.contact.customer_number or '',
            d.subject,
            f'{d.total_net:.2f}'.replace('.', ','),
            f'{d.total_vat:.2f}'.replace('.', ','),
            f'{d.total_gross:.2f}'.replace('.', ','),
            d.status_label,
            d.payment_date.strftime('%d.%m.%Y') if d.payment_date else '',
            d.sent_date.strftime('%d.%m.%Y') if d.sent_date else '',
            d.payment_term_days or '',
        ])

    output.seek(0)
    filename = f'{doc_type}_{year}_export.csv'
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8-sig')),
        mimetype='text/csv',
        download_name=filename,
    )


def _ensure_customer_number(contact):
    if contact.customer_number:
        return contact.customer_number
    seq = NumberSequence.query.filter_by(prefix='K').first()
    contact.customer_number = seq.generate()
    return contact.customer_number


@app.route('/documents/new')
def document_new():
    doc_type = request.args.get('type', 'quote')
    contact_id = request.args.get('contact_id', type=int)
    contacts = Contact.query.order_by(Contact.organisation, Contact.nachname).all()
    return render_template('document_form.html', document=None, doc_type=doc_type,
                           preselect_contact_id=contact_id, contacts=contacts,
                           unit_options=UNIT_OPTIONS, billing_types=BILLING_TYPES,
                           )


def _validate_document_data(data):
    errors = []
    if not data:
        return ['Keine Daten empfangen.'], {}
    if not data.get('contact_id'):
        errors.append('Kontakt ist erforderlich.')
    elif not Contact.query.get(data['contact_id']):
        errors.append('Kontakt nicht gefunden.')
    if not data.get('subject', '').strip():
        errors.append('Betreff ist erforderlich.')

    parsed = {}
    try:
        parsed['date'] = datetime.strptime(data.get('date', date.today().isoformat()), '%Y-%m-%d').date()
    except (ValueError, TypeError):
        errors.append('Ungültiges Datum.')
    for field, label in [('valid_until', 'Gültig bis'), ('due_date', 'Fällig am')]:
        if data.get(field):
            try:
                parsed[field] = datetime.strptime(data[field], '%Y-%m-%d').date()
            except (ValueError, TypeError):
                errors.append(f'Ungültiges Datum bei "{label}".')
        else:
            parsed[field] = None

    for g_idx, g in enumerate(data.get('groups', []), 1):
        for i_idx, item in enumerate(g.get('items', []), 1):
            price = item.get('unit_price', 0)
            qty = item.get('quantity', 0)
            try:
                float(price)
                float(qty)
            except (ValueError, TypeError):
                errors.append(f'Gruppe {g_idx}, Position {i_idx}: ungültiger Preis oder Menge.')

    return errors, parsed


@app.route('/documents/new', methods=['POST'])
def document_create():
    data = request.get_json()
    errors, parsed = _validate_document_data(data)
    if errors:
        return jsonify({'error': ' '.join(errors)}), 400

    doc_type = data.get('doc_type', 'quote')
    contact = Contact.query.get(data['contact_id'])
    _ensure_customer_number(contact)

    doc = Document(
        contact_id=contact.id,
        doc_type=doc_type,
        number=f'_PENDING_{uuid.uuid4().hex[:8]}',
        subject=data.get('subject', ''),
        status='draft',
        date=parsed['date'],
        valid_until=parsed.get('valid_until'),
        due_date=parsed.get('due_date'),
        conditions_text=data.get('conditions_text', ''),
        service_period=(data.get('service_period') or None),
        payment_term_days=data.get('payment_term_days', 14),
        internal_notes=data.get('internal_notes', ''),
        discount_type=data.get('discount_type'),
        discount_value=data.get('discount_value'),
    )
    db.session.add(doc)
    db.session.flush()

    for g_idx, group_data in enumerate(data.get('groups', [])):
        group = LineItemGroup(
            document_id=doc.id,
            position=g_idx,
            title=group_data.get('title', ''),
            description=group_data.get('description', ''),
        )
        db.session.add(group)
        db.session.flush()

        for i_idx, item_data in enumerate(group_data.get('items', [])):
            item = LineItem(
                group_id=group.id,
                position=i_idx,
                title=item_data.get('title', ''),
                description=item_data.get('description', ''),
                quantity=item_data.get('quantity', 1),
                unit=item_data.get('unit', 'Einmalig'),
                unit_price=item_data.get('unit_price', 0),
                vat_rate=item_data.get('vat_rate', 20),
                billing_type=item_data.get('billing_type', 'onetime'),
            )
            db.session.add(item)

    prefix = 'A' if doc_type == 'quote' else 'R'
    seq = NumberSequence.query.filter_by(prefix=prefix).first()
    doc.number = seq.generate()

    db.session.commit()

    if doc_type == 'quote' and not contact.stammdaten_verifiziert:
        _create_stammdaten_check_task(contact, doc)

    return jsonify({'success': True, 'id': doc.id, 'number': doc.number})


@app.route('/documents/<int:id>')
def document_view(id):
    doc = Document.query.get_or_404(id)
    return render_template('document_view.html', document=doc, today=date.today())


@app.route('/documents/<int:id>/edit')
def document_edit(id):
    doc = Document.query.get_or_404(id)
    if doc.status != 'draft':
        flash('Nur Entwürfe können bearbeitet werden.', 'error')
        return redirect(url_for('document_view', id=id))
    contacts = Contact.query.order_by(Contact.organisation, Contact.nachname).all()

    doc_json = {
        'contactId': str(doc.contact_id),
        'subject': doc.subject,
        'docDate': doc.date.isoformat(),
        'validUntil': doc.valid_until.isoformat() if doc.valid_until else '',
        'dueDate': doc.due_date.isoformat() if doc.due_date else '',
        'conditionsText': doc.conditions_text or '',
        'servicePeriod': doc.service_period or '',
        'paymentTermDays': doc.payment_term_days or 14,
        'internalNotes': doc.internal_notes or '',
        'updatedAt': doc.updated_at.isoformat() if doc.updated_at else '',
        'discountType': doc.discount_type or 'percent',
        'discountValue': float(doc.discount_value) if doc.discount_value else 0,
        'groups': [{
            'id': g.id,
            'title': g.title,
            'description': g.description or '',
            'items': [{
                'id': item.id,
                'title': item.title,
                'description': item.description or '',
                'quantity': float(item.quantity),
                'unit': item.unit,
                'unit_price': float(item.unit_price),
                'vat_rate': float(item.vat_rate),
                'billing_type': item.billing_type,
            } for item in g.items]
        } for g in doc.groups],
    }

    return render_template('document_form.html', document=doc, doc_type=doc.doc_type,
                           preselect_contact_id=doc.contact_id, contacts=contacts,
                           unit_options=UNIT_OPTIONS, billing_types=BILLING_TYPES,
                           doc_json=doc_json)


@app.route('/documents/<int:id>/edit', methods=['POST'])
def document_update(id):
    doc = Document.query.get_or_404(id)
    if doc.status != 'draft':
        return jsonify({'error': 'Nur Entwürfe bearbeitbar'}), 400

    data = request.get_json()
    errors, parsed = _validate_document_data(data)
    if errors:
        return jsonify({'error': ' '.join(errors)}), 400

    client_updated_at = data.get('updated_at')
    if client_updated_at and doc.updated_at:
        from datetime import datetime as dt
        try:
            client_ts = dt.fromisoformat(client_updated_at)
            if abs((doc.updated_at - client_ts).total_seconds()) > 1:
                return jsonify({'error': 'Dieses Dokument wurde in der Zwischenzeit geändert. Bitte Seite neu laden.'}), 409
        except (ValueError, TypeError):
            pass

    doc.contact_id = data['contact_id']
    doc.subject = data.get('subject', '')
    doc.date = parsed['date']
    doc.valid_until = parsed.get('valid_until')
    doc.due_date = parsed.get('due_date')
    doc.conditions_text = data.get('conditions_text', '')
    doc.service_period = (data.get('service_period') or None)
    doc.payment_term_days = data.get('payment_term_days', 14)
    doc.internal_notes = data.get('internal_notes', '')
    doc.discount_type = data.get('discount_type')
    doc.discount_value = data.get('discount_value')

    incoming_groups = data.get('groups', [])
    incoming_group_ids = {g['id'] for g in incoming_groups if g.get('id')}
    existing_groups = {g.id: g for g in doc.groups}

    for gid, group in existing_groups.items():
        if gid not in incoming_group_ids:
            db.session.delete(group)
    db.session.flush()

    for g_idx, group_data in enumerate(incoming_groups):
        gid = group_data.get('id')
        if gid and gid in existing_groups:
            group = existing_groups[gid]
            group.position = g_idx
            group.title = group_data.get('title', '')
            group.description = group_data.get('description', '')
        else:
            group = LineItemGroup(
                document_id=doc.id, position=g_idx,
                title=group_data.get('title', ''),
                description=group_data.get('description', ''),
            )
            db.session.add(group)
            db.session.flush()

        incoming_items = group_data.get('items', [])
        incoming_item_ids = {i['id'] for i in incoming_items if i.get('id')}
        existing_items = {i.id: i for i in group.items}

        for iid, item in existing_items.items():
            if iid not in incoming_item_ids:
                db.session.delete(item)

        for i_idx, item_data in enumerate(incoming_items):
            iid = item_data.get('id')
            if iid and iid in existing_items:
                item = existing_items[iid]
                item.group_id = group.id
                item.position = i_idx
                item.title = item_data.get('title', '')
                item.description = item_data.get('description', '')
                item.quantity = item_data.get('quantity', 1)
                item.unit = item_data.get('unit', 'Einmalig')
                item.unit_price = item_data.get('unit_price', 0)
                item.vat_rate = item_data.get('vat_rate', 20)
                item.billing_type = item_data.get('billing_type', 'onetime')
            else:
                item = LineItem(
                    group_id=group.id, position=i_idx,
                    title=item_data.get('title', ''),
                    description=item_data.get('description', ''),
                    quantity=item_data.get('quantity', 1),
                    unit=item_data.get('unit', 'Einmalig'),
                    unit_price=item_data.get('unit_price', 0),
                    vat_rate=item_data.get('vat_rate', 20),
                    billing_type=item_data.get('billing_type', 'onetime'),
                )
                db.session.add(item)

    db.session.commit()
    return jsonify({'success': True, 'id': doc.id})


@app.route('/documents/<int:id>/delete', methods=['POST'])
def document_delete(id):
    doc = Document.query.get_or_404(id)
    if doc.status != 'draft':
        flash('Nur Entwürfe können gelöscht werden.', 'error')
        return redirect(url_for('document_view', id=id))
    number = doc.number
    type_label = doc.type_label

    if doc.doc_type == 'invoice':
        seq = NumberSequence.query.filter_by(prefix='R').first()
        try:
            num_int = int(number.split('-')[-1])
        except (ValueError, AttributeError):
            num_int = None
        is_tail = seq is not None and num_int is not None and num_int == seq.next_number - 1
        if not is_tail:
            flash(f'Entwurf {number} ist nicht die zuletzt vergebene R-Nummer — Löschen '
                  f'würde eine Lücke reißen. Setze ihn stattdessen auf „storniert" '
                  f'(die Nummer bleibt erhalten).', 'error')
            return redirect(url_for('document_view', id=id))
        seq.next_number -= 1

    db.session.delete(doc)
    db.session.commit()
    flash(f'{type_label} {number} gelöscht.', 'success')
    return redirect(url_for('documents'))


@app.route('/documents/<int:id>/send-email', methods=['POST'])
def document_send_email(id):
    doc = Document.query.get_or_404(id)

    email_to = doc.contact.billing_email
    if not email_to:
        return jsonify({'error': 'Kontakt hat keine E-Mail-Adresse'}), 400

    company = CompanySettings.get_all()
    html_content = _render_document_html(doc)
    pdf_data, error = _generate_pdf(html_content, company)
    if error:
        return jsonify({'error': error}), 500

    if doc.doc_type == 'invoice':
        try:
            from zugferd import embed_zugferd_xml
            pdf_data = embed_zugferd_xml(pdf_data, doc, company)
        except Exception as e:
            app.logger.error(f'ZUGFeRD-Einbettung fehlgeschlagen fuer {doc.number}: {e}', exc_info=True)

    pdf_dir = tempfile.mkdtemp()
    pdf_path = os.path.join(pdf_dir, _beleg_dateiname(doc))
    with open(pdf_path, 'wb') as f:
        f.write(pdf_data)

    email_script = os.path.join(os.path.expanduser('~'), 'scripts', 'email_draft.py')
    if not os.path.exists(email_script):
        return jsonify({'error': 'email_draft.py nicht gefunden'}), 500

    vorname = doc.contact.vorname
    if vorname:
        greeting = f'Liebe {vorname}' if doc.contact.anrede == 'Frau' else f'Lieber {vorname}'
    elif doc.contact.organisation:
        greeting = f'Liebes {doc.contact.organisation} Team'
    else:
        greeting = 'Hallo zusammen'

    company = CompanySettings.get_all()
    subject = f'{doc.type_label} {doc.number} — {doc.subject}'

    signature = company.get('email_signature', '') or ''
    sig_block = f"\n\n{signature}" if signature else ''

    if doc.doc_type == 'invoice':
        body = f"""{greeting},

im Anhang findest du die Rechnung {doc.number} für {doc.subject}.

Zahlungsziel sind {doc.payment_term_days} Tage, aber falls du mehr Zeit brauchst, gib mir einfach kurz ein Zeichen.

Wenn dir irgendetwas auffällt oder du Fragen hast, sag einfach kurz Bescheid, das klären wir sofort.

Ich danke für die Zusammenarbeit.

Sonnige Grüße
{company['owner_name'].split()[0] if company.get('owner_name') else 'Max'}{sig_block}"""
    else:
        valid_line = ''
        if doc.valid_until:
            valid_line = f'Das Angebot ist gültig bis {doc.valid_until.strftime("%d.%m.%Y")}.\n\n'
        body = f"""{greeting},

im Anhang findest du das Angebot {doc.number} für {doc.subject}.

{valid_line}Falls du Fragen hast oder etwas angepasst werden soll, sag einfach kurz Bescheid — das klären wir direkt gemeinsam.

Ich freue mich auf deine Rückmeldung.

Sonnige Grüße
{company['owner_name'].split()[0] if company.get('owner_name') else 'Max'}{sig_block}"""

    try:
        result = subprocess.run(
            ['python', email_script, '--to', email_to, '--subject', subject, '--body', body, '--attachment', pdf_path],
            capture_output=True, text=True, errors='replace', timeout=15
        )
        if result.returncode != 0:
            return jsonify({'error': f'E-Mail-Fehler: {result.stderr[:200]}'}), 500
    except Exception as e:
        return jsonify({'error': f'E-Mail-Fehler: {str(e)[:200]}'}), 500
    finally:
        shutil.rmtree(pdf_dir, ignore_errors=True)

    status_changed = False
    if doc.status == 'draft':
        log = DocumentStatusLog(document_id=doc.id, from_status='draft', to_status='sent',
                                note='Automatisch per E-Mail-Versand')
        db.session.add(log)
        doc.status = 'sent'
        doc.sent_date = date.today()
        if doc.doc_type == 'invoice' and not doc.due_date:
            doc.due_date = date.today() + timedelta(days=doc.payment_term_days or 14)
        db.session.commit()
        status_changed = True

    return jsonify({
        'success': True,
        'message': f'E-Mail-Entwurf an {email_to} erstellt',
        'to': email_to,
        'status_changed': status_changed,
    })


@app.route('/documents/<int:id>/status', methods=['POST'])
def document_status(id):
    doc = Document.query.get_or_404(id)
    new_status = request.form.get('status')

    valid_transitions = {
        'quote': {
            'draft': ['sent', 'cancelled'],
            'sent': ['accepted', 'rejected', 'cancelled'],
            'accepted': ['sent', 'converted'],
            'converted': ['accepted'],
            'rejected': ['sent'],
            'expired': ['sent'],
            'cancelled': ['draft', 'sent'],
        },
        'invoice': {
            'draft': ['sent', 'cancelled'],
            'sent': ['paid', 'cancelled'],
            'overdue': ['paid', 'cancelled'],
            'paid': ['sent'],
            'cancelled': ['draft', 'sent'],
        },
    }
    allowed = valid_transitions.get(doc.doc_type, {}).get(doc.status, [])
    if new_status not in allowed:
        flash(f'Ungültige Statusänderung: {doc.status} → {new_status}', 'error')
        return redirect(url_for('document_view', id=id))

    if new_status == 'sent':
        doc.sent_date = date.today()
        if doc.doc_type == 'invoice':
            doc.due_date = date.today() + timedelta(days=doc.payment_term_days or 14)

    if new_status == 'accepted' and doc.doc_type == 'quote':
        contact = doc.contact
        if contact.status != 'won':
            old_status = contact.status
            contact.status = 'won'
            contact.deal_value = doc.total_gross
            sc = StatusChange(contact_id=contact.id,
                              from_status=old_status,
                              to_status='won',
                              deal_value_at_change=doc.total_gross)
            db.session.add(sc)

        for group in doc.groups:
            for item in group.items:
                if item.title:
                    task = Task(
                        title=f'{item.title}',
                        description=f'Aus Angebot {doc.number} — {group.title}',
                        status='todo',
                        priority='mittel',
                        contact_id=doc.contact_id,
                        created_by='System',
                    )
                    db.session.add(task)

        recurring_monthly = sum(
            (item.quantity or 0) * (item.unit_price or 0)
            for group in doc.groups for item in group.items
            if item.billing_type == 'recurring' and (item.unit or '').strip() == 'Monatlich'
        )
        has_other_recurring = any(
            item.billing_type == 'recurring' and (item.unit or '').strip() != 'Monatlich'
            for group in doc.groups for item in group.items
        )
        existing_retainer = Retainer.query.filter_by(contact_id=doc.contact_id, active=True).first()
        if recurring_monthly > 0 and not existing_retainer:
            rate = contact.hourly_rate or 0
            db.session.add(Retainer(
                contact_id=doc.contact_id,
                name=f'Retainer aus {doc.number} — Betrag/Intervall prüfen',
                monthly_fee=recurring_monthly,
                monthly_hours=round(recurring_monthly / rate, 1) if rate else 0,
                billing_interval_months=1,
                start_date=date.today(),
                active=True,
            ))
            flash(f'Retainer aus {doc.number} angelegt: {recurring_monthly:.0f} €/Mo — bitte Betrag & Intervall prüfen.', 'success')
        elif has_other_recurring and not existing_retainer:
            flash(f'{doc.number} hat wiederkehrende Positionen (nicht monatlich) — Retainer ggf. manuell anlegen.', 'info')

    if new_status == 'rejected' and doc.doc_type == 'quote':
        contact = doc.contact
        if contact.status not in ('won', 'lost'):
            old_contact_status = contact.status
            contact.status = 'lost'
            contact.lost_reason = 'no_need'
            sc = StatusChange(contact_id=contact.id,
                              from_status=old_contact_status,
                              to_status='lost',
                              deal_value_at_change=contact.deal_value)
            db.session.add(sc)

    if new_status == 'paid':
        doc.payment_date = date.today()

    storno_id = None
    if new_status == 'cancelled' and doc.doc_type == 'invoice' and doc.status in ('sent', 'overdue'):
        storno_id = _create_storno(doc)

    old_status = doc.status
    doc.status = new_status

    log = DocumentStatusLog(document_id=doc.id, from_status=old_status, to_status=new_status)
    if storno_id:
        log.note = 'Stornorechnung automatisch erstellt'
    db.session.add(log)

    db.session.commit()
    flash(f'{doc.type_label} {doc.number} → {doc.status_label}', 'success')
    if storno_id:
        flash(f'Stornorechnung erstellt.', 'success')
        return redirect(url_for('document_view', id=storno_id))
    return redirect(url_for('document_view', id=id))


def _create_storno(original):
    storno = Document(
        contact_id=original.contact_id,
        doc_type='invoice',
        number=f'_PENDING_{uuid.uuid4().hex[:8]}',
        subject=f'Storno zu: {original.number} — {original.subject}',
        status='draft',
        date=date.today(),
        service_period=original.service_period,
        conditions_text=f'Stornorechnung zu {original.number} vom {original.date.strftime("%d.%m.%Y")}.',
        payment_term_days=original.payment_term_days,
        internal_notes=f'Automatisch erstellt bei Stornierung von {original.number}.',
        source_document_id=original.id,
    )
    db.session.add(storno)
    db.session.flush()

    for group in original.groups:
        new_group = LineItemGroup(
            document_id=storno.id, position=group.position,
            title=group.title, description=group.description,
        )
        db.session.add(new_group)
        db.session.flush()
        for item in group.items:
            new_item = LineItem(
                group_id=new_group.id, position=item.position,
                title=item.title, description=item.description,
                quantity=-abs(float(item.quantity)),
                unit=item.unit, unit_price=item.unit_price,
                vat_rate=item.vat_rate, billing_type=item.billing_type,
            )
            db.session.add(new_item)

    seq = NumberSequence.query.filter_by(prefix='R').first()
    storno.number = seq.generate()

    return storno.id


@app.route('/documents/<int:id>/convert', methods=['POST'])
def document_convert(id):
    quote = Document.query.get_or_404(id)
    if quote.doc_type != 'quote' or quote.status not in QUOTE_WON_STATUSES:
        flash('Nur angenommene Angebote können umgewandelt werden.', 'error')
        return redirect(url_for('document_view', id=id))

    term_days = quote.payment_term_days or 14
    invoice = Document(
        contact_id=quote.contact_id,
        doc_type='invoice',
        number=f'_PENDING_{uuid.uuid4().hex[:8]}',
        subject=f'Rechnung zu: {quote.subject}',
        status='draft',
        date=date.today(),
        due_date=date.today() + timedelta(days=term_days),
        payment_term_days=term_days,
        conditions_text=quote.conditions_text,
        internal_notes=quote.internal_notes,
        discount_type=quote.discount_type,
        discount_value=quote.discount_value,
        source_document_id=quote.id,
    )
    db.session.add(invoice)
    db.session.flush()

    for group in quote.groups:
        new_group = LineItemGroup(
            document_id=invoice.id, position=group.position,
            title=group.title, description=group.description,
        )
        db.session.add(new_group)
        db.session.flush()
        for item in group.items:
            new_item = LineItem(
                group_id=new_group.id, position=item.position,
                title=item.title, description=item.description,
                quantity=item.quantity, unit=item.unit,
                unit_price=item.unit_price, vat_rate=item.vat_rate,
                billing_type=item.billing_type,
            )
            db.session.add(new_item)

    seq = NumberSequence.query.filter_by(prefix='R').first()
    invoice.number = seq.generate()

    if quote.status != 'converted':
        db.session.add(DocumentStatusLog(
            document_id=quote.id, from_status=quote.status, to_status='converted',
            note=f'Rechnung {invoice.number} erstellt',
        ))
        quote.status = 'converted'

    db.session.commit()
    flash(f'Rechnung {invoice.number} aus Angebot {quote.number} erstellt.', 'success')
    return redirect(url_for('document_view', id=invoice.id))


@app.route('/documents/<int:id>/duplicate', methods=['POST'])
def document_duplicate(id):
    source = Document.query.get_or_404(id)

    dup = Document(
        contact_id=source.contact_id,
        doc_type=source.doc_type,
        number=f'_PENDING_{uuid.uuid4().hex[:8]}',
        subject=source.subject,
        status='draft',
        date=date.today(),
        conditions_text=source.conditions_text,
        payment_term_days=source.payment_term_days,
        internal_notes=source.internal_notes,
        discount_type=source.discount_type,
        discount_value=source.discount_value,
    )
    db.session.add(dup)
    db.session.flush()

    for group in source.groups:
        new_group = LineItemGroup(
            document_id=dup.id, position=group.position,
            title=group.title, description=group.description,
        )
        db.session.add(new_group)
        db.session.flush()
        for item in group.items:
            new_item = LineItem(
                group_id=new_group.id, position=item.position,
                title=item.title, description=item.description,
                quantity=item.quantity, unit=item.unit,
                unit_price=item.unit_price, vat_rate=item.vat_rate,
                billing_type=item.billing_type,
            )
            db.session.add(new_item)

    prefix = 'A' if source.doc_type == 'quote' else 'R'
    seq = NumberSequence.query.filter_by(prefix=prefix).first()
    dup.number = seq.generate()

    db.session.commit()
    flash(f'{dup.type_label} {dup.number} als Kopie erstellt.', 'success')
    return redirect(url_for('document_edit', id=dup.id))


@app.route('/documents/<int:id>/preview')
def document_preview(id):
    doc = Document.query.get_or_404(id)
    return _render_document_html(doc)



def _beleg_dateiname(doc, dunning_level=0):
    roh = f'{doc.number}_{doc.contact.display_name}_{doc.date.strftime("%Y-%m-%d")}'
    if dunning_level:
        roh += f'_Mahnung{dunning_level}'
    sicher = re.sub(r'[^\w.\-]+', '-', roh, flags=re.UNICODE).strip('-')
    return f'{sicher}.pdf'


def _dunning_text(doc, level):
    datum = doc.date.strftime('%d.%m.%Y')
    betrag = format_euro_filter(doc.total_gross)
    faellig = doc.due_date.strftime('%d.%m.%Y') if doc.due_date else None
    faellig_satz = f' Fällig war sie am {faellig}.' if faellig else ''

    if level == 1:
        return (f'Kurze Erinnerung in eigener Sache: die Rechnung {doc.number} '
                f'vom {datum} über {betrag} € ist noch offen.{faellig_satz}')
    if level == 2:
        return (f'Die Rechnung {doc.number} vom {datum} über {betrag} € ist trotz '
                f'meiner Zahlungserinnerung noch offen.{faellig_satz} '
                f'Ich bitte dich um Überweisung in den nächsten Tagen.')
    return (f'Die Rechnung {doc.number} vom {datum} über {betrag} € ist weiterhin '
            f'unbeglichen.{faellig_satz} Das ist meine letzte Erinnerung per Mail.')


def _dunning_email(doc, level, company):
    dunning_info = DUNNING_LEVELS[level]
    kern = _dunning_text(doc, level)

    konto = f"IBAN {company['iban']}"
    if company.get('bank_name') and company.get('bank_account_holder'):
        konto += f" ({company['bank_name']}, {company['bank_account_holder']})"

    if level == 1:
        schluss = ('Falls deine Zahlung inzwischen schon unterwegs war, ist diese Mail '
                   'gegenstandslos. Und wenn etwas unklar ist oder es gerade nicht passt, '
                   'sag einfach kurz Bescheid, dann klären wir das direkt gemeinsam.')
    elif level == 2:
        schluss = ('Falls es einen Grund für die Verzögerung gibt, sag einfach kurz '
                   'Bescheid, dann finden wir eine Lösung.')
    else:
        schluss = ('Wenn es dazu etwas zu klären gibt, meld dich bitte kurz bei mir.')

    body = f"""{_du_anrede(doc.contact)}

{kern}

Die Rechnung findest du noch einmal im Anhang. Überweisen kannst du auf:
{konto}

{schluss}

Sonnige Grüße
{company['owner_name']}
{company['company_name']}
{company['phone']}"""

    return f'{dunning_info["title"]}: Rechnung {doc.number}', body


def _create_dunning_draft(doc, level, company):
    email_to = doc.contact.billing_email
    if not email_to:
        return False, 'keine E-Mail-Adresse am Kontakt'

    email_script = os.path.join(os.path.expanduser('~'), 'scripts', 'email_draft.py')
    if not os.path.exists(email_script):
        return False, 'email_draft.py nicht gefunden'

    try:
        html_content = _render_document_html(doc, level, DUNNING_LEVELS[level]['title'],
                                            _dunning_text(doc, level))
        pdf_data, error = _generate_pdf(html_content, company)
        if error or not pdf_data:
            return False, f'PDF-Fehler ({error or "leeres PDF"})'
    except Exception:
        return False, 'PDF-Fehler'

    pdf_dir = tempfile.mkdtemp()
    pdf_path = os.path.join(pdf_dir, _beleg_dateiname(doc, level))
    with open(pdf_path, 'wb') as f:
        f.write(pdf_data)

    try:
        subject, body = _dunning_email(doc, level, company)
        result = subprocess.run(
            ['python', email_script, '--to', email_to, '--subject', subject,
             '--body', body, '--attachment', pdf_path],
            capture_output=True, text=True, errors='replace', timeout=15
        )
        if result.returncode != 0:
            return False, 'email_draft.py mit Fehler beendet'
        return True, None
    except Exception:
        return False, 'E-Mail-Entwurf fehlgeschlagen'
    finally:
        shutil.rmtree(pdf_dir, ignore_errors=True)


@app.route('/documents/<int:id>/dunning', methods=['POST'])
def document_dunning(id):
    doc = Document.query.get_or_404(id)
    if doc.doc_type != 'invoice' or doc.status != 'overdue':
        flash('Mahnungen nur für überfällige Rechnungen möglich.', 'error')
        return redirect(url_for('document_view', id=id))

    next_level = doc.dunning_level + 1
    if next_level > 3:
        flash('Maximale Mahnstufe erreicht.', 'error')
        return redirect(url_for('document_view', id=id))

    doc.dunning_level = next_level
    doc.dunning_last_sent = date.today()

    log = DocumentStatusLog(document_id=doc.id, from_status='overdue', to_status='overdue',
                            note=f'{DUNNING_LEVELS[next_level]["title"]} erstellt (Stufe {next_level})')
    db.session.add(log)
    db.session.commit()

    dunning_info = DUNNING_LEVELS[next_level]
    flash(f'{dunning_info["title"]} für {doc.number} erstellt.', 'success')

    ok, fehler = _create_dunning_draft(doc, next_level, CompanySettings.get_all())
    if ok:
        flash(f'E-Mail-Entwurf an {doc.contact.billing_email} erstellt.', 'success')
    else:
        flash(f'Kein E-Mail-Entwurf erstellt ({fehler}). {dunning_info["title"]} '
              f'für {doc.number} ist als Stufe {next_level} vermerkt, der Versand '
              f'muss von Hand passieren.', 'warning')

    return redirect(url_for('document_pdf', id=id, dunning=next_level))


@app.route('/documents/<int:id>/pdf')
def document_pdf(id):
    doc = Document.query.get_or_404(id)

    dunning_level = request.args.get('dunning', 0, type=int)

    if not dunning_level and doc.status != 'draft' and doc.pdf_cache and doc.pdf_cached_at:
        if doc.updated_at and doc.pdf_cached_at >= doc.updated_at:
            filename = f'{doc.number}_{doc.contact.display_name}_{doc.date.strftime("%Y-%m-%d")}.pdf'
            return send_file(io.BytesIO(doc.pdf_cache), mimetype='application/pdf', download_name=filename)
    dunning_title = ''
    dunning_text = ''
    if dunning_level and dunning_level in DUNNING_LEVELS:
        dunning_title = DUNNING_LEVELS[dunning_level]['title']
        dunning_text = _dunning_text(doc, dunning_level)

    company = CompanySettings.get_all()
    html_content = _render_document_html(doc, dunning_level, dunning_title, dunning_text)
    pdf_data, error = _generate_pdf(html_content, company)

    if error:
        flash(error, 'error')
        return redirect(url_for('document_view', id=id))

    if doc.doc_type == 'invoice' and not dunning_level:
        try:
            from zugferd import embed_zugferd_xml
            pdf_data = embed_zugferd_xml(pdf_data, doc, company)
        except Exception as e:
            app.logger.error(f'ZUGFeRD-Einbettung fehlgeschlagen fuer {doc.number}: {e}', exc_info=True)
            flash(f'E-Rechnungs-XML konnte nicht eingebettet werden ({doc.number}). '
                  f'PDF wurde ohne ZUGFeRD erzeugt.', 'error')

    if not dunning_level and doc.status != 'draft':
        doc.pdf_cache = pdf_data
        doc.pdf_cached_at = datetime.now(timezone.utc).replace(tzinfo=None)
        db.session.commit()

    filename = f'{doc.number}_{doc.contact.display_name}_{doc.date.strftime("%Y-%m-%d")}'
    if dunning_level:
        filename += f'_Mahnung{dunning_level}'
    filename += '.pdf'

    return send_file(io.BytesIO(pdf_data), mimetype='application/pdf', download_name=filename)



def parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(value.strip(), '%Y-%m-%d').date()
    except (ValueError, AttributeError):
        return None



def _spawn_recurring_task(task):
    if not task.recurrence_days:
        return None
    new_due = stellwerk_write.next_recurrence_date(task.due_date or date.today(), task.recurrence_days)
    max_pos = db.session.query(db.func.coalesce(db.func.max(Task.position), -1)).filter(Task.status == 'todo').scalar()
    new_task = Task(
        title=task.title,
        description=task.description,
        status='todo',
        priority=task.priority,
        effort=task.effort,
        impact=task.impact,
        category=task.category,
        due_date=new_due,
        contact_id=task.contact_id,
        created_by='System',
        recurrence_days=task.recurrence_days,
        position=max_pos + 1,
    )
    new_task.calculate_score()
    db.session.add(new_task)
    db.session.flush()
    old_tags = db.session.query(task_tags.c.tag).filter(task_tags.c.task_id == task.id).all()
    for (tag_name,) in old_tags:
        db.session.execute(task_tags.insert().values(task_id=new_task.id, tag=tag_name))
    return new_task


@app.route('/board')
def board():
    from sqlalchemy import case

    category_filter = request.args.get('category')
    priority_filter = request.args.get('priority')
    tag_filter = request.args.get('tag')
    contact_filter = request.args.get('contact_id', type=int)
    search = request.args.get('search', '').strip()

    base_query = Task.query.options(joinedload(Task.contact))

    if search:
        conds = [Task.title.ilike(f'%{search}%')]
        id_q = search.lstrip('#').strip()
        if id_q.isdigit():
            conds.append(Task.id == int(id_q))
        base_query = base_query.filter(db.or_(*conds))
    if category_filter:
        base_query = base_query.filter(Task.category == category_filter)
    if priority_filter:
        base_query = base_query.filter(Task.priority == priority_filter)
    if contact_filter:
        base_query = base_query.filter(Task.contact_id == contact_filter)
    if tag_filter:
        base_query = base_query.filter(Task.id.in_(
            db.session.query(task_tags.c.task_id).filter(task_tags.c.tag == tag_filter)
        ))

    order = [
        case((Task.due_date.is_(None), 0), else_=1),
        Task.due_date.asc(),
        case((Task.due_date.is_(None), Task.id)).asc(),
        case((Task.score.is_(None), 1), else_=0),
        Task.score.desc(),
        Task.position,
    ]

    done_order = [
        case((Task.completed_at.is_(None), 1), else_=0),
        Task.completed_at.desc(),
        Task.id.desc(),
    ]

    today = date.today()
    scheduled_cutoff = today + timedelta(days=BOARD_SCHEDULED_LEAD_DAYS)
    hide_scheduled = not search
    within_lead = db.or_(Task.due_date.is_(None), Task.due_date <= scheduled_cutoff)

    board_data = {}
    for stage_key, stage_label, stage_color in BOARD_STAGES:
        stage_query = base_query.filter(Task.status == stage_key)
        if hide_scheduled and stage_key != 'erledigt':
            stage_query = stage_query.filter(within_lead)
        stage_query = stage_query.order_by(*(done_order if stage_key == 'erledigt' else order))
        if stage_key == 'erledigt' and not search:
            stage_query = stage_query.limit(20)
        board_data[stage_key] = stage_query.all()

    scheduled_count = base_query.filter(
        Task.status != 'erledigt',
        Task.due_date.isnot(None),
        Task.due_date > scheduled_cutoff,
    ).count() if hide_scheduled else 0

    categories = db.session.query(Task.category).filter(Task.category.isnot(None)).distinct().all()
    categories = sorted([c[0] for c in categories])

    tags_list = db.session.query(task_tags.c.tag).distinct().all()
    tags_list = sorted([t[0] for t in tags_list])

    return render_template('board.html',
                           board_data=board_data,
                           stages=BOARD_STAGES,
                           categories=categories,
                           tags_list=tags_list,
                           priority_options=BOARD_PRIORITY_OPTIONS,
                           effort_options=BOARD_EFFORT_OPTIONS,
                           impact_options=BOARD_IMPACT_OPTIONS,
                           created_by_options=CREATED_BY_OPTIONS,
                           category_filter=category_filter,
                           priority_filter=priority_filter,
                           tag_filter=tag_filter,
                           contact_filter=contact_filter,
                           search=search,
                           recurrence_presets=RECURRENCE_PRESETS,
                           scheduled_count=scheduled_count,
                           today=today)


@app.route('/mandate')
def mandate():
    today = date.today()
    prio_rank = {'hoch': 0, 'mittel': 1, 'niedrig': 2}

    retainer_by_contact = {}
    for r in Retainer.query.filter_by(active=True).all():
        retainer_by_contact.setdefault(r.contact_id, r)
    mandate_ids = {c.id for c in Contact.query.filter_by(mandat_aktiv=True).all()}
    ohne_flag_ids = set(retainer_by_contact) - mandate_ids
    retainer_ohne_flag = sorted(
        Contact.query.filter(Contact.id.in_(ohne_flag_ids)).all() if ohne_flag_ids else [],
        key=lambda c: (c.organisation or c.display_name or '').lower())

    def group_of(c):
        modell = c.abrechnungsmodell or ''
        org = (c.organisation or '').lower()
        if modell == 'agentur_partner' or 'partneragentur' in org:
            return (1, 'Agentur-Partner')
        if modell == 'pro_bono':
            return (2, 'Pro bono')
        return (0, 'DW-Mandate')

    contacts = Contact.query.filter(Contact.id.in_(mandate_ids)).all() if mandate_ids else []
    partneragentur_subs = sorted(
        [(c.id, c.organisation or c.display_name) for c in contacts if c.abrechnungsmodell == 'agentur_partner'],
        key=lambda t: (t[1] or '').lower())

    cards = []
    for c in contacts:
        open_tasks = [t for t in c.tasks if t.status != 'erledigt']
        open_tasks.sort(key=lambda t: (
            0 if (t.due_date and t.due_date < today) else 1,
            t.due_date or date.max,
            prio_rank.get(t.priority, 1),
            -(t.score or 0),
        ))
        grp_order, grp_label = group_of(c)
        is_dach = grp_label == 'Agentur-Partner' and 'partneragentur' in (c.organisation or '').lower()
        cards.append({
            'contact': c,
            'retainer': retainer_by_contact.get(c.id),
            'tasks': open_tasks,
            'next_id': open_tasks[0].id if open_tasks else None,
            'overdue': sum(1 for t in open_tasks if t.due_date and t.due_date < today),
            'group_order': grp_order,
            'group': grp_label,
            'erinnerung': c.erinnerung,
            'erinnerungstext': c.erinnerungstext,
            'is_dach': is_dach,
            'subs': partneragentur_subs if is_dach else [],
        })

    cards.sort(key=lambda m: (m['group_order'], (m['contact'].organisation or m['contact'].display_name or '').lower()))
    stage_labels = {s[0]: s[1] for s in BOARD_STAGES}
    total_tasks = sum(len(m['tasks']) for m in cards)
    total_overdue = sum(m['overdue'] for m in cards)
    return render_template('mandate.html', cards=cards, stage_labels=stage_labels,
                           total_tasks=total_tasks, total_overdue=total_overdue, today=today,
                           retainer_ohne_flag=retainer_ohne_flag)


@app.route('/board/cleanup')
def board_cleanup():
    open_stages = [s[0] for s in BOARD_STAGES if s[0] != 'erledigt']
    tasks_list = (Task.query.options(joinedload(Task.contact))
                  .filter(Task.status.in_(open_stages))
                  .order_by(Task.updated_at.asc())
                  .all())
    stale_count = sum(1 for t in tasks_list if t.is_stale)
    stage_labels = {s[0]: s[1] for s in BOARD_STAGES}
    return render_template('board_cleanup.html',
                           tasks=tasks_list,
                           stale_count=stale_count,
                           stage_labels=stage_labels,
                           threshold=TASK_STALE_THRESHOLD_DAYS,
                           today=date.today())


@app.route('/board/scheduled')
def board_scheduled():
    today = date.today()
    scheduled_cutoff = today + timedelta(days=BOARD_SCHEDULED_LEAD_DAYS)
    tasks_list = (Task.query.options(joinedload(Task.contact))
                  .filter(Task.status != 'erledigt',
                          Task.due_date.isnot(None),
                          Task.due_date > scheduled_cutoff)
                  .order_by(Task.due_date.asc())
                  .all())
    stage_labels = {s[0]: s[1] for s in BOARD_STAGES}
    return render_template('board_scheduled.html',
                           tasks=tasks_list,
                           stage_labels=stage_labels,
                           lead_days=BOARD_SCHEDULED_LEAD_DAYS,
                           cutoff=scheduled_cutoff,
                           today=today)


@app.route('/board/tasks/new', methods=['GET', 'POST'])
def board_task_new():
    if request.method == 'POST':
        status = request.form.get('status', 'backlog')
        if status == 'idee':
            status = 'backlog'
        max_pos = db.session.query(db.func.coalesce(db.func.max(Task.position), -1)).filter(Task.status == status).scalar()

        recurrence_raw = request.form.get('recurrence_days')
        recurrence_days = int(recurrence_raw) if recurrence_raw else None

        task = Task(
            title=request.form.get('title', '').strip(),
            description=request.form.get('description', '').strip() or None,
            status=status,
            priority=request.form.get('priority', 'mittel'),
            effort=request.form.get('effort') or None,
            impact=request.form.get('impact') or None,
            category=request.form.get('category', '').strip() or None,
            due_date=parse_date(request.form.get('due_date')),
            contact_id=request.form.get('contact_id', type=int) or None,
            created_by=request.form.get('created_by', 'Max'),
            recurrence_days=recurrence_days,
            position=max_pos + 1
        )

        if not task.title:
            flash('Titel ist erforderlich', 'error')
            contacts_list = Contact.query.order_by(Contact.organisation, Contact.nachname).all()
            return render_template('board_task_form.html', task=None, contacts=contacts_list,
                                   stages=BOARD_STAGES, priority_options=BOARD_PRIORITY_OPTIONS,
                                   effort_options=BOARD_EFFORT_OPTIONS, impact_options=BOARD_IMPACT_OPTIONS,
                                   created_by_options=CREATED_BY_OPTIONS,
                                   recurrence_presets=RECURRENCE_PRESETS)

        task.calculate_score()
        db.session.add(task)
        db.session.commit()

        tags_str = request.form.get('tags', '').strip()
        if tags_str:
            for tag_name in [t.strip() for t in tags_str.split(',') if t.strip()]:
                db.session.execute(task_tags.insert().values(task_id=task.id, tag=tag_name))
            db.session.commit()

        flash('Task erstellt', 'success')
        return redirect(url_for('board'))

    contacts_list = Contact.query.order_by(Contact.organisation, Contact.nachname).all()
    preselect_status = request.args.get('status', 'backlog')
    preselect_contact = request.args.get('contact_id', type=int)

    return render_template('board_task_form.html', task=None, contacts=contacts_list,
                           stages=BOARD_STAGES, priority_options=BOARD_PRIORITY_OPTIONS,
                           effort_options=BOARD_EFFORT_OPTIONS, impact_options=BOARD_IMPACT_OPTIONS,
                           created_by_options=CREATED_BY_OPTIONS,
                           recurrence_presets=RECURRENCE_PRESETS,
                           preselect_status=preselect_status, preselect_contact=preselect_contact)


@app.route('/board/tasks/<int:id>', methods=['GET', 'POST'])
def board_task_view(id):
    task = Task.query.get_or_404(id)

    if request.method == 'POST':
        task.title = request.form.get('title', '').strip()
        task.description = request.form.get('description', '').strip() or None

        old_status = task.status
        task.status = request.form.get('status', task.status)
        task.priority = request.form.get('priority', task.priority)
        task.effort = request.form.get('effort') or None
        task.impact = request.form.get('impact') or None
        task.category = request.form.get('category', '').strip() or None
        task.due_date = parse_date(request.form.get('due_date'))
        task.contact_id = request.form.get('contact_id', type=int) or None
        task.created_by = request.form.get('created_by', task.created_by)

        recurrence_raw = request.form.get('recurrence_days')
        task.recurrence_days = int(recurrence_raw) if recurrence_raw else None

        if task.status == 'erledigt' and old_status != 'erledigt':
            task.completed_at = datetime.now(timezone.utc).replace(tzinfo=None)
            _spawn_recurring_task(task)
        elif task.status != 'erledigt' and old_status == 'erledigt':
            task.completed_at = None

        task.calculate_score()

        db.session.execute(task_tags.delete().where(task_tags.c.task_id == task.id))
        tags_str = request.form.get('tags', '').strip()
        if tags_str:
            for tag_name in [t.strip() for t in tags_str.split(',') if t.strip()]:
                db.session.execute(task_tags.insert().values(task_id=task.id, tag=tag_name))

        db.session.commit()
        flash('Task aktualisiert', 'success')
        return redirect(url_for('board'))

    comments = TaskComment.query.filter_by(task_id=id).order_by(TaskComment.created_at).all()
    contacts_list = Contact.query.order_by(Contact.organisation, Contact.nachname).all()

    current_tags = db.session.query(task_tags.c.tag).filter(task_tags.c.task_id == id).all()
    current_tags_str = ', '.join([t[0] for t in current_tags])

    return render_template('board_task_detail.html', task=task, comments=comments,
                           contacts=contacts_list, current_tags=current_tags_str,
                           stages=BOARD_STAGES, priority_options=BOARD_PRIORITY_OPTIONS,
                           effort_options=BOARD_EFFORT_OPTIONS, impact_options=BOARD_IMPACT_OPTIONS,
                           created_by_options=CREATED_BY_OPTIONS, recurrence_presets=RECURRENCE_PRESETS,
                           today=date.today())


@app.route('/board/tasks/<int:id>/delete', methods=['POST'])
def board_task_delete(id):
    task = Task.query.get_or_404(id)
    if task.recurrence_days:
        Task.query.filter(
            Task.title == task.title,
            Task.recurrence_days == task.recurrence_days,
            Task.id != task.id
        ).update({Task.recurrence_days: None})
    db.session.execute(task_tags.delete().where(task_tags.c.task_id == id))
    db.session.delete(task)
    db.session.commit()
    flash('Task gelöscht', 'success')
    return redirect(url_for('board'))


@app.route('/board/tasks/<int:id>/move', methods=['POST'])
def board_task_move(id):
    task = Task.query.get_or_404(id)
    data = request.get_json()
    new_status = data.get('status')
    new_position = data.get('position', 0)

    valid_statuses = [s[0] for s in BOARD_STAGES]
    if not new_status or new_status not in valid_statuses:
        return jsonify({'success': False, 'error': 'Invalid status'}), 400

    old_status = task.status

    try:
        Task.query.filter(
            Task.status == new_status,
            Task.position >= new_position,
            Task.id != task.id
        ).update({Task.position: Task.position + 1})

        task.status = new_status
        task.position = new_position

        if new_status == 'erledigt' and old_status != 'erledigt':
            task.completed_at = datetime.now(timezone.utc).replace(tzinfo=None)
            _spawn_recurring_task(task)
        elif new_status != 'erledigt' and old_status == 'erledigt':
            task.completed_at = None

        task.calculate_score()
        db.session.commit()
    except Exception:
        db.session.rollback()
        return jsonify({'success': False, 'error': 'Database error'}), 500

    return jsonify({'success': True})


@app.route('/board/tasks/<int:id>/quick-edit', methods=['POST'])
def board_task_quick_edit(id):
    task = Task.query.get_or_404(id)
    data = request.get_json()

    old_status = task.status

    if 'title' in data:
        task.title = data['title']
    if 'status' in data:
        task.status = data['status']
    if 'priority' in data:
        task.priority = data['priority']
    if 'effort' in data:
        task.effort = data['effort'] or None
    if 'impact' in data:
        task.impact = data['impact'] or None
    if 'category' in data:
        task.category = data['category'] or None
    if 'due_date' in data:
        task.due_date = parse_date(data['due_date'])
    if 'description' in data:
        task.description = data['description'] or None
    if 'recurrence_days' in data:
        task.recurrence_days = int(data['recurrence_days']) if data['recurrence_days'] else None

    if task.status == 'erledigt' and old_status != 'erledigt':
        task.completed_at = datetime.now(timezone.utc).replace(tzinfo=None)
        _spawn_recurring_task(task)
    elif task.status != 'erledigt' and old_status == 'erledigt':
        task.completed_at = None

    task.calculate_score()
    db.session.commit()

    return jsonify({'success': True, 'score': task.display_score})


@app.route('/board/tasks/<int:id>/comments', methods=['POST'])
def board_task_comment(id):
    task = Task.query.get_or_404(id)
    content = request.form.get('content', '').strip()
    author = request.form.get('author', 'Max')

    if content:
        comment = TaskComment(task_id=id, author=author, content=content)
        db.session.add(comment)
        db.session.commit()
        flash('Kommentar hinzugefügt', 'success')

    return redirect(url_for('board_task_view', id=id))


@app.route('/board/api/categories')
def board_api_categories():
    categories = db.session.query(Task.category).filter(Task.category.isnot(None)).distinct().all()
    return jsonify(sorted([c[0] for c in categories]))


@app.route('/board/api/tags')
def board_api_tags():
    tags_result = db.session.query(task_tags.c.tag).distinct().all()
    return jsonify(sorted([t[0] for t in tags_result]))



@app.route('/api/dashboard-widgets', methods=['POST'])
def api_dashboard_widgets():
    data = request.get_json()
    widget = data.get('widget', '')
    visible = data.get('visible', True)

    current = CompanySettings.get('dashboard_hidden_widgets') or ''
    hidden = [w.strip() for w in current.split(',') if w.strip()]

    if visible and widget in hidden:
        hidden.remove(widget)
    elif not visible and widget not in hidden:
        hidden.append(widget)

    row = CompanySettings.query.get('dashboard_hidden_widgets')
    if row:
        row.value = ','.join(hidden)
    else:
        db.session.add(CompanySettings(key='dashboard_hidden_widgets', value=','.join(hidden)))
    db.session.commit()
    return jsonify({'success': True})


@app.route('/api/notifications')
def api_notifications():
    today = date.today()
    DEAD_STATUSES = ['lost', 'closed_no_response']
    overdue_reminders = Contact.query.filter(
        Contact.erinnerung < today,
        Contact.erinnerung.isnot(None),
        Contact.status.notin_(DEAD_STATUSES)
    ).count()
    today_reminders = Contact.query.filter(
        Contact.erinnerung == today,
        Contact.status.notin_(DEAD_STATUSES)
    ).count()
    overdue_payments = ExpectedPayment.query.filter(
        ExpectedPayment.status == 'pending', ExpectedPayment.expected_date < today
    ).count()
    overdue_invoices = Document.query.filter(
        Document.doc_type == 'invoice', Document.status == 'overdue'
    ).count()
    stale_quotes = Document.query.filter(
        Document.doc_type == 'quote', Document.status == 'sent'
    ).all()
    stale_count = sum(1 for q in stale_quotes if (today - (q.sent_date or q.date)).days > 14)
    retainer_due = 0
    for r in Retainer.query.filter_by(active=True).all():
        last = r.last_billed_date or r.start_date
        if (today.year - last.year) * 12 + today.month - last.month >= r.billing_interval_months:
            retainer_due += 1

    total = overdue_reminders + today_reminders + overdue_payments + overdue_invoices + stale_count + retainer_due
    items = []
    if overdue_reminders:
        items.append({'label': f'{overdue_reminders} Erinnerung{"en" if overdue_reminders != 1 else ""} überfällig', 'color': 'red', 'url': '/'})
    if today_reminders:
        items.append({'label': f'{today_reminders} heute fällig', 'color': 'amber', 'url': '/'})
    if overdue_payments:
        items.append({'label': f'{overdue_payments} Zahlung{"en" if overdue_payments != 1 else ""} überfällig', 'color': 'red', 'url': '/'})
    if overdue_invoices:
        items.append({'label': f'{overdue_invoices} Rechnung{"en" if overdue_invoices != 1 else ""} überfällig', 'color': 'red', 'url': '/documents?tab=invoices&status=overdue'})
    if stale_count:
        items.append({'label': f'{stale_count} Angebot{"e" if stale_count != 1 else ""} ohne Reaktion', 'color': 'amber', 'url': '/'})
    if retainer_due:
        items.append({'label': f'{retainer_due} Retainer-Abrechnung{"en" if retainer_due != 1 else ""} fällig', 'color': 'violet', 'url': '/'})
    overdue_tasks = Task.query.filter(Task.due_date < today, Task.status != 'erledigt').count()
    badges = {}
    if overdue_tasks > 0:
        badges['board'] = overdue_tasks
    if overdue_invoices > 0:
        badges['documents'] = overdue_invoices

    return jsonify({'total': total, 'items': items, 'badges': badges})


@app.route('/api/briefing')
def api_briefing():
    today = date.today()

    open_tasks = Task.query.filter(Task.status != 'erledigt').order_by(Task.due_date).all()
    task_rows = [{
        'id': t.id,
        'title': t.title,
        'priority': t.priority,
        'due_date': t.due_date.isoformat() if t.due_date else None,
        'contact': t.contact.organisation if t.contact else None,
        'category': t.category,
        'overdue': bool(t.due_date and t.due_date < today),
    } for t in open_tasks]

    billing_rows = _gather_billing_rows(today, _get_default_hourly_rate())
    billing_alerts = [{
        'contact_id': r['contact'].id,
        'contact': r['contact'].organisation,
        'ampel': r['ampel'],
        'action': r['action'],
        'delivered_h': r['delivered_h'],
        'budget_h': r['budget_h'],
    } for r in billing_rows if r['ampel'] in ('rot', 'gelb')]

    wiki_rows = [{
        'slug': w.slug,
        'title': w.name,
        'entry_type': w.entry_type,
        'trigger_phrases': w.trigger_phrases,
    } for w in WikiEntry.query.filter_by(status='active').all()]



    return jsonify({
        'generated_at': datetime.now().isoformat(timespec='seconds'),
        'tasks': task_rows,
        'tasks_overdue': [t for t in task_rows if t['overdue']],
        'billing_alerts': billing_alerts,
        'wiki_registry': wiki_rows,
    })


@app.route('/api/contacts/check-duplicate')
def api_check_duplicate():
    org = request.args.get('organisation', '').strip()
    email = request.args.get('email', '').strip()
    vorname = request.args.get('vorname', '').strip()
    nachname = request.args.get('nachname', '').strip()
    exclude_id = request.args.get('exclude_id', type=int)

    duplicates = []
    conditions = []
    if email:
        conditions.append(db.or_(
            Contact.mail_person.ilike(email),
            Contact.mail_organisation.ilike(email)
        ))
    if org:
        conditions.append(Contact.organisation.ilike(f'%{org}%'))
    if vorname and nachname:
        conditions.append(db.and_(
            Contact.vorname.ilike(f'%{vorname}%'),
            Contact.nachname.ilike(f'%{nachname}%')
        ))

    if conditions:
        query = Contact.query.filter(db.or_(*conditions))
        if exclude_id:
            query = query.filter(Contact.id != exclude_id)
        duplicates = [{
            'id': c.id,
            'name': c.display_name,
            'email': c.primary_email or '',
            'status': c.status,
            'url': url_for('contact_view', id=c.id),
        } for c in query.limit(5).all()]

    return jsonify(duplicates)


@app.route('/api/contacts/search')
def api_contacts_search():
    q = request.args.get('q', '').strip()
    results = []

    if q:
        search_term = f'%{q}%'
        contacts_result = Contact.query.filter(
            db.or_(
                Contact.organisation.ilike(search_term),
                Contact.vorname.ilike(search_term),
                Contact.nachname.ilike(search_term),
                Contact.mail_person.ilike(search_term),
                Contact.mail_organisation.ilike(search_term),
            )
        ).order_by(Contact.updated_at.desc()).limit(6).all()
        for c in contacts_result:
            results.append({
                'name': c.display_name,
                'sub': c.full_name if c.organisation and c.full_name else (c.primary_email or ''),
                'initial': c.display_name[0:1].upper() if c.display_name else '?',
                'status': c.status,
                'url': url_for('contact_view', id=c.id),
                'type': 'contact',
            })

        docs_result = Document.query.filter(
            db.or_(Document.number.ilike(search_term), Document.subject.ilike(search_term))
        ).order_by(Document.date.desc()).limit(4).all()
        for d in docs_result:
            results.append({
                'name': f'{d.number} — {d.subject[:40]}',
                'sub': d.contact.display_name if d.contact else '',
                'initial': 'A' if d.doc_type == 'quote' else 'R',
                'status': d.status,
                'url': url_for('document_view', id=d.id),
                'type': 'document',
            })

        tasks_result = Task.query.filter(Task.title.ilike(search_term)).limit(3).all()
        for t in tasks_result:
            results.append({
                'name': t.title[:50],
                'sub': t.contact.display_name if t.contact else 'Kein Kontakt',
                'initial': 'T',
                'status': t.status,
                'url': url_for('board_task_view', id=t.id),
                'type': 'task',
            })
    else:
        contacts_result = Contact.query.order_by(Contact.updated_at.desc()).limit(8).all()
        for c in contacts_result:
            results.append({
                'name': c.display_name,
                'sub': c.full_name if c.organisation and c.full_name else (c.primary_email or ''),
                'initial': c.display_name[0:1].upper() if c.display_name else '?',
                'status': c.status,
                'url': url_for('contact_view', id=c.id),
                'type': 'contact',
            })

    return jsonify(results)



def _render_snippet(raw_snippet):
    escaped = str(escape(raw_snippet))
    escaped = escaped.replace(
        search_fts.SNIPPET_START,
        '<mark class="bg-gold-200/70 text-gray-900 dark:bg-gold-500/40 dark:text-white rounded px-0.5">'
    )
    escaped = escaped.replace(search_fts.SNIPPET_END, '</mark>')
    return Markup(escaped)


@app.route('/search')
def global_search():
    q = request.args.get('q', '').strip()
    contacts, tasks, documents, wiki_items = [], [], [], []

    if q:
        raw = search_fts.search_all(q)

        for hit in raw['contacts']:
            c = Contact.query.get(hit['id'])
            if c:
                c.search_snippet = _render_snippet(hit['snippet'])
                contacts.append(c)

        for hit in raw['tasks']:
            t = Task.query.get(hit['id'])
            if t:
                t.search_snippet = _render_snippet(hit['snippet'])
                tasks.append(t)

        for hit in raw['documents']:
            d = Document.query.get(hit['id'])
            if d:
                d.search_snippet = _render_snippet(hit['snippet'])
                documents.append(d)

        for hit in raw['wiki_pages']:
            wp = WikiPage.query.get(hit['id'])
            if wp:
                wp.search_snippet = _render_snippet(hit['snippet'])
                wiki_items.append(('page', wp))

        for hit in raw['wiki_entries']:
            we = WikiEntry.query.get(hit['id'])
            if we:
                we.search_snippet = _render_snippet(hit['snippet'])
                wiki_items.append(('entry', we))

    total = len(contacts) + len(tasks) + len(documents) + len(wiki_items)

    return render_template('search.html', q=q, total=total,
                           contacts=contacts, tasks=tasks,
                           documents=documents, wiki_items=wiki_items,
                           entry_types=ENTRY_TYPES)



@app.route('/report')
def monthly_report():
    today = date.today()
    month = request.args.get('month', today.month, type=int)
    year = request.args.get('year', today.year, type=int)

    month_start = date(year, month, 1)
    if month == 12:
        month_end = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        month_end = date(year, month + 1, 1) - timedelta(days=1)

    paid_invoices = Document.query.filter(
        Document.doc_type == 'invoice',
        Document.status == 'paid',
        Document.payment_date >= month_start,
        Document.payment_date <= month_end,
    ).all()
    revenue_net = sum(d.total_net for d in paid_invoices)
    revenue_gross = sum(d.total_gross for d in paid_invoices)

    new_contacts = Contact.query.filter(
        Contact.created_at >= datetime.combine(month_start, datetime.min.time()),
        Contact.created_at <= datetime.combine(month_end, datetime.max.time()),
    ).count()

    won_changes = StatusChange.query.filter(
        StatusChange.to_status == 'won',
        StatusChange.changed_at >= datetime.combine(month_start, datetime.min.time()),
        StatusChange.changed_at <= datetime.combine(month_end, datetime.max.time()),
    ).all()

    quotes_sent = Document.query.filter(
        Document.doc_type == 'quote',
        Document.status.in_(['sent', 'rejected', 'expired', *QUOTE_WON_STATUSES]),
        Document.date >= month_start,
        Document.date <= month_end,
    ).count()
    quotes_accepted = Document.query.filter(
        Document.doc_type == 'quote',
        Document.status.in_(QUOTE_WON_STATUSES),
        Document.date >= month_start,
        Document.date <= month_end,
    ).count()

    time_entries = TimeEntry.query.filter(
        TimeEntry.datum >= month_start,
        TimeEntry.datum <= month_end,
        TimeEntry.timer_start.is_(None),
    ).all()
    total_hours = sum(e.dauer_minuten for e in time_entries) / 60

    open_invoices = Document.query.filter(
        Document.doc_type == 'invoice',
        Document.status.in_(['sent', 'overdue']),
    ).all()
    open_invoice_sum = sum(d.total_net for d in open_invoices)

    customer_revenue = {}
    for inv in paid_invoices:
        name = inv.contact.display_name
        customer_revenue[name] = customer_revenue.get(name, 0) + inv.total_net
    top_customers = sorted(customer_revenue.items(), key=lambda x: x[1], reverse=True)[:5]

    month_names = ['', 'Januar', 'Februar', 'März', 'April', 'Mai', 'Juni',
                   'Juli', 'August', 'September', 'Oktober', 'November', 'Dezember']

    return render_template('report.html',
        month=month, year=year, month_name=month_names[month],
        month_start=month_start, month_end=month_end,
        paid_invoices=paid_invoices, revenue_net=revenue_net, revenue_gross=revenue_gross,
        new_contacts=new_contacts, won_changes=won_changes,
        quotes_sent=quotes_sent, quotes_accepted=quotes_accepted,
        total_hours=round(total_hours, 1), time_entries=time_entries,
        open_invoices=open_invoices, open_invoice_sum=open_invoice_sum,
        top_customers=top_customers, today=today)


@app.route('/documents/bulk-dunning', methods=['POST'])
def bulk_dunning():
    overdue = Document.query.filter(
        Document.doc_type == 'invoice',
        Document.status == 'overdue',
    ).all()

    sent = 0
    skipped = 0
    errors = []

    company = CompanySettings.get_all()

    for inv in overdue:
        ok, fehler = _create_dunning_draft(inv, 1, company)
        if ok:
            sent += 1
        elif fehler == 'keine E-Mail-Adresse am Kontakt':
            skipped += 1
        else:
            errors.append(f'{inv.number}: {fehler}')

    if sent > 0:
        flash(f'{sent} Zahlungserinnerung{"en" if sent != 1 else ""} als E-Mail-Entwurf erstellt.', 'success')
    if skipped > 0:
        flash(f'{skipped} übersprungen (keine E-Mail-Adresse).', 'warning')
    if errors:
        flash(f'Fehler bei: {", ".join(errors[:3])}', 'error')

    return redirect(url_for('documents', tab='invoices', status='overdue'))



WIKI_ERZEUGT = 'erzeugt'
WIKI_ERZEUGT_BEREICHE = ('kunden', 'werkzeuge', 'prozesse', 'entscheidungen', 'zeitstrahl')
WIKI_ERZEUGT_BEREICH_TITEL = {'kunden': 'Kunden', 'werkzeuge': 'Werkzeuge', 'prozesse': 'Prozesse',
                              'entscheidungen': 'Entscheidungen', 'zeitstrahl': 'Zeitstrahl'}


def _erzeugt_bereich(slug):
    teil = (slug or '').split('-', 2)
    return teil[1] if len(teil) > 1 and teil[0] == WIKI_ERZEUGT and teil[1] in WIKI_ERZEUGT_BEREICHE else None


def _erzeugt_quellen(page_id):
    try:
        return [dict(r._mapping) for r in db.session.execute(db.text(
            "SELECT quelle_pfad, gebaut_am FROM wiki_quellen_hash WHERE seite_id = :id ORDER BY quelle_pfad"),
            {'id': page_id})]
    except Exception:
        db.session.rollback()
        return []


def _erzeugt_html(inhalt):
    import markdown
    from markdown.extensions.toc import slugify_unicode
    md = markdown.Markdown(extensions=['tables', 'sane_lists', 'toc'],
                           extension_configs={'toc': {'toc_depth': '2-3', 'slugify': slugify_unicode}})
    md.preprocessors.deregister('html_block')
    md.inlinePatterns.deregister('html')
    html = md.convert(inhalt or '')
    return Markup(html), (Markup(md.toc) if md.toc_tokens else None)


def _erzeugt_gesperrt(page):
    if page.category != WIKI_ERZEUGT:
        return None
    flash('Diese Seite ist erzeugt. Bitte an der Quelle ändern, der nächste Füll-Lauf baut sie neu.', 'error')
    return redirect(url_for('wiki_page_view', id=page.id))


@app.route('/wiki')
def wiki():
    q = request.args.get('q', '').strip()
    tab = request.args.get('tab', 'pages')
    category = request.args.get('category', '')
    entry_type = request.args.get('type', '')
    lb = request.args.get('lb', '')

    pages = WikiPage.query
    entries = WikiEntry.query

    if q:
        pages = pages.filter(WikiPage.title.ilike(f'%{q}%') | WikiPage.content.ilike(f'%{q}%'))
        entries = entries.filter(
            WikiEntry.name.ilike(f'%{q}%') |
            WikiEntry.description.ilike(f'%{q}%') |
            WikiEntry.trigger_phrases.ilike(f'%{q}%')
        )
    if category:
        pages = pages.filter(WikiPage.category == category)
    if entry_type:
        entries = entries.filter(WikiEntry.entry_type == entry_type)
    if lb:
        entries = entries.filter(WikiEntry.leistungsbereich == lb)

    erzeugt = pages.filter(WikiPage.category == WIKI_ERZEUGT).order_by(WikiPage.title).all()
    erzeugt_nach_bereich = {b: [p for p in erzeugt if _erzeugt_bereich(p.slug) == b] for b in WIKI_ERZEUGT_BEREICHE}
    pages = pages.filter(db.or_(WikiPage.category.is_(None), WikiPage.category != WIKI_ERZEUGT))
    pages = pages.order_by(WikiPage.pinned.desc(), WikiPage.title).all()
    entries = entries.order_by(WikiEntry.entry_type, WikiEntry.name).all()

    all_categories = sorted(set(p.category for p in WikiPage.query.with_entities(WikiPage.category).distinct()
                                if p.category and p.category != WIKI_ERZEUGT))

    stale_cutoff = datetime.now() - timedelta(days=30)
    stale_pages = WikiPage.query.filter(WikiPage.updated_at < stale_cutoff,
                                        db.or_(WikiPage.category.is_(None), WikiPage.category != WIKI_ERZEUGT)
                                        ).order_by(WikiPage.updated_at).all()
    draft_entries = WikiEntry.query.filter(WikiEntry.status == 'draft').all()
    deprecated_entries = WikiEntry.query.filter(WikiEntry.status == 'deprecated').all()

    return render_template('wiki.html',
                           pages=pages, entries=entries,
                           tab=tab, q=q,
                           category=category, entry_type=entry_type, lb=lb,
                           all_categories=all_categories,
                           leistungsbereiche=LEISTUNGSBEREICHE,
                           entry_types=ENTRY_TYPES,
                           entry_statuses=ENTRY_STATUSES,
                           stale_pages=stale_pages,
                           draft_entries=draft_entries,
                           deprecated_entries=deprecated_entries,
                           erzeugt_anzahl=len(erzeugt),
                           erzeugt_nach_bereich=erzeugt_nach_bereich,
                           erzeugt_bereich_titel=WIKI_ERZEUGT_BEREICH_TITEL)


@app.route('/wiki/pages/new', methods=['GET', 'POST'])
def wiki_page_new():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        if not title:
            flash('Titel ist Pflichtfeld', 'error')
            return redirect(url_for('wiki_page_new'))
        if request.form.get('category', '').strip() == WIKI_ERZEUGT:
            flash('Die Kategorie „erzeugt“ vergibt nur der Füll-Lauf.', 'error')
            return redirect(url_for('wiki_page_new'))
        page = WikiPage(
            title=title,
            slug=WikiPage.slugify(title),
            content=request.form.get('content', ''),
            category=request.form.get('category', '').strip() or None,
            pinned='pinned' in request.form,
        )
        db.session.add(page)
        db.session.flush()
        _save_file_links_for_page(page, request.form)
        db.session.commit()
        flash(f'Wiki-Seite "{title}" erstellt', 'success')
        return redirect(url_for('wiki_page_view', id=page.id))
    return render_template('wiki_form.html', page=None, mode='page',
                           all_categories=sorted(set(p.category for p in WikiPage.query.with_entities(WikiPage.category).distinct() if p.category)))


@app.route('/wiki/pages/<int:id>')
def wiki_page_view(id):
    page = WikiPage.query.get_or_404(id)
    erzeugt = page.category == WIKI_ERZEUGT
    erzeugt_html, erzeugt_toc = _erzeugt_html(page.content) if erzeugt else (None, None)
    return render_template('wiki_view.html', page=page, mode='page', erzeugt=erzeugt,
                           erzeugt_quellen=_erzeugt_quellen(page.id) if erzeugt else [],
                           erzeugt_html=erzeugt_html, erzeugt_toc=erzeugt_toc)


@app.route('/wiki/pages/<int:id>/edit', methods=['GET', 'POST'])
def wiki_page_edit(id):
    page = WikiPage.query.get_or_404(id)
    gesperrt = _erzeugt_gesperrt(page)
    if gesperrt:
        return gesperrt
    if request.method == 'POST':
        if request.form.get('category', '').strip() == WIKI_ERZEUGT:
            flash('Die Kategorie „erzeugt“ vergibt nur der Füll-Lauf.', 'error')
            return redirect(url_for('wiki_page_edit', id=page.id))
        page.title = request.form.get('title', '').strip() or page.title
        page.slug = WikiPage.slugify(page.title)
        page.content = request.form.get('content', '')
        page.category = request.form.get('category', '').strip() or None
        page.pinned = 'pinned' in request.form
        WikiFileLink.query.filter_by(wiki_page_id=page.id).delete()
        _save_file_links_for_page(page, request.form)
        db.session.commit()
        flash('Wiki-Seite aktualisiert', 'success')
        return redirect(url_for('wiki_page_view', id=page.id))
    return render_template('wiki_form.html', page=page, mode='page',
                           all_categories=sorted(set(p.category for p in WikiPage.query.with_entities(WikiPage.category).distinct() if p.category)))


@app.route('/wiki/pages/<int:id>/delete', methods=['POST'])
def wiki_page_delete(id):
    page = WikiPage.query.get_or_404(id)
    gesperrt = _erzeugt_gesperrt(page)
    if gesperrt:
        return gesperrt
    title = page.title
    db.session.delete(page)
    db.session.commit()
    flash(f'Wiki-Seite "{title}" gelöscht', 'success')
    return redirect(url_for('wiki'))



@app.route('/wiki/entries/new', methods=['GET', 'POST'])
def wiki_entry_new():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        if not name:
            flash('Name ist Pflichtfeld', 'error')
            return redirect(url_for('wiki_entry_new'))
        entry = WikiEntry(
            name=name,
            slug=WikiPage.slugify(name),
            entry_type=request.form.get('entry_type', 'skill'),
            description=request.form.get('description', ''),
            trigger_phrases=request.form.get('trigger_phrases', ''),
            example_call=request.form.get('example_call', ''),
            leistungsbereich=request.form.get('leistungsbereich', ''),
            file_path=request.form.get('file_path', ''),
            status=request.form.get('status', 'active'),
            notes=request.form.get('notes', ''),
        )
        db.session.add(entry)
        db.session.flush()
        _save_file_links_for_entry(entry, request.form)
        db.session.commit()
        flash(f'{ENTRY_TYPES.get(entry.entry_type, "Eintrag")} "{name}" erstellt', 'success')
        return redirect(url_for('wiki_entry_view', id=entry.id))
    return render_template('wiki_form.html', page=None, mode='entry',
                           leistungsbereiche=LEISTUNGSBEREICHE,
                           entry_types=ENTRY_TYPES,
                           entry_statuses=ENTRY_STATUSES)


@app.route('/wiki/entries/<int:id>')
def wiki_entry_view(id):
    entry = WikiEntry.query.get_or_404(id)
    return render_template('wiki_view.html', entry=entry, mode='entry',
                           leistungsbereiche=LEISTUNGSBEREICHE,
                           entry_types=ENTRY_TYPES,
                           entry_statuses=ENTRY_STATUSES)


@app.route('/wiki/entries/<int:id>/edit', methods=['GET', 'POST'])
def wiki_entry_edit(id):
    entry = WikiEntry.query.get_or_404(id)
    if request.method == 'POST':
        entry.name = request.form.get('name', '').strip() or entry.name
        entry.slug = WikiPage.slugify(entry.name)
        entry.entry_type = request.form.get('entry_type', entry.entry_type)
        entry.description = request.form.get('description', '')
        entry.trigger_phrases = request.form.get('trigger_phrases', '')
        entry.example_call = request.form.get('example_call', '')
        entry.leistungsbereich = request.form.get('leistungsbereich', '')
        entry.file_path = request.form.get('file_path', '')
        entry.status = request.form.get('status', entry.status)
        entry.notes = request.form.get('notes', '')
        WikiFileLink.query.filter_by(wiki_entry_id=entry.id).delete()
        _save_file_links_for_entry(entry, request.form)
        db.session.commit()
        flash('Eintrag aktualisiert', 'success')
        return redirect(url_for('wiki_entry_view', id=entry.id))
    return render_template('wiki_form.html', page=entry, mode='entry',
                           leistungsbereiche=LEISTUNGSBEREICHE,
                           entry_types=ENTRY_TYPES,
                           entry_statuses=ENTRY_STATUSES)


@app.route('/wiki/entries/<int:id>/delete', methods=['POST'])
def wiki_entry_delete(id):
    entry = WikiEntry.query.get_or_404(id)
    name = entry.name
    db.session.delete(entry)
    db.session.commit()
    flash(f'Eintrag "{name}" gelöscht', 'success')
    return redirect(url_for('wiki'))


@app.route('/api/wiki/recent', methods=['POST'])
def wiki_recent_items():
    data = request.get_json(silent=True) or {}
    page_ids = data.get('page_ids', [])[:8]
    entry_ids = data.get('entry_ids', [])[:8]
    result = []
    for pid in page_ids:
        p = WikiPage.query.get(pid)
        if p:
            result.append({'type': 'page', 'id': p.id, 'title': p.title, 'category': p.category or '',
                           'url': url_for('wiki_page_view', id=p.id)})
    for eid in entry_ids:
        e = WikiEntry.query.get(eid)
        if e:
            result.append({'type': 'entry', 'id': e.id, 'title': e.name, 'entry_type': e.entry_type,
                           'url': url_for('wiki_entry_view', id=e.id)})
    return jsonify(result)


def _save_file_links_for_page(page, form):
    i = 0
    while True:
        label = form.get(f'fl_label_{i}', '').strip()
        fpath = form.get(f'fl_path_{i}', '').strip()
        if not label and not fpath:
            break
        if label and fpath:
            link = WikiFileLink(label=label, file_path=fpath,
                                description=form.get(f'fl_desc_{i}', '').strip(),
                                wiki_page_id=page.id)
            db.session.add(link)
        i += 1


def _save_file_links_for_entry(entry, form):
    i = 0
    while True:
        label = form.get(f'fl_label_{i}', '').strip()
        fpath = form.get(f'fl_path_{i}', '').strip()
        if not label and not fpath:
            break
        if label and fpath:
            link = WikiFileLink(label=label, file_path=fpath,
                                description=form.get(f'fl_desc_{i}', '').strip(),
                                wiki_entry_id=entry.id)
            db.session.add(link)
        i += 1



INCOMING_PDF_DIR = os.environ.get('STELLWERK_INCOMING_PDF_DIR',
                                  os.path.join(BASE_DIR, 'uploads', 'incoming_invoices'))


def _ensure_incoming_pdf_dir():
    os.makedirs(INCOMING_PDF_DIR, exist_ok=True)


def _calc_vat_amount(net, vat_rate):
    return round(float(net) * float(vat_rate) / 100.0, 2)


def _parse_decimal(val, default=0):
    if val is None or val == '':
        return default
    try:
        return float(str(val).replace(',', '.'))
    except (ValueError, TypeError):
        return default


def _parse_iso_date(val):
    if not val:
        return None
    try:
        return datetime.strptime(val, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None


@app.route('/incoming-invoices')
def incoming_invoices():
    status_filter = request.args.get('status', 'all')
    supplier_filter = request.args.get('supplier_id', type=int)
    category_filter = request.args.get('category_id', type=int)
    year_filter = request.args.get('year', type=int)
    search = request.args.get('q', '').strip()
    beleg_filter = request.args.get('beleg', 'all')

    query = IncomingInvoice.query.options(
        joinedload(IncomingInvoice.supplier),
        joinedload(IncomingInvoice.category),
    )

    if status_filter == 'open':
        query = query.filter(IncomingInvoice.status == 'open')
    elif status_filter == 'paid':
        query = query.filter(IncomingInvoice.status == 'paid')
    elif status_filter == 'overdue':
        query = query.filter(IncomingInvoice.status == 'open',
                             IncomingInvoice.due_date < date.today())

    if beleg_filter in ('fehlt', 'vorhanden', 'geprueft'):
        query = query.filter(IncomingInvoice.beleg_status == beleg_filter)
    elif beleg_filter == 'missing':
        query = query.filter(IncomingInvoice.beleg_status == 'fehlt')

    if supplier_filter:
        query = query.filter(IncomingInvoice.supplier_id == supplier_filter)
    if category_filter:
        query = query.filter(IncomingInvoice.expense_category_id == category_filter)
    if year_filter:
        query = query.filter(db.extract('year', IncomingInvoice.posting_date) == year_filter)
    if search:
        query = query.filter(db.or_(
            IncomingInvoice.number.ilike(f'%{search}%'),
            IncomingInvoice.subject.ilike(f'%{search}%'),
            IncomingInvoice.external_invoice_number.ilike(f'%{search}%'),
        ))

    invoices_list = query.order_by(IncomingInvoice.posting_date.desc(),
                                    IncomingInvoice.id.desc()).all()

    summary = {
        'count': len(invoices_list),
        'total_net': round(sum(float(i.net_amount) for i in invoices_list), 2),
        'total_vat': round(sum(float(i.vat_amount) for i in invoices_list), 2),
        'total_gross': round(sum(float(i.gross_amount) for i in invoices_list), 2),
        'open_net': round(sum(float(i.net_amount) for i in invoices_list
                               if i.status == 'open'), 2),
        'open_count': sum(1 for i in invoices_list if i.status == 'open'),
        'overdue_count': sum(1 for i in invoices_list if i.is_overdue),
        'beleg_fehlt_count': sum(1 for i in invoices_list if i.beleg_status == 'fehlt'),
        'beleg_fehlt_gross': round(sum(float(i.gross_amount) for i in invoices_list
                                       if i.beleg_status == 'fehlt'), 2),
        'beleg_vorhanden_count': sum(1 for i in invoices_list if i.beleg_status == 'vorhanden'),
        'beleg_geprueft_count': sum(1 for i in invoices_list if i.beleg_status == 'geprueft'),
    }

    suppliers = Contact.query.filter(Contact.contact_type.in_(['supplier', 'both'])) \
                              .order_by(Contact.organisation).all()
    categories = ExpenseCategory.query.filter_by(is_active=True) \
                                       .order_by(ExpenseCategory.sort_order).all()

    return render_template('incoming_invoices/list.html',
                           invoices=invoices_list,
                           summary=summary,
                           suppliers=suppliers,
                           categories=categories,
                           status_filter=status_filter,
                           supplier_filter=supplier_filter,
                           category_filter=category_filter,
                           year_filter=year_filter,
                           search=search,
                           beleg_filter=beleg_filter,
                           today=date.today(),
                           INCOMING_INVOICE_STATUSES=INCOMING_INVOICE_STATUSES,
                           BELEG_STATUSES=BELEG_STATUSES)


@app.route('/incoming-invoices/new')
def incoming_invoice_new():
    suppliers = Contact.query.filter(Contact.contact_type.in_(['supplier', 'both'])) \
                              .order_by(Contact.organisation).all()
    categories = ExpenseCategory.query.filter_by(is_active=True) \
                                       .order_by(ExpenseCategory.sort_order).all()
    return render_template('incoming_invoices/form.html',
                           invoice=None,
                           suppliers=suppliers,
                           categories=categories,
                           PAYMENT_METHODS=PAYMENT_METHODS,
                           INCOMING_INVOICE_STATUSES=INCOMING_INVOICE_STATUSES,
                           today=date.today())


@app.route('/incoming-invoices/new', methods=['POST'])
def incoming_invoice_create():
    form = request.form

    supplier_id = form.get('supplier_id', type=int)
    if not supplier_id:
        flash('Lieferant ist erforderlich.', 'error')
        return redirect(url_for('incoming_invoice_new'))

    supplier = Contact.query.get(supplier_id)
    if not supplier:
        flash('Lieferant nicht gefunden.', 'error')
        return redirect(url_for('incoming_invoice_new'))

    if supplier.contact_type == 'customer':
        supplier.contact_type = 'both'
    elif not supplier.contact_type:
        supplier.contact_type = 'supplier'

    invoice_date = _parse_iso_date(form.get('invoice_date')) or date.today()
    delivery_date = _parse_iso_date(form.get('delivery_date'))
    posting_date = _parse_iso_date(form.get('posting_date')) or invoice_date

    net_amount = _parse_decimal(form.get('net_amount'))
    vat_rate = _parse_decimal(form.get('vat_rate'), default=20)
    vat_amount_input = form.get('vat_amount', '').strip()
    if vat_amount_input:
        vat_amount = _parse_decimal(vat_amount_input)
    else:
        vat_amount = _calc_vat_amount(net_amount, vat_rate)
    gross_amount = round(net_amount + vat_amount, 2)

    gross_provided = form.get('gross_amount', '').strip()
    if gross_provided:
        gp = _parse_decimal(gross_provided)
        if abs(gp - gross_amount) > 0.01:
            flash(f'Brutto stimmt nicht: erwartet {gross_amount:.2f}, übergeben {gp:.2f}.', 'error')
            return redirect(url_for('incoming_invoice_new'))
        gross_amount = gp

    is_vat_deductible = form.get('is_vat_deductible') == 'on'
    vat_exemption_reason = form.get('vat_exemption_reason', '').strip() or None
    if vat_amount == 0 and not vat_exemption_reason and float(vat_rate) == 0:
        flash('Bei 0% USt ist ein Befreiungsgrund Pflicht (z.B. "Reverse Charge §19 UStG").', 'error')
        return redirect(url_for('incoming_invoice_new'))

    seq = NumberSequence.query.filter_by(prefix='E').first()
    if not seq:
        flash('Nummern-Sequenz E fehlt — Migration unvollständig.', 'error')
        return redirect(url_for('incoming_invoice_new'))
    number = seq.generate()

    inv = IncomingInvoice(
        number=number,
        supplier_id=supplier_id,
        expense_category_id=form.get('expense_category_id', type=int) or None,
        external_invoice_number=form.get('external_invoice_number', '').strip(),
        subject=form.get('subject', '').strip(),
        invoice_date=invoice_date,
        delivery_date=delivery_date,
        posting_date=posting_date,
        due_date=_parse_iso_date(form.get('due_date')),
        payment_date=_parse_iso_date(form.get('payment_date')),
        net_amount=net_amount,
        vat_rate=vat_rate,
        vat_amount=vat_amount,
        gross_amount=gross_amount,
        vat_exemption_reason=vat_exemption_reason,
        is_vat_deductible=is_vat_deductible,
        payment_method=form.get('payment_method') or None,
        status='paid' if form.get('payment_date') else 'open',
        internal_notes=form.get('internal_notes', '').strip() or None,
    )

    if not inv.external_invoice_number:
        flash('Lieferanten-Rechnungsnummer ist Pflicht (§11 UStG).', 'error')
        return redirect(url_for('incoming_invoice_new'))
    if not inv.subject:
        flash('Betreff/Beschreibung ist Pflicht.', 'error')
        return redirect(url_for('incoming_invoice_new'))

    try:
        db.session.add(inv)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f'Speichern fehlgeschlagen: {e}', 'error')
        return redirect(url_for('incoming_invoice_new'))

    pdf_file = request.files.get('pdf_file')
    if pdf_file and pdf_file.filename:
        _store_incoming_pdf(inv, pdf_file)
        db.session.commit()

    flash(f'Eingangsrechnung {inv.number} erfasst.', 'success')
    return redirect(url_for('incoming_invoice_view', id=inv.id))


@app.route('/incoming-invoices/<int:id>')
def incoming_invoice_view(id):
    inv = IncomingInvoice.query.options(
        joinedload(IncomingInvoice.supplier),
        joinedload(IncomingInvoice.category),
    ).get_or_404(id)
    return render_template('incoming_invoices/detail.html', invoice=inv, today=date.today(),
                           BELEG_STATUSES=BELEG_STATUSES)


@app.route('/incoming-invoices/<int:id>/edit')
def incoming_invoice_edit(id):
    inv = IncomingInvoice.query.get_or_404(id)
    suppliers = Contact.query.filter(Contact.contact_type.in_(['supplier', 'both'])) \
                              .order_by(Contact.organisation).all()
    categories = ExpenseCategory.query.filter_by(is_active=True) \
                                       .order_by(ExpenseCategory.sort_order).all()
    return render_template('incoming_invoices/form.html',
                           invoice=inv,
                           suppliers=suppliers,
                           categories=categories,
                           PAYMENT_METHODS=PAYMENT_METHODS,
                           INCOMING_INVOICE_STATUSES=INCOMING_INVOICE_STATUSES,
                           today=date.today())


@app.route('/incoming-invoices/<int:id>/edit', methods=['POST'])
def incoming_invoice_update(id):
    inv = IncomingInvoice.query.get_or_404(id)
    form = request.form

    inv.supplier_id = form.get('supplier_id', type=int) or inv.supplier_id
    inv.expense_category_id = form.get('expense_category_id', type=int) or None
    inv.external_invoice_number = form.get('external_invoice_number', inv.external_invoice_number).strip()
    inv.subject = form.get('subject', inv.subject).strip()
    inv.invoice_date = _parse_iso_date(form.get('invoice_date')) or inv.invoice_date
    inv.delivery_date = _parse_iso_date(form.get('delivery_date'))
    inv.posting_date = _parse_iso_date(form.get('posting_date')) or inv.posting_date
    inv.due_date = _parse_iso_date(form.get('due_date'))
    inv.payment_date = _parse_iso_date(form.get('payment_date'))

    inv.net_amount = _parse_decimal(form.get('net_amount'), default=float(inv.net_amount))
    inv.vat_rate = _parse_decimal(form.get('vat_rate'), default=float(inv.vat_rate))
    vat_input = form.get('vat_amount', '').strip()
    if vat_input:
        inv.vat_amount = _parse_decimal(vat_input)
    else:
        inv.vat_amount = _calc_vat_amount(float(inv.net_amount), float(inv.vat_rate))
    inv.gross_amount = round(float(inv.net_amount) + float(inv.vat_amount), 2)

    inv.vat_exemption_reason = form.get('vat_exemption_reason', '').strip() or None
    inv.is_vat_deductible = form.get('is_vat_deductible') == 'on'
    inv.payment_method = form.get('payment_method') or None
    inv.status = 'paid' if inv.payment_date else (form.get('status', 'open'))
    inv.internal_notes = form.get('internal_notes', '').strip() or None

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f'Update fehlgeschlagen: {e}', 'error')
        return redirect(url_for('incoming_invoice_edit', id=id))

    pdf_file = request.files.get('pdf_file')
    if pdf_file and pdf_file.filename:
        _store_incoming_pdf(inv, pdf_file)
        db.session.commit()

    flash(f'{inv.number} aktualisiert.', 'success')
    return redirect(url_for('incoming_invoice_view', id=inv.id))


@app.route('/incoming-invoices/<int:id>/mark-paid', methods=['POST'])
def incoming_invoice_mark_paid(id):
    inv = IncomingInvoice.query.get_or_404(id)
    pay_date = _parse_iso_date(request.form.get('payment_date')) or date.today()
    inv.payment_date = pay_date
    inv.status = 'paid'
    if request.form.get('payment_method'):
        inv.payment_method = request.form.get('payment_method')
    db.session.commit()
    flash(f'{inv.number} als bezahlt markiert ({pay_date.isoformat()}).', 'success')
    return redirect(url_for('incoming_invoice_view', id=id))


@app.route('/incoming-invoices/<int:id>/beleg-status', methods=['POST'])
def incoming_invoice_beleg_status(id):
    inv = IncomingInvoice.query.get_or_404(id)
    new_status = request.form.get('beleg_status', '').strip()
    valid = {code for code, _ in BELEG_STATUSES}
    if new_status not in valid:
        flash('Ungueltiger Beleg-Status.', 'error')
        return redirect(url_for('incoming_invoice_view', id=id))
    if new_status in ('vorhanden', 'geprueft') and not inv.document_file_path:
        flash('Ohne hinterlegtes PDF kann der Beleg nicht als vorhanden/geprueft markiert werden.', 'error')
        return redirect(url_for('incoming_invoice_view', id=id))
    inv.beleg_status = new_status
    db.session.commit()
    flash(f'{inv.number}: Beleg-Status = {inv.beleg_status_label}.', 'success')
    return redirect(url_for('incoming_invoice_view', id=id))


def _store_incoming_pdf(inv, file_storage):
    _ensure_incoming_pdf_dir()
    filename = file_storage.filename or ''
    if not filename.lower().endswith('.pdf'):
        flash('Nur PDF-Dateien erlaubt.', 'error')
        return False
    mime = (file_storage.mimetype or '').lower()
    if mime and mime != 'application/pdf':
        flash(f'Ungueltiger MIME-Type: {mime} (erwartet application/pdf).', 'error')
        return False
    target = os.path.join(INCOMING_PDF_DIR, f'{inv.number}.pdf')
    file_storage.save(target)
    size = os.path.getsize(target)
    if size == 0:
        os.remove(target)
        flash('PDF ist leer.', 'error')
        return False
    inv.document_file_path = filename
    if inv.beleg_status == 'fehlt':
        inv.beleg_status = 'vorhanden'
    return True


@app.route('/incoming-invoices/<int:id>/upload-pdf', methods=['POST'])
def incoming_invoice_upload_pdf(id):
    inv = IncomingInvoice.query.get_or_404(id)
    pdf_file = request.files.get('pdf_file')
    if not pdf_file or not pdf_file.filename:
        flash('Keine Datei ausgewaehlt.', 'error')
        return redirect(url_for('incoming_invoice_view', id=id))
    if _store_incoming_pdf(inv, pdf_file):
        db.session.commit()
        flash('PDF gespeichert.', 'success')
    return redirect(url_for('incoming_invoice_view', id=id))


@app.route('/incoming-invoices/<int:id>/pdf')
def incoming_invoice_pdf(id):
    inv = IncomingInvoice.query.get_or_404(id)
    pdf_path = os.path.join(INCOMING_PDF_DIR, f'{inv.number}.pdf')
    if not os.path.exists(pdf_path):
        flash('PDF nicht gefunden.', 'error')
        return redirect(url_for('incoming_invoice_view', id=id))
    download_name = f'{inv.number}_{inv.supplier.display_name}_{inv.invoice_date.strftime("%Y-%m-%d")}.pdf'
    return send_file(pdf_path, mimetype='application/pdf', download_name=download_name)


@app.route('/incoming-invoices/<int:id>/delete', methods=['POST'])
def incoming_invoice_delete(id):
    inv = IncomingInvoice.query.get_or_404(id)
    if inv.status == 'paid':
        flash('Bezahlte Eingangsrechnungen koennen nicht geloescht werden — Storno-Pflicht.', 'error')
        return redirect(url_for('incoming_invoice_view', id=id))
    number = inv.number
    pdf_path = os.path.join(INCOMING_PDF_DIR, f'{inv.number}.pdf')
    db.session.delete(inv)
    db.session.commit()
    if os.path.exists(pdf_path):
        try:
            os.remove(pdf_path)
        except OSError:
            pass
    flash(f'Eingangsrechnung {number} geloescht.', 'success')
    return redirect(url_for('incoming_invoices'))



if __name__ == '__main__':
    print("\n" + "="*50)
    print("  Sales CRM is running!")
    print("  Open: http://localhost:5000")
    print("="*50 + "\n")
    try:
        _autopost_on_start()
    except Exception as _autopost_err:
        print(f'Content-Poster uebersprungen: {_autopost_err}', file=sys.stderr)
    app.run(host=os.environ.get('STELLWERK_HOST', '127.0.0.1'),
            debug=os.environ.get('STELLWERK_DEBUG') == '1', port=5000, use_reloader=False)
