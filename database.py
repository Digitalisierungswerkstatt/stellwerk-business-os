import os
from datetime import date
from openpyxl import load_workbook
from models import db, Contact, Interaction, Task


def init_db(app):
    db.init_app(app)
    with app.app_context():
        db.create_all()
        _run_migrations()
        _seed_email_templates()
        _seed_tool_types()
        _seed_expense_categories()


def _run_migrations():
    migrations = [
        "ALTER TABLE contacts ADD COLUMN steuernummer VARCHAR(100)",
        "ALTER TABLE contacts ADD COLUMN gemeinnuetzig BOOLEAN DEFAULT 0",
        "ALTER TABLE contacts ADD COLUMN strasse VARCHAR(300)",
        "ALTER TABLE contacts ADD COLUMN customer_number VARCHAR(20)",
        "ALTER TABLE contacts ADD COLUMN gruendungsdatum DATE",
        "ALTER TABLE time_entries ADD COLUMN recurring_days INTEGER",
        "ALTER TABLE time_entries ADD COLUMN last_recurred DATE",
        "ALTER TABLE time_entries ADD COLUMN billed_document_id INTEGER REFERENCES documents(id)",
        "ALTER TABLE contacts ADD COLUMN uid VARCHAR(100)",
        "ALTER TABLE customer_accounts ADD COLUMN client_login_email VARCHAR(200)",
        "ALTER TABLE customer_accounts ADD COLUMN client_login_password VARCHAR(300)",
        "ALTER TABLE customer_accounts ADD COLUMN done BOOLEAN DEFAULT 0",
        "ALTER TABLE contacts ADD COLUMN mail_rechnung VARCHAR(200)",
        "ALTER TABLE contacts ADD COLUMN firmenbuchnummer VARCHAR(100)",
        "ALTER TABLE contacts ADD COLUMN stammdaten_verifiziert BOOLEAN DEFAULT 0",
        "ALTER TABLE contacts ADD COLUMN rechnung_name VARCHAR(300)",
        "ALTER TABLE contacts ADD COLUMN rechnung_uid VARCHAR(100)",
        "ALTER TABLE contacts ADD COLUMN rechnung_strasse VARCHAR(300)",
        "ALTER TABLE contacts ADD COLUMN rechnung_plz VARCHAR(20)",
        "ALTER TABLE contacts ADD COLUMN rechnung_stadt VARCHAR(100)",
        "ALTER TABLE contacts ADD COLUMN rechnung_land VARCHAR(100)",
        "ALTER TABLE documents ADD COLUMN service_period TEXT",
    ]
    for sql in migrations:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()

    try:
        db.session.execute(db.text(
            "CREATE UNIQUE INDEX IF NOT EXISTS ix_contacts_customer_number ON contacts(customer_number)"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    from models import NumberSequence
    for prefix, start in [('A', 11), ('R', 6), ('K', 1)]:
        existing = NumberSequence.query.filter_by(prefix=prefix).first()
        if not existing:
            seq = NumberSequence(prefix=prefix, next_number=start, padding=5)
            db.session.add(seq)
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        cursor = db.session.execute(db.text(
            "SELECT name FROM sqlite_master WHERE type='table' AND name='task_comments'"
        ))
        if not cursor.fetchone():
            db.session.execute(db.text("DROP TABLE IF EXISTS tasks"))
            db.session.execute(db.text("""
                CREATE TABLE tasks (
                    id INTEGER PRIMARY KEY,
                    title VARCHAR(500) NOT NULL,
                    description TEXT,
                    status VARCHAR(20) NOT NULL DEFAULT 'idee',
                    priority VARCHAR(10) NOT NULL DEFAULT 'mittel',
                    effort VARCHAR(5),
                    impact VARCHAR(10),
                    score FLOAT,
                    category VARCHAR(100),
                    due_date DATE,
                    contact_id INTEGER REFERENCES contacts(id),
                    created_by VARCHAR(50) NOT NULL DEFAULT 'Max',
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    updated_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    completed_at DATETIME,
                    position INTEGER NOT NULL DEFAULT 0
                )
            """))
            db.session.execute(db.text("""
                CREATE TABLE task_comments (
                    id INTEGER PRIMARY KEY,
                    task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
                    author VARCHAR(50) NOT NULL,
                    content TEXT NOT NULL,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP
                )
            """))
            db.session.execute(db.text("""
                CREATE TABLE task_tags (
                    task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
                    tag VARCHAR(100) NOT NULL,
                    PRIMARY KEY (task_id, tag)
                )
            """))
            db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("ALTER TABLE tasks ADD COLUMN recurrence_days INTEGER"))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("ALTER TABLE service_templates ADD COLUMN use_count INTEGER NOT NULL DEFAULT 0"))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("ALTER TABLE contacts ADD COLUMN hourly_rate FLOAT"))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("ALTER TABLE expected_payments ADD COLUMN received_date DATE"))
        db.session.commit()
    except Exception:
        db.session.rollback()

    for col in [
        "ALTER TABLE documents ADD COLUMN internal_notes TEXT",
        "ALTER TABLE documents ADD COLUMN discount_type VARCHAR(20)",
        "ALTER TABLE documents ADD COLUMN discount_value DECIMAL(10,2)",
    ]:
        try:
            db.session.execute(db.text(col))
            db.session.commit()
        except Exception:
            db.session.rollback()

    try:
        db.session.execute(db.text("ALTER TABLE documents ADD COLUMN sent_date DATE"))
        db.session.commit()
    except Exception:
        db.session.rollback()

    for sql in [
        "ALTER TABLE documents ADD COLUMN payment_term_days INTEGER NOT NULL DEFAULT 14",
        "ALTER TABLE conditions_templates ADD COLUMN payment_term_days INTEGER NOT NULL DEFAULT 14",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()

    for sql in [
        "ALTER TABLE documents ADD COLUMN pdf_cache BLOB",
        "ALTER TABLE documents ADD COLUMN pdf_cached_at DATETIME",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()

    try:
        db.session.execute(db.text("""
            CREATE TABLE IF NOT EXISTS company_settings (
                key VARCHAR(100) PRIMARY KEY,
                value TEXT NOT NULL DEFAULT ''
            )
        """))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("""
            CREATE TABLE IF NOT EXISTS document_status_log (
                id INTEGER PRIMARY KEY,
                document_id INTEGER NOT NULL REFERENCES documents(id),
                from_status VARCHAR(20),
                to_status VARCHAR(20) NOT NULL,
                changed_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                note TEXT
            )
        """))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("""
            CREATE TABLE IF NOT EXISTS retainers (
                id INTEGER PRIMARY KEY,
                contact_id INTEGER NOT NULL REFERENCES contacts(id),
                name VARCHAR(300) NOT NULL,
                monthly_fee FLOAT NOT NULL,
                monthly_hours FLOAT NOT NULL,
                billing_interval_months INTEGER NOT NULL DEFAULT 6,
                hourly_rate FLOAT,
                start_date DATE NOT NULL,
                end_date DATE,
                active BOOLEAN DEFAULT 1,
                last_billed_date DATE,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP
            )
        """))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("ALTER TABLE retainers ADD COLUMN last_fee_period VARCHAR(7)"))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("ALTER TABLE time_entries ADD COLUMN billed_document_id INTEGER REFERENCES documents(id)"))
        db.session.commit()
    except Exception:
        db.session.rollback()

    for sql in [
        "ALTER TABLE contacts ADD COLUMN tier VARCHAR(20)",
        "ALTER TABLE contacts ADD COLUMN pipeline_phase VARCHAR(10)",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()


    try:
        db.session.execute(db.text(
            "ALTER TABLE contacts ADD COLUMN contact_type VARCHAR(10) NOT NULL DEFAULT 'customer'"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("""
            CREATE TABLE IF NOT EXISTS expense_categories (
                id INTEGER PRIMARY KEY,
                code VARCHAR(30) UNIQUE NOT NULL,
                name VARCHAR(100) NOT NULL,
                gmbh_account_code VARCHAR(10),
                ear_group VARCHAR(50),
                is_active BOOLEAN NOT NULL DEFAULT 1,
                sort_order INTEGER NOT NULL DEFAULT 100
            )
        """))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text("""
            CREATE TABLE IF NOT EXISTS incoming_invoices (
                id INTEGER PRIMARY KEY,
                number VARCHAR(20) UNIQUE NOT NULL,
                supplier_id INTEGER NOT NULL REFERENCES contacts(id),
                expense_category_id INTEGER REFERENCES expense_categories(id),
                external_invoice_number VARCHAR(100) NOT NULL,
                subject VARCHAR(500) NOT NULL,
                invoice_date DATE NOT NULL,
                delivery_date DATE,
                posting_date DATE NOT NULL,
                due_date DATE,
                payment_date DATE,
                net_amount NUMERIC(12,2) NOT NULL DEFAULT 0,
                vat_rate NUMERIC(5,2) NOT NULL DEFAULT 20,
                vat_amount NUMERIC(12,2) NOT NULL DEFAULT 0,
                gross_amount NUMERIC(12,2) NOT NULL DEFAULT 0,
                vat_exemption_reason VARCHAR(200),
                is_vat_deductible BOOLEAN NOT NULL DEFAULT 1,
                payment_method VARCHAR(30),
                document_file_path VARCHAR(500),
                status VARCHAR(20) NOT NULL DEFAULT 'open',
                internal_notes TEXT,
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                CONSTRAINT uq_supplier_external_invoice UNIQUE (supplier_id, external_invoice_number)
            )
        """))
        db.session.commit()
    except Exception:
        db.session.rollback()

    for sql in [
        "CREATE INDEX IF NOT EXISTS idx_incoming_supplier ON incoming_invoices(supplier_id)",
        "CREATE INDEX IF NOT EXISTS idx_incoming_posting_date ON incoming_invoices(posting_date)",
        "CREATE INDEX IF NOT EXISTS idx_incoming_category ON incoming_invoices(expense_category_id)",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()

    try:
        db.session.execute(db.text("""
            CREATE TABLE IF NOT EXISTS incoming_line_items (
                id INTEGER PRIMARY KEY,
                incoming_invoice_id INTEGER NOT NULL REFERENCES incoming_invoices(id) ON DELETE CASCADE,
                position INTEGER NOT NULL DEFAULT 0,
                description VARCHAR(500),
                net_amount NUMERIC(12,2) NOT NULL DEFAULT 0,
                vat_rate NUMERIC(5,2) NOT NULL DEFAULT 20,
                vat_amount NUMERIC(12,2) NOT NULL DEFAULT 0,
                expense_category_id INTEGER REFERENCES expense_categories(id)
            )
        """))
        db.session.commit()
    except Exception:
        db.session.rollback()

    from models import NumberSequence
    if not NumberSequence.query.filter_by(prefix='E').first():
        db.session.add(NumberSequence(prefix='E', next_number=1, padding=5))
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()

    if not NumberSequence.query.filter_by(prefix='H').first():
        db.session.add(NumberSequence(prefix='H', next_number=1, padding=5))
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()

    try:
        db.session.execute(db.text(
            "ALTER TABLE incoming_invoices ADD COLUMN beleg_status VARCHAR(20) NOT NULL DEFAULT 'fehlt'"
        ))
        db.session.commit()
        db.session.execute(db.text(
            "UPDATE incoming_invoices SET beleg_status='vorhanden' "
            "WHERE document_file_path IS NOT NULL AND document_file_path != ''"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text(
            "ALTER TABLE contacts ADD COLUMN abrechnungsmodell VARCHAR(50)"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text(
            "ALTER TABLE contacts ADD COLUMN abrechnung_terminiert_bis DATE"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text(
            "ALTER TABLE contacts ADD COLUMN has_social_mandate BOOLEAN NOT NULL DEFAULT 0"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text(
            "ALTER TABLE contacts ADD COLUMN mandat_aktiv BOOLEAN NOT NULL DEFAULT 0"
        ))
        db.session.execute(db.text(
            "UPDATE contacts SET mandat_aktiv = 1 WHERE "
            "id IN (SELECT contact_id FROM retainers WHERE active = 1) "
            "OR (status IN ('won', 'active', 'aktiv') AND abrechnungsmodell IN "
            "('retainer', 'retainer_trueup', 'agentur_partner', 'pro_bono', "
            "'projekt_fix', 'stundenbasis'))"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text(
            "ALTER TABLE content_calendar ADD COLUMN outbox_job VARCHAR(80)"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text(
            "ALTER TABLE contacts ADD COLUMN archived BOOLEAN NOT NULL DEFAULT 0"
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()

    for sql in [
        """
        CREATE TABLE IF NOT EXISTS bank_imports (
            id INTEGER PRIMARY KEY,
            filename VARCHAR(300) NOT NULL,
            schema_name VARCHAR(50),
            imported_at DATETIME DEFAULT CURRENT_TIMESTAMP,
            row_count INTEGER NOT NULL DEFAULT 0,
            new_count INTEGER NOT NULL DEFAULT 0,
            duplicate_count INTEGER NOT NULL DEFAULT 0,
            period_from DATE,
            period_to DATE
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS bank_transactions (
            id INTEGER PRIMARY KEY,
            import_id INTEGER NOT NULL REFERENCES bank_imports(id),
            booking_date DATE NOT NULL,
            value_date DATE,
            amount NUMERIC(12,2) NOT NULL,
            currency VARCHAR(3) NOT NULL DEFAULT 'EUR',
            partner_name VARCHAR(300),
            partner_iban VARCHAR(50),
            reference TEXT,
            tx_type VARCHAR(50),
            row_hash VARCHAR(64) NOT NULL UNIQUE,
            match_status VARCHAR(20) NOT NULL DEFAULT 'unmatched',
            matched_document_id INTEGER REFERENCES documents(id),
            matched_incoming_invoice_id INTEGER REFERENCES incoming_invoices(id),
            confirmed_at DATETIME,
            note TEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_bank_tx_booking_date ON bank_transactions(booking_date)",
        "CREATE INDEX IF NOT EXISTS idx_bank_tx_match_status ON bank_transactions(match_status)",
        "CREATE INDEX IF NOT EXISTS idx_bank_tx_import ON bank_transactions(import_id)",
        "ALTER TABLE bank_transactions ADD COLUMN foreign_amount NUMERIC(12,2)",
        "ALTER TABLE bank_transactions ADD COLUMN foreign_currency VARCHAR(3)",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()

    for sql in [
        """
        CREATE TABLE IF NOT EXISTS claude_limit_snapshots (
            id INTEGER PRIMARY KEY,
            taken_at DATETIME NOT NULL,
            profile VARCHAR(20) NOT NULL,
            account_uuid VARCHAR(64),
            five_hour_pct REAL,
            seven_day_pct REAL,
            seven_day_resets_at DATETIME,
            error TEXT,
            anlass VARCHAR(20) NOT NULL
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_cls_taken_at ON claude_limit_snapshots(taken_at)",
        """
        CREATE TABLE IF NOT EXISTS fuell_jobs (
            id INTEGER PRIMARY KEY,
            job VARCHAR(1) NOT NULL,
            profile VARCHAR(20) NOT NULL,
            account_uuid VARCHAR(64),
            modus VARCHAR(20) NOT NULL DEFAULT 'halbautomatisch',
            started_at DATETIME NOT NULL,
            ended_at DATETIME,
            pct_before REAL,
            pct_after REAL,
            budget_punkte REAL,
            status VARCHAR(20) NOT NULL DEFAULT 'laeuft',
            ergebnis TEXT,
            cursor TEXT,
            rolle_id INTEGER,
            gegenstand_id INTEGER
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_fuell_jobs_status ON fuell_jobs(status)",
        """
        CREATE TABLE IF NOT EXISTS fuell_job_items (
            id INTEGER PRIMARY KEY,
            fuell_job_id INTEGER NOT NULL REFERENCES fuell_jobs(id),
            task_id INTEGER REFERENCES tasks(id) ON DELETE SET NULL,
            idee_id INTEGER,
            aktion VARCHAR(30) NOT NULL,
            stufe INTEGER,
            beschreibung_hash VARCHAR(64),
            kommentar_id INTEGER,
            created_at DATETIME NOT NULL
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_fuell_items_job ON fuell_job_items(fuell_job_id)",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()

    for sql in [
        """
        CREATE TABLE IF NOT EXISTS board_rollen (
            id INTEGER PRIMARY KEY,
            name VARCHAR(100) NOT NULL UNIQUE,
            art VARCHAR(10) NOT NULL,
            leitfrage TEXT NOT NULL,
            quellen TEXT,
            prozess_readme VARCHAR(300),
            passt_auf TEXT NOT NULL,
            mindestanteil_pct REAL NOT NULL DEFAULT 2,
            aktiv INTEGER NOT NULL DEFAULT 1
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS ideen_gegenstaende (
            id INTEGER PRIMARY KEY,
            art VARCHAR(30) NOT NULL,
            name VARCHAR(200) NOT NULL,
            spur VARCHAR(30) NOT NULL,
            contact_id INTEGER REFERENCES contacts(id) ON DELETE SET NULL,
            partneragentur INTEGER NOT NULL DEFAULT 0,
            zuletzt_durchdacht DATETIME,
            notiz TEXT,
            aktiv INTEGER NOT NULL DEFAULT 1,
            frist DATE
        )
        """,
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_gegenstand_kontakt ON ideen_gegenstaende(art, contact_id) "
        "WHERE contact_id IS NOT NULL",
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_gegenstand_name ON ideen_gegenstaende(art, name) "
        "WHERE contact_id IS NULL",
        """
        CREATE TABLE IF NOT EXISTS ideen (
            id INTEGER PRIMARY KEY,
            spur VARCHAR(30) NOT NULL,
            art VARCHAR(20) NOT NULL,
            contact_id INTEGER REFERENCES contacts(id) ON DELETE SET NULL,
            rolle_id INTEGER REFERENCES board_rollen(id),
            gegenstand_id INTEGER REFERENCES ideen_gegenstaende(id),
            herkunft VARCHAR(20) NOT NULL,
            herkunft_profil VARCHAR(20),
            unverifiziert INTEGER NOT NULL DEFAULT 0,
            dw_bereich VARCHAR(20),
            titel VARCHAR(500) NOT NULL,
            problem TEXT NOT NULL,
            vorschlag TEXT NOT NULL,
            test_beschreibung TEXT,
            test_kosten_h REAL,
            test_kosten_eur REAL NOT NULL DEFAULT 0,
            test_erfolgskriterium TEXT,
            test_dauer_tage INTEGER,
            wirkung_wert REAL,
            wirkung_einheit VARCHAR(10),
            wirkung_beleg TEXT,
            erfolgschance_pct REAL,
            aufwand_umsetzung_h REAL,
            frist DATE,
            prio_manuell VARCHAR(10) NOT NULL DEFAULT 'normal',
            max_bewertung INTEGER,
            max_notiz TEXT,
            status VARCHAR(30) NOT NULL DEFAULT 'offen',
            verwurf_grund TEXT,
            task_id INTEGER REFERENCES tasks(id) ON DELETE SET NULL,
            test_ergebnis VARCHAR(20),
            wirkung_ist REAL,
            wirkung_ist_am DATE,
            vorgaenger_id INTEGER REFERENCES ideen(id),
            unterschied_zur_verworfenen TEXT,
            created_at DATETIME NOT NULL,
            updated_at DATETIME NOT NULL,
            geaendert_von VARCHAR(40)
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_ideen_status ON ideen(status)",
        """
        CREATE TABLE IF NOT EXISTS ideen_verlauf (
            id INTEGER PRIMARY KEY,
            idee_id INTEGER NOT NULL REFERENCES ideen(id) ON DELETE CASCADE,
            von VARCHAR(40) NOT NULL,
            art VARCHAR(20) NOT NULL,
            text TEXT,
            created_at DATETIME NOT NULL
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_ideen_verlauf_idee ON ideen_verlauf(idee_id)",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()


    for sql in [
        """
        CREATE TABLE IF NOT EXISTS wiki_quellen_hash (
            id INTEGER PRIMARY KEY,
            quelle_pfad TEXT NOT NULL UNIQUE,
            sha256 VARCHAR(64) NOT NULL,
            seite_id INTEGER NOT NULL REFERENCES wiki_pages(id) ON DELETE CASCADE,
            gebaut_am DATETIME NOT NULL
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_wiki_quellen_seite ON wiki_quellen_hash(seite_id)",
    ]:
        try:
            db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()

    try:
        db.session.execute(db.text("""
            CREATE TABLE IF NOT EXISTS work_days (
                datum DATE PRIMARY KEY,
                minuten INTEGER NOT NULL DEFAULT 0,
                notiz TEXT,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
            )
        """))
        db.session.commit()
    except Exception:
        db.session.rollback()

    _setup_search_fts()


_SEARCH_FTS_SPECS = [
    ('contacts_fts', 'contacts', [
        'organisation', 'vorname', 'nachname',
        'mail_person', 'mail_organisation', 'kommentare',
    ]),
    ('tasks_fts', 'tasks', ['title', 'description']),
    ('documents_fts', 'documents', ['number', 'subject']),
    ('wiki_pages_fts', 'wiki_pages', ['title', 'content']),
    ('wiki_entries_fts', 'wiki_entries', ['name', 'description', 'trigger_phrases']),
]


def _setup_search_fts():
    for fts_table, content_table, cols in _SEARCH_FTS_SPECS:
        fts_cols = ', '.join(cols)
        new_vals = ', '.join(f'new.{c}' for c in cols)
        old_vals = ', '.join(f'old.{c}' for c in cols)

        already_existed = bool(db.session.execute(db.text(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=:t"
        ), {"t": fts_table}).scalar())

        try:
            db.session.execute(db.text(
                f"CREATE VIRTUAL TABLE IF NOT EXISTS {fts_table} USING fts5("
                f"{fts_cols}, content='{content_table}', content_rowid='id')"
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()
            continue

        try:
            db.session.execute(db.text(f"""
                CREATE TRIGGER IF NOT EXISTS {content_table}_fts_ai AFTER INSERT ON {content_table} BEGIN
                    INSERT INTO {fts_table}(rowid, {fts_cols}) VALUES (new.id, {new_vals});
                END
            """))
            db.session.execute(db.text(f"""
                CREATE TRIGGER IF NOT EXISTS {content_table}_fts_ad AFTER DELETE ON {content_table} BEGIN
                    INSERT INTO {fts_table}({fts_table}, rowid, {fts_cols}) VALUES ('delete', old.id, {old_vals});
                END
            """))
            db.session.execute(db.text(f"""
                CREATE TRIGGER IF NOT EXISTS {content_table}_fts_au AFTER UPDATE ON {content_table} BEGIN
                    INSERT INTO {fts_table}({fts_table}, rowid, {fts_cols}) VALUES ('delete', old.id, {old_vals});
                    INSERT INTO {fts_table}(rowid, {fts_cols}) VALUES (new.id, {new_vals});
                END
            """))
            db.session.commit()
        except Exception:
            db.session.rollback()

        if not already_existed:
            try:
                db.session.execute(db.text(
                    f"INSERT INTO {fts_table}(rowid, {fts_cols}) "
                    f"SELECT id, {fts_cols} FROM {content_table}"
                ))
                db.session.commit()
            except Exception:
                db.session.rollback()



COLUMN_MAPPING = {
    'organisation': ['organisation', 'firma', 'company', 'firmenname', 'unternehmen'],
    'organisationsform': ['organisationsform', 'rechtsform', 'firmenart', 'legal form', 'company type'],
    'website': ['website', 'web', 'url', 'homepage'],
    'mail_organisation': ['mail organisation', 'mail org', 'firmen-email', 'company email'],
    'telefon_organisation': ['telefon organisation', 'tel org', 'firmen-telefon', 'company phone'],
    'in_a_nutshell': ['in a nutshell', 'nutshell', 'beschreibung', 'description'],
    'land': ['land', 'country'],
    'bundesland': ['bundesland', 'state', 'province'],
    'plz': ['plz', 'postleitzahl', 'zip', 'postal code', 'zip code'],
    'stadt': ['stadt', 'city', 'ort'],
    'strasse': ['strasse', 'straße', 'strasse', 'street', 'adresse', 'address'],
    'sprache_1': ['sprache 1', 'sprache1', 'language 1', 'hauptsprache'],
    'sprache_2': ['sprache 2', 'sprache2', 'language 2', 'zweitsprache'],
    'vorname': ['vorname', 'first name', 'firstname'],
    'nachname': ['nachname', 'last name', 'lastname', 'name'],
    'mail_person': ['mail person', 'email', 'e-mail', 'mail'],
    'telefon_person': ['telefon person', 'telefon', 'phone', 'tel'],
    'position': ['position', 'job title', 'rolle', 'title'],
    'anrede': ['anrede', 'salutation'],
    'anrede_en': ['anrede en', 'salutation en', 'english salutation'],
    'titel_vor': ['titel vor', 'titel', 'title', 'academic title'],
    'titel_nach': ['titel nach', 'suffix', 'nachgestellter titel'],
    'linkedin': ['linkedin', 'linkedin url', 'linkedin profile'],
    'gruendungsdatum': ['gruendungsdatum', 'gründungsdatum', 'founding date', 'eintragungsdatum', 'firmenbuch datum', 'founded'],
    'steuernummer': ['steuernummer', 'tax number', 'tax id', 'ust-id', 'vat', 'steuer-nr'],
    'gemeinnuetzig': ['gemeinnuetzig', 'gemeinnützig', 'non-profit', 'nonprofit', 'non profit'],
    'kontaktiert': ['kontaktiert', 'contacted', 'contact date'],
    'wie': ['wie', 'how', 'contact method', 'kontaktart'],
    'kommentare': ['kommentare', 'comments', 'notes', 'notizen', 'anmerkungen'],
    'erinnerung': ['erinnerung', 'reminder', 'reminder date'],
    'erinnerungstext': ['erinnerungstext', 'reminder text', 'reminder note'],
    'status': ['status', 'phase', 'stage'],
    'deal_value': ['deal_value', 'deal value', 'dealwert', 'deal wert', 'wert', 'value'],
    'deal_probability': ['deal_probability', 'deal probability', 'wahrscheinlichkeit', 'probability'],
    'lost_reason': ['lost_reason', 'lost reason', 'verlustgrund', 'grund'],
    'quelle': ['quelle', 'source', 'lead source', 'herkunft'],
    'branche': ['branche', 'industry', 'branchengruppe'],
    'prioritaet': ['prioritaet', 'priorität', 'priority', 'prio'],
    'customer_number': ['customer_number', 'kundennummer', 'kunden-nr', 'customer no'],
}

UPDATABLE_FIELDS = [
    'organisation', 'organisationsform', 'website', 'mail_organisation',
    'telefon_organisation', 'in_a_nutshell', 'gruendungsdatum', 'steuernummer', 'gemeinnuetzig',
    'land', 'bundesland', 'plz',
    'stadt', 'strasse', 'sprache_1', 'sprache_2', 'vorname', 'nachname', 'mail_person',
    'telefon_person', 'position', 'anrede', 'anrede_en', 'titel_vor',
    'titel_nach', 'linkedin', 'kontaktiert', 'wie', 'kommentare',
    'erinnerung', 'erinnerungstext', 'status',
    'deal_value', 'deal_probability', 'lost_reason',
    'quelle', 'branche', 'prioritaet',
]

FLOAT_FIELDS = {'deal_value'}
INT_FIELDS = {'deal_probability'}
BOOL_FIELDS = {'gemeinnuetzig'}



def parse_date(val):
    if val is None:
        return None
    if hasattr(val, 'date'):
        return val.date() if callable(getattr(val, 'date')) else val
    if isinstance(val, str):
        from datetime import datetime
        for fmt in ['%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y']:
            try:
                return datetime.strptime(val, fmt).date()
            except ValueError:
                pass
    return None


def clean_str(val):
    if val is None:
        return None
    return str(val).strip() or None


def parse_float(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).strip().replace(',', '.'))
    except (ValueError, TypeError):
        return None


def parse_int(val):
    if val is None:
        return None
    if isinstance(val, int):
        return val
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return None


def parse_bool(val):
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    s = str(val).strip().lower()
    return s in ('1', 'true', 'yes', 'ja', 'x', 'y')


def normalize_status(status):
    if not status:
        return 'new'
    status = str(status).lower().strip()
    status_map = {
        'new': ['new', 'neu', 'open', 'offen'],
        'contacted': ['contacted', 'kontaktiert', 'called', 'angerufen'],
        'qualified': ['qualified', 'qualifiziert', 'interested', 'interessiert'],
        'proposal': ['proposal', 'angebot', 'offer'],
        'negotiation': ['negotiation', 'verhandlung'],
        'won': ['won', 'gewonnen', 'closed', 'kunde'],
        'lost': ['lost', 'verloren', 'rejected', 'abgelehnt']
    }
    for standard, variants in status_map.items():
        if status in variants:
            return standard
    return 'new'


def _parse_row_data(row, field_indices):
    data = {}
    for field, idx in field_indices.items():
        if idx < len(row) and row[idx] is not None:
            data[field] = row[idx]
    return data


def _build_contact_kwargs(data):
    date_fields = {'kontaktiert', 'erinnerung', 'gruendungsdatum'}
    kwargs = {}
    for field in UPDATABLE_FIELDS:
        val = data.get(field)
        if field in date_fields:
            kwargs[field] = parse_date(val)
        elif field == 'status':
            kwargs[field] = normalize_status(val)
        elif field in FLOAT_FIELDS:
            kwargs[field] = parse_float(val)
        elif field in INT_FIELDS:
            kwargs[field] = parse_int(val)
        elif field in BOOL_FIELDS:
            kwargs[field] = parse_bool(val)
        else:
            kwargs[field] = clean_str(val)
    return kwargs



def _build_lookup_dicts():
    all_contacts = Contact.query.all()
    email_lookup = {}
    name_lookup = {}
    org_lookup = {}

    for c in all_contacts:
        for email in [c.mail_person, c.mail_organisation]:
            if email:
                email_lookup[email.strip().lower()] = c

        if c.organisation and c.nachname and c.vorname:
            key = (c.organisation.strip().lower(),
                   c.nachname.strip().lower(),
                   c.vorname.strip().lower())
            name_lookup[key] = c

        if c.organisation:
            org_lookup[c.organisation.strip().lower()] = c

    return email_lookup, name_lookup, org_lookup


def _fuzzy_match(s1, s2, threshold=0.8):
    if not s1 or not s2:
        return False
    s1 = s1.strip().lower()
    s2 = s2.strip().lower()
    if s1 == s2:
        return True
    if s1 in s2 or s2 in s1:
        return True
    for suffix in [' gmbh & co kg', ' gmbh', ' e.u.', ' e.u', ' og', ' kg', ' ag', ' ug', ' e.v.', ' verein']:
        s1 = s1.removesuffix(suffix)
        s2 = s2.removesuffix(suffix)
    if s1 == s2:
        return True
    if len(s1) < 3 or len(s2) < 3:
        return False
    trigrams1 = set(s1[i:i+3] for i in range(len(s1) - 2))
    trigrams2 = set(s2[i:i+3] for i in range(len(s2) - 2))
    if not trigrams1 or not trigrams2:
        return False
    intersection = trigrams1 & trigrams2
    union = trigrams1 | trigrams2
    return len(intersection) / len(union) >= threshold


def _find_existing_contact(data, email_lookup, name_lookup, org_lookup=None):
    for email_field in ['mail_person', 'mail_organisation']:
        email = clean_str(data.get(email_field))
        if email:
            match = email_lookup.get(email.lower())
            if match:
                return match

    org = clean_str(data.get('organisation'))
    nachname = clean_str(data.get('nachname'))
    vorname = clean_str(data.get('vorname'))
    if org and nachname and vorname:
        key = (org.lower(), nachname.lower(), vorname.lower())
        match = name_lookup.get(key)
        if match:
            return match

    if org and org_lookup:
        match = org_lookup.get(org.lower())
        if match:
            return match

    if org and org_lookup:
        org_lower = org.lower()
        for existing_org, contact in org_lookup.items():
            if _fuzzy_match(org_lower, existing_org):
                return contact

    return None


def _update_contact_fields(contact, data):
    date_fields = {'kontaktiert', 'erinnerung', 'gruendungsdatum'}
    changed = False

    for field in UPDATABLE_FIELDS:
        raw = data.get(field)
        if raw is None:
            continue

        if field in date_fields:
            new_val = parse_date(raw)
        elif field == 'status':
            new_val = normalize_status(raw)
            if new_val == 'new':
                continue
        elif field in FLOAT_FIELDS:
            new_val = parse_float(raw)
        elif field in INT_FIELDS:
            new_val = parse_int(raw)
        elif field in BOOL_FIELDS:
            new_val = parse_bool(raw)
        else:
            new_val = clean_str(raw)

        if new_val is None and field not in BOOL_FIELDS:
            continue

        if getattr(contact, field) != new_val:
            setattr(contact, field, new_val)
            changed = True

    return changed


def import_from_excel(filepath):
    result = {'created': 0, 'updated': 0, 'skipped': 0, 'error': None}

    if not os.path.exists(filepath):
        result['error'] = f"Datei nicht gefunden: {filepath}"
        return result

    try:
        wb = load_workbook(filepath, read_only=True)
        ws = wb.active

        headers = []
        for cell in ws[1]:
            val = str(cell.value).strip().lower() if cell.value else ''
            headers.append(val)

        field_indices = {}
        for field, possible_names in COLUMN_MAPPING.items():
            for i, header in enumerate(headers):
                if header in possible_names:
                    field_indices[field] = i
                    break

        email_lookup, name_lookup, org_lookup = _build_lookup_dicts()

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(cell is None for cell in row):
                continue

            data = _parse_row_data(row, field_indices)
            if not data:
                result['skipped'] += 1
                continue

            existing = _find_existing_contact(data, email_lookup, name_lookup, org_lookup)

            if existing:
                changed = _update_contact_fields(existing, data)
                if changed:
                    interaction = Interaction(
                        contact_id=existing.id,
                        datum=date.today(),
                        typ='Import',
                        notiz='Import-Update: Kontaktdaten aktualisiert'
                    )
                    db.session.add(interaction)
                    result['updated'] += 1
                else:
                    result['skipped'] += 1
            else:
                kwargs = _build_contact_kwargs(data)
                contact = Contact(**kwargs)
                db.session.add(contact)
                db.session.flush()

                for email in [contact.mail_person, contact.mail_organisation]:
                    if email:
                        email_lookup[email.strip().lower()] = contact
                if contact.organisation and contact.nachname and contact.vorname:
                    key = (contact.organisation.strip().lower(),
                           contact.nachname.strip().lower(),
                           contact.vorname.strip().lower())
                    name_lookup[key] = contact
                if contact.organisation:
                    org_lookup[contact.organisation.strip().lower()] = contact

                result['created'] += 1

        db.session.commit()
        wb.close()

    except Exception as e:
        db.session.rollback()
        result['error'] = str(e)

    return result


def get_pipeline_stats():
    from sqlalchemy import func
    stats = db.session.query(
        Contact.status,
        func.count(Contact.id)
    ).group_by(Contact.status).all()
    return dict(stats)


STANDARD_EMAIL_TEMPLATES = [
    {
        'name': 'Erstansprache',
        'betreff': 'Digitale Sichtbarkeit für {organisation}',
        'text': '{anrede_du}\n\nich bin {mein_name} von {firma}. Ich habe mir eure Website angesehen und ein paar Stellen gefunden, an denen ihr online leichter gefunden würdet.\n\nWenn du magst, zeige ich dir das in einem kurzen Gespräch, unverbindlich und in deinem Tempo. Falls es gerade nicht passt, sag einfach kurz Bescheid.\n\nSonnige Grüße\n\n{signatur}',
    },
    {
        'name': 'Angebot nachfassen',
        'betreff': 'Rückfrage zu Angebot {angebotsnummer}',
        'text': '{anrede_du}\n\nich wollte kurz nachfragen, ob du zum Angebot {angebotsnummer} vom {angebotsdatum} schon zu einer Einschätzung gekommen bist.\n\nFalls Fragen offen sind oder etwas anders zugeschnitten sein soll, klären wir das direkt gemeinsam. Und wenn es gerade nicht passt, ist das auch eine Antwort.\n\nSonnige Grüße\n\n{signatur}',
    },
    {
        'name': 'Rechnung versendet',
        'betreff': 'Rechnung {rechnungsnummer}',
        'text': '{anrede_du}\n\nim Anhang findest du die Rechnung {rechnungsnummer}.\n\nZahlungsziel sind 14 Tage, aber falls du mehr Zeit brauchst, gib mir einfach kurz ein Zeichen.\n\nWenn dir irgendetwas auffällt oder du Fragen hast, sag einfach kurz Bescheid, das klären wir sofort.\n\nIch danke für die Zusammenarbeit.\n\nSonnige Grüße\n\n{signatur}',
    },
    {
        'name': 'Zahlungserinnerung',
        'betreff': 'Zahlungserinnerung: Rechnung {rechnungsnummer}',
        'text': '{anrede_du}\n\nkurze Erinnerung in eigener Sache: die Rechnung {rechnungsnummer} ist noch offen.\n\nFalls deine Zahlung inzwischen schon unterwegs war, ist diese Mail gegenstandslos. Und wenn etwas unklar ist oder es gerade nicht passt, sag einfach kurz Bescheid, dann klären wir das direkt gemeinsam.\n\nSonnige Grüße\n\n{signatur}',
    },
    {
        'name': 'Onboarding',
        'betreff': 'Willkommen: nächste Schritte für {organisation}',
        'text': '{anrede_du}\n\ndanke für dein Vertrauen, ich freue mich auf die Zusammenarbeit.\n\nFür den Start brauche ich diese Zugänge:\n- Google Ads (falls vorhanden)\n- Google Analytics und GTM\n- WordPress-Login\n\nSobald alles eingerichtet ist, melde ich mich. Wenn du bei einem Zugang nicht weiterkommst, sag einfach kurz Bescheid, das klären wir direkt gemeinsam.\n\nSonnige Grüße\n\n{signatur}',
    },
    {
        'name': 'Monatlicher Report',
        'betreff': 'Monatsbericht: {organisation}',
        'text': '{anrede_du}\n\nim Anhang findest du den Monatsbericht für {organisation}.\n\nWenn du die Zahlen lieber gemeinsam durchgehen willst, sag einfach kurz Bescheid, dann nehmen wir uns 20 Minuten dafür.\n\nSonnige Grüße\n\n{signatur}',
    },
]


def _seed_email_templates():
    from models import EmailTemplate
    if EmailTemplate.query.count() > 0:
        return

    for t in STANDARD_EMAIL_TEMPLATES:
        db.session.add(EmailTemplate(**t))
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()


def _seed_tool_types():
    from models import AccountToolType
    if AccountToolType.query.first():
        return

    defaults = [
        ('Google Tag Manager', 'GTM', 'tracking', 'https://tagmanager.google.com/', 10),
        ('Google Analytics 4', 'GA4', 'tracking', 'https://analytics.google.com/', 20),
        ('Meta Pixel', 'Pixel', 'tracking', 'https://business.facebook.com/events_manager', 30),
        ('Google Search Console', 'GSC', 'seo', 'https://search.google.com/search-console', 40),
        ('Bing Webmaster Tools', 'Bing WMT', 'seo', 'https://www.bing.com/webmasters', 50),
        ('Google Unternehmensprofil', 'GBP', 'seo', 'https://business.google.com/', 60),
        ('Google Ads', 'GAds', 'ads', 'https://ads.google.com/', 70),
        ('Google Ad Grants', 'Grants', 'ads', 'https://ads.google.com/', 80),
        ('Microsoft Advertising', 'MS Ads', 'ads', 'https://ads.microsoft.com/', 90),
        ('Meta Business Suite', 'Meta BM', 'ads', 'https://business.facebook.com/', 100),
        ('Facebook-Seite', 'FB', 'social', 'https://www.facebook.com/', 110),
        ('Instagram Business', 'IG', 'social', 'https://www.instagram.com/', 120),
        ('Looker Studio', 'Looker', 'reporting', 'https://lookerstudio.google.com/', 130),
        ('Website Backend', 'Website', 'cms', '', 140),
        ('Newsletter Tool', 'Newsletter', 'email', '', 150),
    ]
    for name, short, cat, url, order in defaults:
        db.session.add(AccountToolType(
            name=name, short_name=short, category=cat,
            default_url=url, sort_order=order
        ))
    db.session.commit()


def _seed_expense_categories():
    from models import ExpenseCategory
    if ExpenseCategory.query.first():
        return

    defaults = [
        ('software', 'Software / Cloud-Abos', '7320', 'Betriebsausgaben', 10),
        ('hosting', 'Hosting, Domains', '7320', 'Betriebsausgaben', 20),
        ('werbung', 'Werbung / Marketing', '7700', 'Betriebsausgaben', 30),
        ('reisekosten', 'Reisekosten', '7350', 'Betriebsausgaben', 40),
        ('fortbildung', 'Fortbildung, Buecher, Konferenzen', '7780', 'Betriebsausgaben', 50),
        ('subunternehmer', 'Fremdleistungen / Subunternehmer', '5700', 'Wareneinsatz', 60),
        ('hardware', 'Hardware / GWG', '7400', 'Betriebsausgaben', 70),
        ('bueromaterial', 'Bueromaterial, Druck', '7600', 'Betriebsausgaben', 80),
        ('verwaltung', 'WKO-Umlage, Steuerberater, Versicherung', '7750', 'Betriebsausgaben', 90),
        ('miete', 'Miete, Nebenkosten', '7100', 'Betriebsausgaben', 100),
        ('sonstiges', 'Sonstiges', '7800', 'Betriebsausgaben', 110),
    ]
    for code, name, account, group, order in defaults:
        db.session.add(ExpenseCategory(
            code=code, name=name, gmbh_account_code=account,
            ear_group=group, sort_order=order
        ))
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
